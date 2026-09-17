"""Download the accepted answers of one immutable human-baseline version."""
from __future__ import annotations

from io import BytesIO
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .compare_statistics import DIMENSION_NAMES


def _text(value: str) -> str:
    """Keep control characters visible, as in the existing OOXML exporters."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]",
                  lambda match: f"\\u{ord(match.group()):04x}", value)


def _append(ws, values) -> None:
    ws.append([_text(value) if isinstance(value, str) else value for value in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            # Explicit strings preserve answers such as '=SUM(...)' without
            # creating executable formulas or altering their original text.
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _json(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def export_baseline_answers(baseline: dict) -> bytes:
    """Build a wide answer sheet from saved cases, labels and human states."""
    wb = Workbook()
    ws = wb.active
    ws.title = "人工基准答案"
    products = baseline["products"]
    labels = {(r["case_key"], r["product_id"], r["dimension_id"]): r
              for r in baseline["labels"]}
    states = {r["case_key"]: r for r in baseline["states"]}
    headers = ["case_id", "query", "公共背景", "提问原图摘要", "场景", "会话ID", "轮次"]
    for n, _ in enumerate(products, 1):
        headers.extend(f"{field}_产品{n}" for field in
                       ("产品名称", "源回答", "源背景", "采集ID", "源摘要", "证据引用"))
        for name in DIMENSION_NAMES.values():
            headers.extend(f"{field}_产品{n}_{name}" for field in
                           ("人工分", "人工评分状态", "人工批注", "人工标注阶段", "人工评分来源"))
        headers.extend((f"人工响应Gate_产品{n}", f"人工安全Gate_产品{n}"))
    headers.extend(f"人工是否适用_{name}" for name in DIMENSION_NAMES.values())
    _append(ws, headers)
    for case in baseline["cases"]:
        case_key = case["case_key"]
        state = states.get(case_key, {})
        row = [case["case_id"], case.get("query", ""), case.get("context", ""),
               _json(case.get("query_hashes", [])), case.get("category", ""),
               case.get("session_group", ""), case.get("turn_index", "")]
        for product in products:
            pid = product["product_id"]
            response = case.get("responses", {}).get(pid, {})
            row.append(product["display_name"])
            row.extend(response.get(field, "") for field in
                       ("answer", "context", "response_id", "sha256", "evidence"))
            for dimension in DIMENSION_NAMES:
                label = labels.get((case_key, pid, dimension), {})
                status = label.get("score_status", "unlabeled")
                score = label.get("score") if status == "scored" else (
                    None if status == "unlabeled" else "N/A")
                row.extend((score, status, label.get("reason"), label.get("annotation_stage"),
                            _json(label["source"]) if label.get("source") else None))
            gates = state.get("gates", {}).get(pid, {})
            row.extend((gates.get("response"), gates.get("safety")))
        row.extend(state.get("applicability", {}).get(dimension) for dimension in DIMENSION_NAMES)
        _append(ws, row)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for column, header in enumerate(headers, 1):
        width = 38 if any(token in header for token in ("query", "背景", "回答", "来源", "批注", "引用")) else 22
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 42
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285A78")

    notes = wb.create_sheet("填写说明")
    metadata = dict(answer_export_version="human-answers-1.0", products=products,
                    human_standard_version=baseline["human_standard_version"])
    note_rows = [
        ["项目", "说明"],
        ["模板元信息", _json(metadata)],
        ["人工基准", baseline["name"]],
        ["基准编号", baseline["baseline_id"]],
        ["基准版本", baseline["version"]],
        ["人工标准", baseline["human_standard_version"]],
        ["人工口径", baseline.get("human_policy_note", "")],
        ["用途", baseline.get("purpose", "")],
        ["产品映射", _json(products)],
        ["答案范围", "该版本已保存的全部 Case 与已接受标签；不包含导入时被排除的评分。"],
        ["分数规则", "数字为最终有效人工分（复核优先，保留 0）；N/A 按评分状态区分原因；未标注分数留空，不补 0。"],
        ["人工评分状态", "scored=已评分；not_applicable=不适用；unverifiable=无法核验；gate_blocked=Gate阻断；unlabeled=未标注；na_unspecified=NA原因未知。"],
        ["Gate/适用性", "保留已保存的显式状态；空白代表未提供。内容准确性人工数字标签仅留档，不计人机数字一致率。"],
        ["人工标签摘要", baseline.get("labels_sha256", "")],
        ["题目目录摘要", baseline.get("case_catalog_sha256", "")],
        ["源文件摘要", baseline.get("source_file_sha256", "")],
        ["导入排除记录", _json(baseline.get("excluded_records", []))],
    ]
    for row in note_rows:
        _append(notes, row)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 100
    notes.freeze_panes = "B2"
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()

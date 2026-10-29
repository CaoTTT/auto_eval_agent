"""Deterministic wide-sheet import with physical column addresses and provenance."""
from __future__ import annotations

from collections import Counter
from contextlib import contextmanager
import hashlib
import json
import math
from pathlib import Path
import re
import uuid
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .human_baselines import (HumanError, HumanScoreLabel, ImportMapping, digest, human_protocol,
                              human_dimensions, human_dimension_names, CHAT_NAMES)

MAX_UPLOAD = 20 * 1024 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 512


@contextmanager
def workbooks(path: Path):
    try:
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 4000 or sum(i.file_size for i in archive.infolist()) > 160 * 1024 * 1024:
                raise HumanError("工作簿解压后过大", 400)
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                yield formulas, cached
            finally:
                cached.close()
        finally:
            formulas.close()
    except HumanError:
        raise
    except Exception as exc:
        raise HumanError(f"无法解析 XLSX：{type(exc).__name__}", 400) from exc


def _check_sheet(ws):
    if ws.sheet_state != "visible":
        raise HumanError("不能导入隐藏工作表", 400)
    if ws.max_row is None or ws.max_column is None:
        # Our exported templates and other valid XLSX files can omit the
        # optional worksheet dimension. Read-only openpyxl then needs a scan
        # before headers and bounds can be read reliably.
        ws.calculate_dimension(force=True)
    if (ws.max_row or 0) > MAX_ROWS + 30 or (ws.max_column or 0) > MAX_COLUMNS:
        raise HumanError(f"工作表上限为 {MAX_ROWS} 行、{MAX_COLUMNS} 列", 400)


def _headers(ws, row: int) -> list[str]:
    return [str(c.value or "") for c in next(ws.iter_rows(min_row=row, max_row=row, max_col=ws.max_column or 1))]


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _workbook_metadata(workbook) -> dict:
    if "填写说明" in workbook:
        for row in workbook["填写说明"].iter_rows(max_row=40, max_col=2, values_only=True):
            if row[0] == "模板元信息":
                try:
                    metadata = json.loads(row[1])
                    return metadata if isinstance(metadata, dict) else {}
                except (ValueError, TypeError):
                    return {}
    return {}


def suggest_mapping(headers: list[str], sheet_name: str, header_row: int = 1, metadata: dict | None = None) -> dict:
    metadata = metadata or {}
    protocol = human_protocol(metadata.get("human_standard_version") or "0.2-simplified", metadata.get("human_standard_id", ""), metadata.get("human_bundle_revision") or None)
    dimension_names = human_dimension_names(protocol)
    if not metadata.get("human_standard_version"):
        dimension_names = {**dimension_names, **CHAT_NAMES}
    cols = {h: get_column_letter(i + 1) for i, h in enumerate(headers) if headers.count(h) == 1}
    products = metadata.get("products") or [{"product_id": "product1", "display_name": "产品1"},
                                          {"product_id": "product2", "display_name": "产品2"}]
    if any("产品3_" in h for h in headers) and len(products) == 2:
        products.append({"product_id": "product3", "display_name": "产品3"})
    labels, responses, gates, inputs = [], {}, {}, {}
    for n, product in enumerate(products, 1):
        pid = product["product_id"]
        responses[pid] = {field: cols[name] for field, name in (
            ("answer", f"源回答_产品{n}"), ("context", f"源背景_产品{n}"),
            ("response_id", f"采集ID_产品{n}"), ("sha256", f"源摘要_产品{n}"),
            ("evidence", f"证据引用_产品{n}")) if name in cols}
        gates[pid] = {g: cols[f"人工{label}Gate_产品{n}"] for g, label in (("response", "响应"), ("safety", "安全"))
                      if f"人工{label}Gate_产品{n}" in cols}
        if f"人工输入状态_产品{n}" in cols:
            inputs[pid] = cols[f"人工输入状态_产品{n}"]
        for dim, label in dimension_names.items():
            mapping = {field: cols[name] for field, name in (
                ("score", f"人工分_产品{n}_{label}"), ("review", f"人工复核分_产品{n}_{label}"),
                ("status", f"人工评分状态_产品{n}_{label}"), ("reason", f"人工批注_产品{n}_{label}"),
                ("applicable", f"人工是否适用_产品{n}_{label}"),
                ("verification_status", f"人工核验状态_产品{n}_{label}")) if name in cols}
            if any(mapping.get(k) for k in ("score", "review", "status")):
                labels.append(dict(product_id=pid, dimension_id=dim, **mapping))
    # The legacy adapter is only suggested when its physical headers still look human-authored.
    legacy = [("understanding", "AW", "AX", "", ""), ("accuracy", "BJ", "BK", "", ""),
              ("service_closure", "BU", "BV", "", ""), ("scenario_fulfillment", "CF", "CG", "", ""),
              ("intuitive_efficiency", "CU", "CV", "CZ", "DA"),
              ("evidence_quality", "DK", "DM", "DR", "DS"), ("guided_recommendation", "ED", "EE", "", "")]
    if protocol.standard_id == "qa_competitor_compare" and not labels and sheet_name == "合并标注结果-review后":
        addresses = [c for spec in legacy for c in spec[1:] if c]
        if all(column_index_from_string(c) <= len(headers) and
               re.search(r"人工|review|复核", headers[column_index_from_string(c) - 1], re.I) for c in addresses):
            products = [{"product_id": "p_xiaoyi", "display_name": "小艺"}, {"product_id": "p_doubao", "display_name": "豆包"}]
            for dim, a, b, ar, br in legacy:
                labels.extend([dict(product_id=products[n]["product_id"], dimension_id=dim, score=s, review=r)
                               for n, (s, r) in enumerate(((a, ar), (b, br)))])
            responses = {products[0]["product_id"]: dict(evidence="C", context="E", answer="F"),
                         products[1]["product_id"]: dict(evidence="D", context="G", answer="H")}
            gates = {}
    return dict(sheet_name=sheet_name, header_row=header_row, case_id=cols.get("case_id", cols.get("query_id", "")),
                query=cols.get("query", cols.get("题目", "")), context=cols.get("公共背景", ""),
                query_hashes=cols.get("提问原图摘要", ""), category=cols.get("场景", ""), products=products,
                session_group=cols.get("会话ID", ""), turn_index=cols.get("轮次", ""),
                history_prefix_sha256=cols.get("历史前缀指纹", ""),
                labels=labels, responses=responses, gates=gates, inputs=inputs,
                applicability={d: cols[f"人工是否适用_{name}"] for d, name in dimension_names.items() if f"人工是否适用_{name}" in cols},
                answer_text_origin="source" if metadata.get("template_version") else "unknown",
                context_origin="source" if metadata.get("template_version") else "unknown",
                na_status="na_unspecified", human_standard_version=metadata.get("human_standard_version", ""),
                human_standard_id=metadata.get("human_standard_id", ""), human_bundle_revision=metadata.get("human_bundle_revision", ""),
                human_result_schema_version=metadata.get("human_result_schema_version") or "",
                policy_note="", purpose="regression", header_signature=digest(headers))


def inspect_workbook(path: Path, sheet_name: str = "", header_row: int = 1) -> dict:
    with workbooks(path) as (wb, cached):
        sheets = [s.title for s in wb if s.sheet_state == "visible" and s.title != "填写说明"]
        if not sheets:
            raise HumanError("没有可导入的可见工作表", 400)
        selected = sheet_name or ("合并标注结果-review后" if "合并标注结果-review后" in sheets else sheets[0])
        if selected not in sheets or not 1 <= header_row <= 30:
            raise HumanError("工作表或表头行无效", 400)
        ws = wb[selected]
        _check_sheet(ws)
        headers = _headers(ws, header_row)
        metadata = _workbook_metadata(cached)
        return dict(sheets=sheets, sheet_name=selected, header_row=header_row,
                    columns=[dict(column=get_column_letter(i + 1), header=h) for i, h in enumerate(headers)],
                    sample=[[str(c or "") for c in row] for row in cached[selected].iter_rows(
                        min_row=header_row + 1, max_row=header_row + 3, max_col=len(headers), values_only=True)],
                    suggested_mapping=suggest_mapping(headers, selected, header_row, metadata))


def _number(value, minimum: int, maximum: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ValueError("分数必须为整数")
    if isinstance(value, str) and not re.fullmatch(r"[+-]?\d+(?:\.0+)?", value.strip()):
        raise ValueError("分数格式错误")
    number = float(value)
    if not math.isfinite(number) or not number.is_integer() or not minimum <= number <= maximum:
        raise ValueError(f"分数必须在 {minimum}—{maximum} 之间")
    return int(number)


def _knowledge_number(value, dimension: str) -> int | float:
    """Read native Excel numeric labels without coercing half marks or bools."""
    from ..judges.knowledge_qa_schema import KNOWLEDGE_QA_ALLOWED_SCORES
    allowed = KNOWLEDGE_QA_ALLOWED_SCORES.get(dimension, ())
    if type(value) is int and value in allowed:
        return value
    if dimension == "accuracy" and type(value) is float and value == 1.5:
        return value
    raise ValueError("知识问答准确性只接受数字0、1、1.5、2；其他维度只接受数字0、1、2，不接受文本")


def parse_human_labels(path: Path, mapping: ImportMapping) -> dict:
    protocol = human_protocol(mapping.human_standard_version, mapping.human_standard_id, mapping.human_bundle_revision or None)
    timeliness = protocol.standard_id == "timeliness_compare"
    knowledge = protocol.standard_id == "knowledge_qa_compare"
    chat = protocol.standard_id in ("qa_chat_compare", "timeliness_compare", "knowledge_qa_compare")
    catalog = "catalog_" + uuid.uuid4().hex[:16]
    cases, labels, states, issues = [], [], [], []
    all_ids = []
    total_rows = 0
    with workbooks(path) as (wb, cache):
        if mapping.sheet_name not in wb:
            raise HumanError("工作表不存在", 400)
        ws, cs = wb[mapping.sheet_name], cache[mapping.sheet_name]
        _check_sheet(ws)
        headers = _headers(ws, mapping.header_row)
        metadata = _workbook_metadata(cache)
        if metadata.get("human_standard_version"):
            recorded_protocol = human_protocol(metadata["human_standard_version"], metadata.get("human_standard_id", ""),
                                               metadata.get("human_bundle_revision") or None)
            frozen_v2 = any(s in ("qa-chat-result-2", "timeliness-result-1", "timeliness-result-2", "timeliness-result-3", "knowledge-qa-result-2") for s in
                            (recorded_protocol.result_schema_version, protocol.result_schema_version))
            if frozen_v2 and (recorded_protocol.id, recorded_protocol.bundle_revision) != (protocol.id, protocol.bundle_revision):
                raise HumanError("模板或基准的冻结标准与所选标准不同；不能仅改版本复用评分，请生成新标准空白模板重新标注")
            if metadata.get("human_result_schema_version") and metadata["human_result_schema_version"] != protocol.result_schema_version:
                raise HumanError("模板或基准结果Schema与所选标准冲突")
        if mapping.header_signature and digest(headers) != mapping.header_signature:
            raise HumanError("表头已改变，不能复用原列地址；请重新映射", 409)
        physical_columns = [mapping.case_id, mapping.query, mapping.context, mapping.query_hashes, mapping.category,
                            mapping.session_group, mapping.turn_index, mapping.history_prefix_sha256,
                            *mapping.applicability.values(), *mapping.inputs.values()]
        for label in mapping.labels:
            physical_columns.extend(getattr(label, k) for k in ("score", "review", "status", "reason", "applicable", "verification_status"))
        for fields in (*mapping.responses.values(), *mapping.gates.values()):
            physical_columns.extend(fields.values())
        for col in filter(None, physical_columns):
            if not re.fullmatch(r"[A-Z]{1,3}", col) or column_index_from_string(col) > len(headers):
                raise HumanError(f"无效物理列地址：{col}")
        for label in mapping.labels:
            for col in (label.score, label.review):
                if col and re.search(r"模型|一致率|胜负|Gate|适用性", headers[column_index_from_string(col) - 1], re.I):
                    raise HumanError(f"{col} 列不是人工数字评分列")
        mapping = mapping.model_copy(update={"header_signature": digest(headers)})
        formula_rows = ws.iter_rows(min_row=mapping.header_row + 1, max_col=len(headers))
        cached_rows = cs.iter_rows(min_row=mapping.header_row + 1, max_col=len(headers))
        for row_no, (raw, cached) in enumerate(zip(formula_rows, cached_rows), mapping.header_row + 1):
            if all(_blank(c.value) for c in raw):
                continue
            total_rows += 1
            if total_rows > MAX_ROWS:
                raise HumanError("数据行数超过上限", 400)
            def get(col):
                if not col:
                    return None
                i = column_index_from_string(col) - 1
                cell = raw[i]
                if cell.data_type == "f":
                    if cached[i].value is None:
                        raise ValueError(f"{col}{row_no} 公式无缓存；请另存计算结果")
                    return cached[i].value
                if cell.data_type == "e":
                    raise ValueError(f"{col}{row_no} Excel 错误值")
                return cell.value
            def issue(message, col="", pid="", dim=""):
                issues.append(dict(sheet=mapping.sheet_name, row=row_no, column=col, case_id=case_id,
                                   product_id=pid, dimension_id=dim, message=str(message)))
            case_id = ""
            try:
                value = get(mapping.case_id)
                if not isinstance(value, str) or not value.strip():
                    raise ValueError("题号必须以非空文本保存（数值题号请在 Excel 中显式转为文本，保留前导零）")
                case_id = value
                all_ids.append(case_id)
                case_key = catalog + ":" + case_id
                query_hashes = get(mapping.query_hashes)
                query_hashes = json.loads(query_hashes) if query_hashes else []
                if not isinstance(query_hashes, list) or not all(isinstance(h, str) for h in query_hashes):
                    raise ValueError("提问原图摘要须为 JSON 字符串数组")
                case = dict(case_key=case_key, case_id=case_id, query=str(get(mapping.query) or ""),
                            context=str(get(mapping.context) or ""), query_hashes=query_hashes,
                            context_known=bool(mapping.context),
                            session_group=str(get(mapping.session_group) or ""),
                            history_prefix_sha256=str(get(mapping.history_prefix_sha256) or ""),
                            turn_index="" if get(mapping.turn_index) is None else str(get(mapping.turn_index)),
                            context_origin=mapping.context_origin, category=str(get(mapping.category) or ""), responses={})
                for product in mapping.products:
                    pid = product.product_id
                    response = {k: str(get(c) or "") for k, c in mapping.responses.get(pid, {}).items()}
                    response["answer_text_origin"] = mapping.answer_text_origin
                    response["context_origin"] = mapping.context_origin
                    response["response_id"] = response.get("response_id") or ""
                    case["responses"][pid] = response
                cases.append(case)
            except (ValueError, TypeError) as exc:
                issue(exc, mapping.case_id)
                continue
            shared, gates, inputs = {}, {}, {}
            try:
                for dim, col in mapping.applicability.items():
                    v = get(col)
                    token = str(v).strip().lower() if not _blank(v) else "unknown"
                    if token not in ("true", "false", "是", "否", "适用", "不适用", "unknown"):
                        raise ValueError(f"{col}{row_no} 适用性值无效")
                    shared[dim] = None if token == "unknown" else token in ("true", "是", "适用")
                for pid, cols in mapping.gates.items():
                    gates[pid] = {}
                    for gate, col in cols.items():
                        v = get(col)
                        token = str(v).strip().lower() if not _blank(v) else "unknown"
                        token = {"通过": "pass", "不通过": "fail", "不确定": "unclear"}.get(token, token)
                        if token not in ("pass", "fail", "unclear", "unknown"):
                            raise ValueError(f"{col}{row_no} Gate 值无效")
                        gates[pid][gate] = token
                for pid, col in mapping.inputs.items():
                    value = get(col)
                    token = str(value).strip().lower() if not _blank(value) else "unknown"
                    if token not in ("complete", "partial", "failed", "unknown"):
                        raise ValueError(f"{col}{row_no} 产品输入状态无效")
                    inputs[pid] = token
                states.append(dict(case_key=case_key, applicability=shared, gates=gates, inputs=inputs,
                                   source=dict(sheet=mapping.sheet_name, row=row_no,
                                               applicability=mapping.applicability, gates=mapping.gates)))
            except ValueError as exc:
                issue(exc)
                cases.pop()
                continue
            start = len(labels)
            for col in mapping.labels:
                try:
                    review, explicit = get(col.review), get(col.status)
                    use_review = not _blank(review)
                    try:
                        base = get(col.score)
                    except ValueError:
                        if not use_review:
                            raise
                        base = str(raw[column_index_from_string(col.score)-1].value)
                    value, chosen = (review, col.review) if use_review else (base, col.score)
                    status = "unlabeled" if _blank(value) else (mapping.na_status if str(value).strip().upper() in ("N/A", "NA") else "scored")
                    if chat and not knowledge and status == "scored" and type(value) is not int:
                        raise ValueError(f"人工分须为{protocol.score_min}—{protocol.score_max}严格整数，不接受文本或小数")
                    score = ((_knowledge_number(value, col.dimension_id) if knowledge else
                              _number(value, protocol.score_min, protocol.score_max)) if status == "scored" else None)
                    if not _blank(explicit):
                        explicit = {"不适用": "not_applicable", "无法核验": "unverifiable", "Gate阻断": "gate_blocked",
                                    "输入失败": "input_failed", "未标注": "unlabeled", "已评分": "scored"}.get(str(explicit).strip(), str(explicit).strip())
                        if timeliness:
                            explicit = {"gate_blocked": "blocked_gate", "证据不足": "insufficient_evidence",
                                        "理解阻断": "blocked_understanding"}.get(explicit, explicit)
                        if knowledge:
                            explicit = {"证据不足": "insufficient_evidence"}.get(explicit, explicit)
                        statuses = (("scored", "not_applicable", "insufficient_evidence", "unverifiable",
                                     "input_failed", "unlabeled", "na_unspecified") if knowledge else
                                    ("scored", "not_applicable", "insufficient_evidence", "blocked_gate",
                                     "blocked_understanding", "input_failed", "unlabeled", "na_unspecified") if timeliness else
                                    ("scored", "not_applicable", "unverifiable", "gate_blocked", "input_failed", "unlabeled", "na_unspecified"))
                        if explicit not in statuses:
                            raise ValueError("人工评分状态无效")
                        if (score is not None and explicit != "scored") or (score is None and explicit == "scored"):
                            raise ValueError("数字评分与显式状态冲突")
                        if status not in ("unlabeled", "na_unspecified", explicit):
                            raise ValueError("NA 转换规则与显式状态冲突")
                        status = explicit
                    pg = gates.get(col.product_id, {})
                    blocked = pg.get("response") == "fail" or (protocol.require_response_pass and pg.get("response") == "unclear") or (
                        not protocol.require_response_pass and pg.get("safety") == "fail")
                    applicable_raw = get(col.applicable)
                    token = str(applicable_raw).strip().lower() if not _blank(applicable_raw) else "unknown"
                    if token not in ("true", "false", "是", "否", "适用", "不适用", "unknown"):
                        raise ValueError("产品适用性值无效")
                    applicable = None if token == "unknown" else token in ("true", "是", "适用")
                    if applicable is None and not col.applicable:
                        applicable = shared.get(col.dimension_id)
                    inferred = True if status == "scored" else False if status == "not_applicable" else None
                    if applicable is not None and inferred is not None and applicable != inferred:
                        raise ValueError("人工状态与产品适用性冲突")
                    if applicable is None:
                        applicable = inferred
                    if knowledge and col.dimension_id != "multimedia_value" and (
                            status == "not_applicable" or applicable is False):
                        raise ValueError("知识问答只有富媒体增益允许不适用，前三维缺证须用证据不足或无法核验")
                    if chat and status in ("unverifiable", "insufficient_evidence") and applicable is False:
                        raise ValueError("无法核验状态不能明确不适用")
                    verification_raw = get(col.verification_status)
                    verification = str(verification_raw).strip().lower() if not _blank(verification_raw) else None
                    if verification not in (None, "verified", "partial", "unverifiable", "not_required"):
                        raise ValueError("人工核验状态无效")
                    if chat and ((status == "scored" and verification == "unverifiable") or
                                 (status == "not_applicable" and verification not in (None, "not_required")) or
                                 (status in ("unverifiable", "insufficient_evidence") and verification not in (None, "unverifiable"))):
                        raise ValueError("人工评分状态与核验状态冲突")
                    if timeliness:
                        if blocked and status not in ("blocked_gate", "unlabeled", "na_unspecified", "input_failed"):
                            raise ValueError("Gate失败后应标为 blocked_gate，不继续评分或判不适用")
                        if col.dimension_id == "understanding" and (status in ("not_applicable", "blocked_understanding") or applicable is False):
                            raise ValueError("理解需求不允许不适用或理解阻断")
                        if status == "scored" and verification == "not_required":
                            raise ValueError("时效性有效评分不能标为无需核验")
                        if col.dimension_id != "understanding" and verification == "partial" and score != 1:
                            raise ValueError("权威性/时效性 partial 仅能对应决定性核心1分")
                        if status in ("blocked_gate", "blocked_understanding") and (applicable is not None or verification is not None):
                            raise ValueError("前置阻断后适用性与核验状态应留空")
                        if status == "blocked_gate" and pg and not blocked and all(pg.get(g) in ("pass", "fail", "unclear") for g in ("response", "safety")):
                            raise ValueError("Gate阻断与已标注Gate不一致")
                        if status == "blocked_understanding" and blocked:
                            raise ValueError("Gate失败优先于理解阻断")
                    if score is not None and (applicable is False or blocked or inputs.get(col.product_id) == "failed"):
                        raise ValueError("人工数字分与适用性/Gate 冲突")
                    if not chat and status == "not_applicable" and shared.get(col.dimension_id) is True:
                        raise ValueError("不适用标签与共享适用性冲突")
                    reason = str(get(col.reason) or "")
                    if chat and use_review and str(review).strip().upper() in ("NA", "N/A") and not _blank(base) and not reason.strip():
                        raise ValueError("复核NA覆盖初始评分必须填写原因")
                    source = dict(sheet=mapping.sheet_name, row=row_no, base_cell=f"{col.score}{row_no}" if col.score else "",
                                  review_cell=f"{col.review}{row_no}" if col.review else "", selected_cell=f"{chosen}{row_no}" if chosen else "",
                                  base_raw=base, review_raw=review, applicable_cell=f"{col.applicable}{row_no}" if col.applicable else "",
                                  applicable_raw=applicable_raw, verification_cell=f"{col.verification_status}{row_no}" if col.verification_status else "",
                                  verification_raw=verification_raw, formula=any(raw[column_index_from_string(c)-1].data_type == "f" for c in (col.score, col.review) if c))
                    labels.append(HumanScoreLabel.model_validate(dict(case_key=case_key, product_id=col.product_id,
                        response_id=case["responses"][col.product_id].get("response_id") or digest([case_key, col.product_id])[:24],
                        dimension_id=col.dimension_id, score_status=status, score=score, applicable=applicable, verification_status=verification,
                        annotation_stage="review" if use_review else "base", reason=reason, source=source),
                        context={"standard_id": protocol.standard_id}).model_dump())
                except (ValueError, TypeError) as exc:
                    issue(exc, col.review or col.score or col.status, col.product_id, col.dimension_id)
            for dim in (() if chat else human_dimensions(protocol)):
                group = [r for r in labels[start:] if r["dimension_id"] == dim]
                if any(r["score_status"] == "not_applicable" for r in group) and any(r["score_status"] == "scored" for r in group):
                    for r in group:
                        r["resolution"] = "pending_conflict"
                    issue("不同产品的不适用与数字分冲突", dim=dim)
            if timeliness:
                for product in mapping.products:
                    group = [r for r in labels[start:] if r["product_id"] == product.product_id]
                    understanding = next((r for r in group if r["dimension_id"] == "understanding"), None)
                    for r in group:
                        if r["dimension_id"] == "understanding" or not understanding:
                            continue
                        u_score = understanding.get("score")
                        conflict = (u_score == 1 and r["score_status"] not in
                                    ("blocked_understanding", "unlabeled", "na_unspecified", "input_failed")) or (
                                    u_score in (2, 3) and r["score_status"] == "blocked_understanding")
                        if conflict:
                            r["resolution"] = "pending_conflict"
                            issue("理解分与后续维度阻断状态冲突", pid=product.product_id, dim=r["dimension_id"])
    duplicate_ids = {key for key, count in Counter(all_ids).items() if count > 1}
    duplicate_keys = {c["case_key"] for c in cases if c["case_id"] in duplicate_ids}
    for key in sorted(duplicate_ids):
        issues.append(dict(sheet=mapping.sheet_name, case_id=key, column=mapping.case_id, message="重复题号；所有重复行均排除"))
    cases = [c for c in cases if c["case_key"] not in duplicate_keys]
    labels = [r for r in labels if r["case_key"] not in duplicate_keys]
    states = [r for r in states if r["case_key"] not in duplicate_keys]
    result = dict(mapping=mapping.model_dump(), cases=cases, labels=labels, states=states, issues=issues,
                  source_file_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                  summary=dict(total_rows=total_rows, case_count=len(cases), label_count=len(labels),
                               statuses=dict(Counter(r["score_status"] for r in labels)), issue_count=len(issues)))
    result["preview_sha256"] = digest(result)
    return result


def build_human_template(snapshot: dict, standard: str) -> bytes:
    from .human_compare import case_identity, response_identity
    from .human_report_export import tables_workbook
    from .compare_statistics import StatisticsTable
    protocol = human_protocol(standard)
    if protocol.standard_id in ("timeliness_compare", "knowledge_qa_compare"):
        from .protocol_identity import recover_protocol_identity, require_snapshot_protocol
        recovered = recover_protocol_identity(snapshot)
        if recovered["executable"] and recovered["evaluation_profile"] == protocol.id:
            # A template downloaded from a historical task must still align
            # with that task's frozen model result contract.
            protocol = require_snapshot_protocol(snapshot)
    dimensions = human_dimensions(protocol, scores=True)
    names_by_dimension = human_dimension_names(protocol)
    timeliness = protocol.standard_id == "timeliness_compare"
    knowledge = protocol.standard_id == "knowledge_qa_compare"
    chat = protocol.standard_id in ("qa_chat_compare", "timeliness_compare", "knowledge_qa_compare")
    detailed_states = chat and not knowledge
    items = snapshot.get("items") or []
    count = max((int(i.get("product_count") or (3 if i.get("video3") or i.get("screenshot3") or i.get("answer3") else 2)) for i in items), default=2)
    names = (snapshot.get("options") or {}).get("product_names") or []
    products = [dict(product_id=f"product{n}", display_name=names[n-1] if n <= len(names) and names[n-1] else f"产品{n}") for n in range(1, count+1)]
    headers = ["case_id", "query", "公共背景", "提问原图摘要", "场景", "会话ID", "轮次", "历史前缀指纹"]
    for n in range(1, count+1):
        headers.extend(f"{name}_产品{n}" for name in ("源回答", "源背景", "采集ID", "源摘要", "证据引用"))
        for dim in dimensions:
            headers.extend(f"{name}_产品{n}_{names_by_dimension[dim]}" for name in
                           (("人工分", "人工复核分", "人工评分状态", "人工批注", "人工是否适用", "人工核验状态") if detailed_states else
                            ("人工分", "人工复核分", "人工评分状态", "人工批注")))
        if not knowledge:
            headers.extend((f"人工响应Gate_产品{n}", f"人工安全Gate_产品{n}"))
        if chat:
            headers.append(f"人工输入状态_产品{n}")
    if not chat:
        headers.extend(f"人工是否适用_{names_by_dimension[d]}" for d in dimensions)
    table = StatisticsTable("人工标注模板", tuple(headers))
    for item in items:
        identity = case_identity(item)
        source = item.get("source_data") or {}
        row = [str(item.get("id") or source.get("query_id") or ""), identity["query"], identity["context"],
               json.dumps(identity["query_hashes"], ensure_ascii=False), item.get("category") or source.get("category") or "",
               item.get("session_group") or item.get("session_id") or source.get("session_group") or source.get("session_id", ""),
               item.get("turn_index",source.get("turn_index","")), item.get("history_prefix_sha256", "")]
        for n in range(1, count+1):
            response = response_identity(item, n)
            row.extend(response.get(k, "") for k in ("answer", "context", "response_id", "sha256", "evidence"))
            row.extend([None] * (len(dimensions)*(6 if detailed_states else 4)
                                 + (0 if knowledge else 2) + (1 if chat else 0)))
        if not chat:
            row.extend([None]*len(dimensions))
        table.rows.append(row)
    meta = dict(template_version="human-wide-1.0", products=products, human_standard_version=protocol.standard_version,
                human_standard_id=protocol.standard_id, human_bundle_revision=protocol.bundle_revision,
                human_result_schema_version=protocol.result_schema_version,
                evaluation_profile=protocol.id, dimensions=list(dimensions), dimension_names=names_by_dimension,
                score_range=[protocol.score_min, protocol.score_max])
    if knowledge:
        from ..judges.knowledge_qa_schema import KNOWLEDGE_QA_ALLOWED_SCORES
        meta["allowed_scores_by_dimension"] = {key: list(values) for key, values in KNOWLEDGE_QA_ALLOWED_SCORES.items()}
    notes = StatisticsTable("填写说明", ("项目", "说明"), rows=[
        ["模板元信息", json.dumps(meta, ensure_ascii=False)], ["人工标准", protocol.display],
        ["分值范围", f"{protocol.score_min}—{protocol.score_max} 的整数"],
        ["填写规则", f"题号须为文本；留空=未标注；review非空优先（包括N/A）；数字只接受{protocol.score_min}—{protocol.score_max}严格整数，N/A默认原因未知。" if protocol.result_schema_version in ("qa-chat-result-2", "timeliness-result-1", "timeliness-result-2", "timeliness-result-3") else "题号须为文本；留空=未标注；review 非空优先（包括 0 和 N/A）。N/A 默认原因未知。"],
        ["状态可选值", "scored / not_applicable / insufficient_evidence / blocked_gate / blocked_understanding / input_failed / unlabeled / na_unspecified" if timeliness else "scored / not_applicable / unverifiable / gate_blocked / input_failed / unlabeled / na_unspecified"],
        ["身份核对", "请核对产品、题号和原始回答；缺失真实题号须先补充。填写说明中的产品 ID 可在导入映射中修正。"],
        ["Gate/适用性", "Gate：pass/fail/unclear；适用性：true/false；留空=未知。"],
        ["标准隔离", "禁止将旧标准人工分或模型分复制为新标准标注；全部评分、理由、Gate与适用性栏初始为空。"],
        ["核验/输入状态", "核验：verified/partial/unverifiable/not_required；输入：complete/partial/failed；留空=未知。"],
    ])
    if knowledge:
        replacements = {
            "分值范围": "准确性：0、1、1.5、2；其余三维：0、1、2。必须填写Excel数字，不能用文本数字。",
            "填写规则": "题号须为文本；留空=未标注；复核分非空优先，保留0及准确性1.5；空分原因在评分状态中明确，不补0。",
            "状态可选值": "scored / not_applicable / insufficient_evidence / unverifiable / input_failed / unlabeled / na_unspecified；仅富媒体允许not_applicable。",
            "Gate/适用性": "知识问答无Gate；前三维始终检查，只有富媒体可不适用。理由应包含证据位置或短原话、实际影响及分档依据；缺证不能补造。",
            "标准隔离": "禁止复制旧标准评分；本模板仅复用原始测评输入，四维分数、状态、理由均需重新人工标注。",
            "核验/输入状态": "模型不额外输出核验结论；人工输入状态可填complete/partial/failed，留空代表未知。",
        }
        notes.rows = [[key, replacements.get(key, value)] for key, value in notes.rows]
    return tables_workbook({"人工标注": [table], "填写说明": [notes]}, plain=True)

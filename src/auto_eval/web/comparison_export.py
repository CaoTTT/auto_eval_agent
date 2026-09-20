"""Pair frozen task results from the same ordered input dataset."""
from __future__ import annotations

import hashlib
import json
from html import unescape
from itertools import chain
from pathlib import Path

from .export_links import ExportLink
from .history import (
    _aligned_results, _results_with_identity, _source_data_for_item,
    _headers, _xlsx_text, export_rows,
)


def validate_pair(left: dict, right: dict) -> str:
    if left.get("task_id") == right.get("task_id"):
        raise ValueError("请选择两个不同的任务")
    if left.get("mode") not in {"compare", "rich_content"} or left.get("mode") != right.get("mode"):
        raise ValueError("两个任务的评测模式必须相同")
    items = left.get("items") or []
    others = right.get("items") or []
    if not items or len(items) != len(others):
        raise ValueError("测试集为空或 case 数量不同，无法合并")
    digest = hashlib.sha256()
    for index, (a, b) in enumerate(zip(items, others), 1):
        def identity(item):
            return json.dumps({"source": _source_data_for_item(item),
                               "session_id": item.get("session_id"),
                               "turn_index": item.get("turn_index")},
                              ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        encoded = identity(a)
        if encoded != identity(b):
            raise ValueError(f"第 {index} 个 case 的输入内容、顺序或会话轮次不同，无法合并")
        digest.update(encoded.encode("utf-8") + b"\n")
    return digest.hexdigest()


def comparison_sheets(left: dict, right: dict) -> dict[str, list[dict]]:
    fingerprint = validate_pair(left, right)
    snapshots = (left, right)
    exports = [export_rows(snapshot) for snapshot in snapshots]
    aligned = [_aligned_results(snapshot, _results_with_identity(snapshot)) for snapshot in snapshots]
    columns = _headers(exports[0]["逐题结果"] + exports[1]["逐题结果"])
    # Input is shared in a separate sheet; result fields are always adjacent A/B.
    identity_columns = {"query_id", "题目", "Query", "数据集序号", "会话ID", "轮次"}
    rows = []
    for index, item in enumerate(left["items"]):
        row = {"数据集序号": index + 1, "query_id": item.get("id", f"q{index}"),
               "题目": item.get("query") or item.get("question") or ""}
        if item.get("session_id"):
            row.update({"会话ID": item["session_id"], "轮次": item.get("turn_index")})
        for field in ("评估状态", "error"):
            for side, label in enumerate(("A", "B")):
                row[f"{field} · {label}"] = aligned[side][index].get(field, "")
        for column in columns:
            if column in identity_columns or column in {"评估状态", "error"}:
                continue
            for side, label in enumerate(("A", "B")):
                row[f"{column} · {label}"] = exports[side]["逐题结果"][index].get(column, "")
        rows.append(row)
    info = []
    for label, snapshot, sheets in zip(("A", "B"), snapshots, exports):
        manifest = snapshot.get("protocol_manifest") or {}
        info.append({"对比侧": label, **sheets["运行信息"][0],
                     "任务ID": snapshot["task_id"], "备注": snapshot.get("note", ""),
                     "评测协议": snapshot.get("evaluation_profile", ""),
                     "评分范围": manifest.get("score_range", "未记录"),
                     "协议信息": manifest, "裁判运行信息": snapshot.get("judge_runtime", {}),
                     "输入内容SHA256": fingerprint,
                     "说明": "按原始输入顺序对齐；A/B 为选择顺序。不同评分范围须分别解读。"
                               "失败和待评估不计为零分；原图请使用单任务完整导出。"})
    return {"逐题对比": rows, "共同测试数据": exports[0]["数据集明细"], "任务说明": info}


def write_comparison_xlsx(left: dict, right: dict, destination: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = comparison_sheets(left, right)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        headers = _headers(rows)
        alignment = Alignment(vertical="top", wrap_text=True)
        values_by_row = chain([headers], ([row.get(key, "") for key in headers] for row in rows))
        for row_number, values in enumerate(values_by_row, 1):
            # All strings are literal, including formulas supplied in dataset cells.
            for column_number, value in enumerate(values, 1):
                hyperlink = str(value) if isinstance(value, ExportLink) else None
                value = value if isinstance(value, (int, float, bool)) else unescape(_xlsx_text(value))
                cell = sheet.cell(row_number, column_number, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = alignment
                if hyperlink:
                    cell.hyperlink = hyperlink
                    cell.style = "Hyperlink"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28 if " · " in header else 22
        sheet.freeze_panes = "D2" if name == "逐题对比" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if name == "逐题对比":
            difference_fill = PatternFill("solid", fgColor="FFF2CC")
            row_count = len(rows) + 1
            for index, header in enumerate(headers, 1):
                if not header.endswith(" · A"):
                    continue
                for row_number in range(2, row_count + 1):
                    a, b = sheet.cell(row_number, index), sheet.cell(row_number, index + 1)
                    if a.value != b.value:
                        a.fill = b.fill = difference_fill
    workbook.save(destination)

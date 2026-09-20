import asyncio
import copy
import io

import httpx
import pytest
from openpyxl import load_workbook

from auto_eval.web import exports, server
from auto_eval.web.comparison_export import comparison_sheets, validate_pair, write_comparison_xlsx
from auto_eval.web.tasks import Task


def snapshots():
    a = {"task_id": "a", "mode": "compare", "dataset_name": "same.jsonl",
         "items": [{"id": "duplicate", "query": "=1+1 <tag>&", "answer1": "x"},
                   {"id": "duplicate", "query": "second", "answer1": "y"}],
         "results": [{"index": 1, "answer1_understanding_score": 2},
                     {"index": 0, "answer1_understanding_score": 5}],
         "protocol_manifest": {"score_range": [1, 5], "bundle_revision": "v1"},
         "judge_runtime": {"judges": [{"model": "model-a"}]}}
    b = copy.deepcopy(a)
    b.update(task_id="b", results=[{"index": 0, "answer1_understanding_score": 3}])
    b["judge_runtime"] = {"judges": [{"model": "model-b"}]}
    b["protocol_manifest"]["bundle_revision"] = "v2"
    return a, b


def test_pairs_are_adjacent_and_results_align_by_input_index(tmp_path):
    a, b = snapshots()
    sheets = comparison_sheets(a, b)
    first, second = sheets["逐题对比"]
    keys = list(first)
    score = "产品1理解需求分"
    assert keys[keys.index(score + " · A") + 1] == score + " · B"
    assert first[score + " · A"] == 5
    assert first[score + " · B"] == 3
    assert second[score + " · A"] == 2
    assert second["评估状态 · B"] == "待评估"
    path = tmp_path / "comparison.xlsx"
    write_comparison_xlsx(a, b, path)
    workbook = load_workbook(path)
    sheet = workbook["逐题对比"]
    assert sheet["C2"].value == "=1+1 <tag>&"
    assert sheet["C2"].data_type == "s"
    assert sheet.freeze_panes == "D2"
    assert sheet.cell(2, keys.index(score + " · A") + 1).fill.fgColor.rgb == "00FFF2CC"
    assert len(workbook["任务说明"]["A"]) == 3


@pytest.mark.parametrize("change", ["query", "order", "count", "mode", "same_task", "turn"])
def test_rejects_mismatched_datasets(change):
    a, b = snapshots()
    if change == "query":
        b["items"][0]["query"] = "changed"
    elif change == "order":
        b["items"].reverse()
    elif change == "count":
        b["items"].pop()
    elif change == "mode":
        b["mode"] = "rich_content"
    elif change == "same_task":
        b["task_id"] = "a"
    else:
        b["items"][0]["turn_index"] = 2
    with pytest.raises(ValueError):
        validate_pair(a, b)


def test_raw_source_ignores_runtime_changes_and_preserves_failure():
    a, b = snapshots()
    for snapshot in (a, b):
        for item in snapshot["items"]:
            item["source_data"] = dict(item)
    b["items"][0]["frames1"] = ["different cache path"]
    b["results"] = [{"index": 0, "error": "judge unavailable"}]
    rows = comparison_sheets(a, b)["逐题对比"]
    assert rows[0]["评估状态 · B"] == "评估失败"
    assert rows[0]["error · B"] == "judge unavailable"


def test_rich_content_supported():
    a, b = snapshots()
    a["mode"] = b["mode"] = "rich_content"
    a["results"] = [{"index": 0, "correctness": "correct"}]
    assert "correctness · A" in comparison_sheets(a, b)["逐题对比"][0]


@pytest.mark.asyncio
async def test_background_comparison_api_and_validation(tmp_path, monkeypatch):
    manager = exports.XlsxExports(tmp_path)
    tasks = {}
    for snapshot in snapshots():
        tasks[snapshot["task_id"]] = Task(id=snapshot["task_id"], mode=snapshot["mode"],
            items=snapshot["items"], options={}, results=snapshot["results"])
    async def peek(key):
        return tasks.get(key)
    monkeypatch.setattr(exports, "peek_task_async", peek)
    monkeypatch.setattr(server, "XLSX_EXPORTS", manager)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="http://test") as client:
        for ids in (["a"], ["a", "a"], ["a", "b", "c"], ["a", " "]):
            assert (await client.post("/api/exports/comparison", json={"task_ids": ids})).status_code == 422
        response = await client.post("/api/exports/comparison", json={"task_ids": ["a", "b"]})
        assert response.status_code == 202
        key = response.json()["export_id"]
        assert manager.create("a", "b")["export_id"] == key
        assert manager.create("a")["export_id"] != key
        await asyncio.gather(*list(manager.workers))
        assert manager.view(key)["status"] == "ready"
        response = await client.get(f"/api/exports/{key}/download")
        assert load_workbook(io.BytesIO(response.content)).sheetnames[0] == "逐题对比"
        tasks["b"].items = [{"query": "wrong"}]
        bad = manager.create("a", "b")["export_id"]
        missing = manager.create("a", "missing")["export_id"]
        await asyncio.gather(*list(manager.workers))
        assert manager.view(bad)["status"] == "error"
        assert manager.view(missing)["error"] == "对比任务不存在或已删除"
    await manager.close()

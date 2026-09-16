from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest

from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baselines import HumanError, HumanStore, ImportMapping, atomic_json
from test_compare_statistics import make_snapshot


def workbook(tmp_path, rows, *, standard="0.2-simplified", extra=None):
    path=tmp_path/"human.xlsx"
    wb=Workbook()
    ws=wb.active
    ws.title="人工标注"
    ws.append(["case_id","query","人工分_产品1_理解需求","人工复核分_产品1_理解需求","人工分_产品2_理解需求",
               "人工评分状态_产品1_理解需求","人工是否适用_理解需求","人工响应Gate_产品1"])
    for row in rows:ws.append(row)
    hidden=wb.create_sheet("WPS图片管理")
    hidden.sheet_state="hidden"
    wb.save(path)
    m=inspect_workbook(path)["suggested_mapping"]
    m.update(human_standard_version=standard,policy_note="专家按指定标准评分")
    if extra:m.update(extra)
    return path,ImportMapping(**m)


def test_review_na_zero_and_conflicts(tmp_path):
    path,m=workbook(tmp_path,[["001","q",3,"N/A",2],["002","q",2,0,1],["003","q",2,None,3],
                             ["004","q",3,True,2],["005","q",1,1.5,2],["006","q",2,"4分（有问题）",2]],standard="0.3")
    result=parse_human_labels(path,m)
    first=[r for r in result["labels"] if r["product_id"]=="product1"]
    assert [(r["score_status"],r["score"]) for r in first]==[("na_unspecified",None),("scored",0),("scored",2)]
    assert result["cases"][0]["case_id"]=="001"
    assert first[0]["source"]["selected_cell"]=="D2"
    assert len(result["issues"])==3
    assert inspect_workbook(path)["sheets"]==["人工标注"]


def test_duplicate_numeric_id_and_state_conflicts_are_excluded(tmp_path):
    path,m=workbook(tmp_path,[["001","q",4,None,3],["001","q",5,None,3],[12,"q",3,None,3],
                             ["002","q",4,None,3,"not_applicable"],["003","q",4,None,3,None,False],
                             ["004","q",4,None,3,None,None,"fail"]])
    result=parse_human_labels(path,m)
    assert all(not r["case_key"].endswith(":001") for r in result["labels"])
    assert not any(r["product_id"]=="product1" for r in result["labels"])
    assert len(result["issues"])>=5
    store=HumanStore(tmp_path/"runs")
    dest=store.path("imports","hi_test","source.xlsx")
    dest.parent.mkdir(parents=True)
    dest.write_bytes(path.read_bytes())
    atomic_json(store.path("imports","hi_test","preview.json"),result)
    with pytest.raises(HumanError,match="解析冲突"):
        store.publish_baseline("hi_test",result["preview_sha256"],"test")
    manifest=store.publish_baseline("hi_test",result["preview_sha256"],"test",valid_only=True)
    assert manifest["excluded_records"]
    old=store.load_baseline(manifest["baseline_id"],1)
    second=store.publish_baseline("hi_test",result["preview_sha256"],"test",manifest["baseline_id"],1,valid_only=True)
    assert second["version"]==2
    assert old==store.load_baseline(manifest["baseline_id"],1)
    with pytest.raises(HumanError,match="版本已改变"):
        store.publish_baseline("hi_test",result["preview_sha256"],"test",manifest["baseline_id"],1,valid_only=True)


def test_header_change_formula_cache_and_duplicate_headers(tmp_path):
    path,m=workbook(tmp_path,[["001","q","=1+2",None,4]])
    result=parse_human_labels(path,m)
    assert "公式无缓存" in result["issues"][0]["message"]
    wb=load_workbook(path)
    wb.active["C1"]="变更表头"
    wb.save(path)
    with pytest.raises(HumanError,match="表头已改变"):
        parse_human_labels(path,m)


def test_review_overrides_uncached_formula_and_shared_na_conflict(tmp_path):
    path,m=workbook(tmp_path,[["001","q","=1+2",4,3],["002","q",2,"N/A",3]],extra={"na_status":"not_applicable"})
    result=parse_human_labels(path,m)
    first=[r for r in result["labels"] if r["case_key"].endswith(":001") and r["product_id"]=="product1"][0]
    assert first["score"]==4 and first["source"]["formula"]
    conflicts=[r for r in result["labels"] if r["case_key"].endswith(":002")]
    assert all(r["resolution"]=="pending_conflict" for r in conflicts)


def test_legacy_152_row_adapter_and_changed_physical_header(tmp_path):
    wb=Workbook();ws=wb.active;ws.title="合并标注结果-review后"
    ws['A1']='query_id';ws['B1']='题目'
    specs=[('AW','AX','',''),('BJ','BK','',''),('BU','BV','',''),('CF','CG','',''),('CU','CV','CZ','DA'),('DK','DM','DR','DS'),('ED','EE','','')]
    for a,b,ar,br in specs:
        for c in (a,b):ws[c+'1']='人工最终分'
        for c in filter(None,(ar,br)):ws[c+'1']='review'
    ws['I1']='胜负结果';ws['J1']='胜负结果'
    for n in range(2,154):
        ws[f'A{n}']=f'{n-1:03}';ws[f'B{n}']=f'问题{n}'
        for a,b,ar,br in specs:
            ws[f'{a}{n}']=3;ws[f'{b}{n}']=4
        if n<=18:ws[f'CZ{n}']=4;ws[f'DA{n}']=5
        if n<=29:ws[f'DR{n}']='N/A' if n==2 else 4;ws[f'DS{n}']='N/A' if n==2 else 5
    path=tmp_path/'legacy.xlsx';wb.save(path)
    mapping=inspect_workbook(path)['suggested_mapping']
    assert len(mapping['labels'])==14
    mapping.update(human_standard_version='0.2-simplified',policy_note='合成人工基准')
    result=parse_human_labels(path,ImportMapping(**mapping))
    assert result['summary']['case_count']==152 and not result['issues']
    assert result['summary']['statuses']['na_unspecified']==2
    assert sum(r['annotation_stage']=='review' for r in result['labels'])==(17+28)*2
    ws['CU1']='模型评分';wb.save(path)
    assert inspect_workbook(path)['suggested_mapping']['labels']==[]


@pytest.mark.parametrize("standard",["0.2-simplified","0.2-simplified-calibrated","0.2-simplified-thinking-exposure","0.3"])
@pytest.mark.parametrize("count",[2,3])
def test_blank_template_roundtrip_without_model_answers(tmp_path,standard,count):
    snapshot=make_snapshot([[1]*count],standard)
    snapshot["items"][0].update(answer1="=SUM(A1:A2)",context1="source")
    data=build_human_template(snapshot,standard)
    wb=load_workbook(BytesIO(data))
    ws=wb["人工标注"]
    headers=[c.value for c in ws[1]]
    assert ws.cell(2,headers.index("源回答_产品1")+1).data_type=="s"
    assert not any("内容准确性" in h for h in headers)
    assert all(ws.cell(2,i+1).value is None for i,h in enumerate(headers) if h.startswith("人工分_"))
    ws.cell(2,headers.index("人工分_产品1_理解需求")+1,1)
    path=tmp_path/"filled.xlsx"
    wb.save(path)
    info=inspect_workbook(path)
    mapping=info["suggested_mapping"]
    mapping["policy_note"]="人工核对标准"
    assert mapping["human_standard_version"]==standard
    assert len(mapping["products"])==count
    result=parse_human_labels(path,ImportMapping(**mapping))
    assert result["issues"]==[]
    assert result["summary"]["statuses"]["scored"]==1

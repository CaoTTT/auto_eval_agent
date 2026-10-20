"""Human annotations use the same frozen three-dimensional timeliness family."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_baselines import HumanError, HumanStore, ImportMapping, atomic_json, read_json, baseline_protocol
from auto_eval.web.human_compare import ComparisonRequest, GenerateRequest, HumanComparisons, align_baseline_to_run
from auto_eval.web.human_report_export import export_report
from auto_eval.web.human_statistics import build_human_statistics
from test_human_compare import example, put_baseline
from test_timeliness_statistics import dataset


def timeliness_example(count=2, revision=None):
    task, snapshot, baseline, mapping = example([[3] * count, [2] * count])
    generated, rows, protocol = dataset([[3] * count, [2] * count], revision=revision)
    snapshot.update(evaluation_profile=protocol.id, protocol_manifest=protocol.public_metadata())
    baseline.update(human_standard_id=protocol.standard_id, human_standard_version=protocol.standard_version,
                    human_bundle_revision=protocol.bundle_revision, human_result_schema_version=protocol.result_schema_version,
                    dimensions=list(protocol.dimensions), score_dimensions=list(protocol.score_dimensions),
                    dimension_names=protocol.dimension_names, score_range=[1, 3])
    for target, generated in zip(snapshot["results"], rows):
        target.update(generated)
    baseline["labels"] = [{**label, "dimension_id": dim, "applicable": True, "verification_status": "verified"}
                          for label in baseline["labels"] for dim in protocol.dimensions]
    return task, snapshot, baseline, mapping


def fill_template(tmp_path, fields, count=2):
    snapshot, _, _ = dataset([([3] * count)])
    path = tmp_path / "timeliness.xlsx"
    path.write_bytes(build_human_template(snapshot, "timeliness_compare@1.0"))
    wb = load_workbook(path)
    ws = wb["人工标注"]
    headers = {c.value: c.column for c in ws[1]}
    assert all(ws.cell(2, col).value is None for name, col in headers.items() if name.startswith("人工"))
    for key, value in fields.items():
        ws.cell(2, headers[key], value)
    wb.save(path)
    wb.close()
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "依专家修订版独立人工评定"
    return path, ImportMapping(**mapping)


@pytest.mark.parametrize("count", (2, 3))
def test_template_import_publish_export_roundtrip_new_statuses(tmp_path, count):
    path, mapping = fill_template(tmp_path, {
        "人工分_产品1_理解需求": 1, "人工是否适用_产品1_理解需求": True,
        "人工核验状态_产品1_理解需求": "verified",
        "人工评分状态_产品1_权威性": "blocked_understanding",
        "人工评分状态_产品1_时效性": "blocked_understanding",
        "人工响应Gate_产品1": "unclear", "人工安全Gate_产品1": "pass",
        "人工分_产品2_理解需求": 3,
        "人工评分状态_产品2_权威性": "not_applicable",
        "人工是否适用_产品2_权威性": False, "人工核验状态_产品2_权威性": "not_required",
        "人工评分状态_产品2_时效性": "insufficient_evidence",
        "人工是否适用_产品2_时效性": True, "人工核验状态_产品2_时效性": "unverifiable",
    }, count)
    assert mapping.human_standard_id == "timeliness_compare"
    assert mapping.human_result_schema_version == "timeliness-result-3"
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    store = HumanStore(tmp_path / "store")
    dest = store.path("imports", "hi_time", "source.xlsx")
    dest.parent.mkdir(parents=True)
    dest.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_time", "preview.json"), parsed)
    manifest = store.publish_baseline("hi_time", parsed["preview_sha256"], "时效性人工基准")
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    path.write_bytes(export_baseline_answers(baseline))
    remapping = inspect_workbook(path)["suggested_mapping"]
    remapping["policy_note"] = "同一冻结标准重导入"
    second = parse_human_labels(path, ImportMapping(**remapping))
    assert not second["issues"]
    assert {r["score_status"] for r in second["labels"]} >= {"scored", "blocked_understanding", "insufficient_evidence", "not_applicable"}


@pytest.mark.parametrize("fields", [
    {"人工分_产品1_理解需求": 4},
    {"人工分_产品1_理解需求": "3"},
    {"人工评分状态_产品1_理解需求": "not_applicable"},
    {"人工分_产品1_权威性": 2, "人工核验状态_产品1_权威性": "partial"},
    {"人工响应Gate_产品1": "fail", "人工评分状态_产品1_时效性": "not_applicable"},
    {"人工分_产品1_理解需求": 1, "人工分_产品1_时效性": 3},
])
def test_import_rejects_out_of_scale_and_contradictory_states(tmp_path, fields):
    path, mapping = fill_template(tmp_path, fields)
    assert parse_human_labels(path, mapping)["issues"]


def test_import_allows_decisive_one_with_partial_and_unknown_applicability(tmp_path):
    path, mapping = fill_template(tmp_path, {
        "人工分_产品1_权威性": 1, "人工核验状态_产品1_权威性": "partial",
        "人工评分状态_产品1_时效性": "insufficient_evidence",
        "人工核验状态_产品1_时效性": "unverifiable",
    })
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    time_label = next(r for r in parsed["labels"] if r["product_id"] == "product1" and r["dimension_id"] == "timeliness")
    assert time_label["applicable"] is None


def test_frozen_baseline_missing_schema_is_not_guessed():
    _, _, baseline, _ = timeliness_example()
    del baseline["human_result_schema_version"]
    with pytest.raises(HumanError, match="Schema"):
        baseline_protocol(baseline)


def test_old_standard_template_cannot_be_relabelled_as_timeliness(tmp_path):
    snapshot, _, _ = dataset()
    path = tmp_path / "old.xlsx"
    path.write_bytes(build_human_template(snapshot, "qa_chat_compare@1.1"))
    workbook = load_workbook(path)
    workbook.save(path)
    workbook.close()
    mapping = inspect_workbook(path)["suggested_mapping"]
    protocol = resolve_compare_protocol("timeliness_compare@1.0")
    mapping.update(human_standard_id=protocol.standard_id, human_standard_version=protocol.standard_version,
                   human_bundle_revision=protocol.bundle_revision, human_result_schema_version=protocol.result_schema_version,
                   policy_note="不能把已有闲聊标签重命名", labels=[c for c in mapping["labels"] if c["dimension_id"] == "understanding"])
    with pytest.raises(HumanError, match="不能仅改版本"):
        parse_human_labels(path, ImportMapping(**mapping))


@pytest.mark.parametrize("revision", (None, "1.0.0", "1.0.1"))
def test_human_alignment_statistics_report_share_reasons_without_winners(tmp_path, revision):
    _, snapshot, baseline, mapping = timeliness_example(3, revision=revision)
    if revision != "1.0.0":
        for result in snapshot["results"]:
            for product in (1, 2, 3):
                for dimension in ("understanding", "authority", "timeliness"):
                    # New model schema never collects this; missing and null
                    # must both preserve numeric consistency with humans.
                    result.pop(f"answer{product}_{dimension}_verification_status")
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["authority"], "same_standard")
    assert all(r["comparable"] for r in rows)
    assert rows[0]["model_reason"] == snapshot["results"][0]["authority_reason"]
    assert rows[0]["model_reason_scope"] == "dimension_shared"
    if revision is None:
        assert rows[0]["model_conflict_reason"] == snapshot["results"][0]["conflict_reason"]
        assert rows[0]["model_has_conflict"] == snapshot["results"][0]["has_conflict"]
        assert rows[0]["model_conflict_status"] in ("有冲突", "无冲突", "不确定")
    else:
        assert "model_has_conflict" not in rows[0]
    metrics = build_human_statistics(rows, ["test"], ["p1", "p2", "p3"], ["authority"])
    assert metrics["pairs"] == metrics["ranks"] == []
    assert all(s["n"] == 2 for s in metrics["scores"])
    service = HumanComparisons(HumanStore(tmp_path))
    put_baseline(service.store, baseline)
    preview = service.create_preview(ComparisonRequest(baseline_id="hb_test", version=1, tasks=[mapping]), [snapshot])
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"], config_sha256=preview["config_sha256"]))
    service.generate(manifest["report_id"])
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    assert report["metrics"]["pairs"] == report["metrics"]["ranks"] == []
    wb = load_workbook(BytesIO(export_report(report)), read_only=True)
    try:
        assert "产品差距对比" not in wb.sheetnames
        headers = next(wb["匹配与排除"].values)
        assert any("模型理由（维度共享完整原文）" in str(v) for v in headers)
        if revision is None:
            conflict_reason = next(i for i, value in enumerate(headers) if str(value).endswith("模型内容冲突理由（同题跨产品）"))
            conflict_verdict = next(i for i, value in enumerate(headers) if str(value).endswith("模型内容冲突（不计入数字一致性）"))
            conflict_status = next(i for i, value in enumerate(headers) if str(value).endswith("模型内容冲突状态（程序判定）"))
            assert conflict_reason < conflict_status < conflict_verdict
        else:
            assert not any("模型内容冲突" in str(v) for v in headers)
        if revision != "1.0.0":
            assert not any("模型核验状态" in str(v) or "模型适用性" in str(v) for v in headers)
            assert any("模型NA原因" in str(v) for v in headers)
        else:
            assert any("模型核验状态" in str(v) for v in headers)
    finally:
        wb.close()


def test_new_human_baseline_cannot_align_old_model_frozen_revision():
    _, snapshot, baseline, mapping = timeliness_example()
    old, results, protocol = dataset(((3, 3), (2, 2)), revision="1.0.0")
    snapshot.update(protocol_manifest=protocol.public_metadata())
    for target, generated in zip(snapshot["results"], results):
        target.update(generated)
    with pytest.raises(HumanError, match="冻结修订号"):
        align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["authority"], "same_standard")


def test_timeliness_template_cannot_relabel_frozen_old_revision(tmp_path):
    old, _, _ = dataset(revision="1.0.0")
    path = tmp_path / "frozen-old-time.xlsx"
    path.write_bytes(build_human_template(old, "timeliness_compare@1.0"))
    mapping = inspect_workbook(path)["suggested_mapping"]
    assert mapping["human_bundle_revision"] == "1.0.0"
    assert mapping["human_result_schema_version"] == "timeliness-result-1"
    protocol = resolve_compare_protocol("timeliness_compare@1.0")
    mapping.update(human_bundle_revision=protocol.bundle_revision,
                   human_result_schema_version=protocol.result_schema_version,
                   policy_note="不能只改冻结修订")
    with pytest.raises(HumanError, match="不能仅改版本"):
        parse_human_labels(path, ImportMapping(**mapping))


def test_new_downloaded_template_can_be_imported_without_excel_resave(tmp_path):
    snapshot, _, _ = dataset()
    path = tmp_path / "fresh-download.xlsx"
    path.write_bytes(build_human_template(snapshot, "timeliness_compare@1.0"))
    mapping = inspect_workbook(path)["suggested_mapping"]
    assert mapping["human_result_schema_version"] == "timeliness-result-3"
    assert len(mapping["labels"]) == 6
    mapping["policy_note"] = "直接下载的空白模板"
    parsed = parse_human_labels(path, ImportMapping(**mapping))
    assert not parsed["issues"]
    assert len(parsed["labels"]) == 6
    assert all(label["score_status"] == "unlabeled" for label in parsed["labels"])


def test_old_timeliness_human_baseline_still_aligns_old_frozen_model():
    _, snapshot, baseline, mapping = timeliness_example(revision="1.0.0")
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["authority"], "same_standard")
    assert all(r["comparable"] for r in rows)
    assert all(r["model_result_schema_version"] == "timeliness-result-1" for r in rows)
    assert all(r["model_verification_status"] == "verified" for r in rows)
    metrics = build_human_statistics(rows, ["test"], ["p1", "p2"], ["authority"])
    assert all(metric["n"] == 2 for metric in metrics["scores"])


@pytest.mark.parametrize("count,verdict", ((2, "no"), (3, "yes")))
def test_human_context_marks_failed_collection_without_rewriting_raw_conflict(count, verdict):
    _, snapshot, baseline, mapping = timeliness_example(count)
    snapshot["results"][0].update(has_conflict=verdict, conflict_reason="保留原始跨产品对比判断。",
                                 **{f"answer{count}_input_status": "failed"})
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["authority"], "same_standard")
    failed_case = [row for row in rows if row["case_id"] == "case-0"]
    assert failed_case
    assert all(row["model_conflict_status"] == "采集失败，待核验" for row in failed_case)
    assert all(row["model_has_conflict"] == verdict for row in failed_case)
    assert all(row["model_conflict_reason"] == "保留原始跨产品对比判断。" for row in failed_case)

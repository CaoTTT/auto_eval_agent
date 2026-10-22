"""Protocol selection, frozen history, reporting and diagnostic integration."""
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PIL import Image

from auto_eval.config import AppConfig, JudgeConfig, VisualExtractionConfig, VisualModeProfile
from auto_eval.judges.base import JudgeOutputParseError
from auto_eval.judges.compare_1022_contract import example_1022_output, validate_1022_output
from auto_eval.judges.compare_protocols import (
    DEFAULT_COMPARE_PROTOCOL_ID, V03_COMPARE_PROTOCOL_ID, V1022_COMPARE_PROTOCOL_ID,
    resolve_compare_protocol,
)
from auto_eval.judges.visual_compare_judge import visual_compare_result_fields
from auto_eval.web import history, runner, server
from auto_eval.web.tasks import Task, _task_from_snapshot


def result_for_report(count=2, index=0):
    data = validate_1022_output(example_1022_output(count), count)
    data["has_conflict"] = data.pop("content_conflict")
    data["product_count"] = count
    protocol = resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID)
    observation = protocol.observation_model.model_validate(data)
    result = visual_compare_result_fields(observation, protocol)
    result.update(index=index, item_id=f"case{index}", query="你好")
    return result


def test_config_exposes_new_label_without_changing_default(monkeypatch):
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[JudgeConfig(name="offline")]))
    profiles = server.api_config()["evaluation_profiles"]
    selected = next(item for item in profiles if item["id"] == V1022_COMPARE_PROTOCOL_ID)
    assert selected["display"] == "V0.2简化版（1022优化）"
    assert selected["standard_version"] == "0.2-simplified-1022"
    assert selected["bundle_revision"] == "0.2.1"
    assert selected["output_contract"] == "compare_1022_v1"
    assert selected["score_range"] == [1, 5]
    assert selected["input_modalities"] == ["text", "text_image"]
    assert next(item["id"] for item in profiles if item["status"] == "stable") == DEFAULT_COMPARE_PROTOCOL_ID
    assert resolve_compare_protocol(None).id == DEFAULT_COMPARE_PROTOCOL_ID
    assert resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID).request_metadata()["output_contract"] == "compare_1022_v1"


@pytest.mark.parametrize("revision", ["0.2.0", "0.3.0", "0.3.1", "9.9"])
def test_new_protocol_never_borrows_legacy_revision(revision):
    with pytest.raises(ValueError, match="冻结"):
        resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID, revision)


def test_existing_legacy_protocol_revisions_remain_resolvable():
    assert resolve_compare_protocol(DEFAULT_COMPARE_PROTOCOL_ID, "0.2.0").bundle_revision == "0.2.0"
    assert resolve_compare_protocol(V03_COMPARE_PROTOCOL_ID, "0.3.0").bundle_revision == "0.3.0"
    assert resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID, "0.2.1").display == "V0.2简化版（1022优化）"


@pytest.mark.asyncio
async def test_eval_api_freezes_new_protocol_and_snapshot_restores_it(monkeypatch, tmp_path):
    cfg = AppConfig(
        judges=[JudgeConfig(name="offline", compare_structured_output="json_object")],
        visual_modes={"rich_content": VisualModeProfile(
            extraction=VisualExtractionConfig(algorithm_version="test"),
        )},
    )
    monkeypatch.setattr(server, "cfg", lambda: cfg)
    created = []

    def new_task(mode, items, options, dataset_name="", **kwargs):
        task = Task(id="1022-task", mode=mode, items=items, options=options,
                    dataset_name=dataset_name, evaluation_profile=kwargs["evaluation_profile"],
                    protocol_manifest=kwargs["protocol_manifest"],
                    judge_runtime=kwargs["judge_runtime"])
        created.append(task)
        return task

    monkeypatch.setattr(server, "new_task", new_task)
    monkeypatch.setattr(server.EVAL_SCHEDULER, "enqueue", lambda *_: 1)
    response = await server.api_eval(server.EvalReq(
        mode="compare", evaluation_profile=V1022_COMPARE_PROTOCOL_ID,
        items=[{"id": "greeting", "query": "你好", "video1": "a.mp4", "video2": "b.mp4"}],
    ))
    assert response["evaluation_profile"] == V1022_COMPARE_PROTOCOL_ID
    task = created[0]
    task.status = "done"
    task.results = [result_for_report()]
    task.summary = runner._summarize(task)
    monkeypatch.setattr(history, "HISTORY_DIR", tmp_path / "history")
    assert history.save_task(task)
    snapshot = history.load_snapshot(task.id)
    restored = _task_from_snapshot(snapshot, task.id)
    assert restored.evaluation_profile == V1022_COMPARE_PROTOCOL_ID
    assert restored.protocol_manifest == task.protocol_manifest
    assert restored.protocol_manifest["output_contract"] == "compare_1022_v1"
    assert restored.results == task.results
    assert restored.summary["standard_version"] == "0.2-simplified-1022"
    assert restored.protocol_manifest["judges"][0]["compare_structured_output"] == "json_object"
    assert restored.judge_runtime["judges"][0]["compare_structured_output"] == "json_object"
    cfg.judges[0].compare_structured_output = "off"
    captured = []

    def make_client(configuration):
        captured.append(configuration)
        return object()

    monkeypatch.setattr(runner, "JudgeClient", make_client)
    runner._make_item_evaluator(restored, cfg)
    assert captured[0].compare_structured_output == "json_object"


@pytest.mark.parametrize("count", [2, 3])
def test_summary_and_excel_preserve_new_version_and_five_point_scores(count):
    protocol = resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID)
    result = result_for_report(count)
    task = Task(id="report", mode="compare", items=[{"id": "case0", "query": "你好"}],
                options={}, results=[result], evaluation_profile=protocol.id,
                protocol_manifest=protocol.public_metadata(), status="done")
    summary = runner._summarize(task)
    assert summary["standard_version"] == "0.2-simplified-1022"
    assert summary["understanding_answer1_avg"] == 5
    assert summary["understanding_scored"] == 1
    assert summary["accuracy_aggregation_enabled"] is False
    if count == 3:
        assert summary["understanding_answer3_avg"] == 5
    row = history._visual_compare_export_rows([result])[0]
    assert row["标准版本"] == "0.2-simplified-1022"
    assert row["评测协议"] == V1022_COMPARE_PROTOCOL_ID
    assert row["产品1理解需求分"] == 5
    assert row["理解需求理由"] == result["understanding_reason"]
    workbook = load_workbook(BytesIO(history.build_xlsx(history.task_to_snapshot(task))), read_only=True)
    try:
        rows = [row for sheet in workbook for row in sheet.iter_rows(values_only=True)]
        assert any("标准版本" in row for row in rows)
        assert any("0.2-simplified-1022" in row for row in rows)
        assert any(V1022_COMPARE_PROTOCOL_ID in row for row in rows)
        assert any("产品3理解需求分" in row for row in rows) == (count == 3)
    finally:
        workbook.close()


@pytest.mark.parametrize("category", ["json_syntax", "schema", "state", "truncated"])
def test_failure_log_keeps_raw_and_separates_output_error_categories(tmp_path, monkeypatch, category):
    monkeypatch.setattr(runner, "RUNS_DIR", tmp_path)
    raw = '{"answer1_understanding_score":'
    error = JudgeOutputParseError(
        "离线失败样例", raw_output=raw, repair_output="", judge="offline", model="fake",
        error_category=category, finish_reason="length" if category == "truncated" else None,
    )
    runner._write_eval_error("1022", 0, {"id": "a", "query": "q"}, error, request_id="req")
    record = json.loads((tmp_path / "eval_errors.jsonl").read_text(encoding="utf-8"))
    assert record["stage"] == f"judge_output_{category}"
    assert record["error_category"] == category
    assert record["original_model_output"] == raw
    assert record["original_output_length"] == len(raw)
    assert record["repair_model_output"] == ""
    assert record["finish_reason"] == ("length" if category == "truncated" else None)


@pytest.mark.asyncio
@pytest.mark.parametrize("category", ["json_syntax", "schema", "state"])
async def test_runner_returns_diagnosable_failure_without_retries(tmp_path, monkeypatch, category):
    """Keep a rejected judgment out of aggregates, but retain it in the error log."""
    output = example_1022_output()
    if category == "schema":
        output["standard_id"] = "model-invented-metadata"
    elif category == "state":
        output["answer1_understanding_score"] = None
    raw = json.dumps(output, ensure_ascii=False)
    if category == "json_syntax":
        raw = raw[:-1]
    calls = []

    class Client:
        persona = "离线裁判"
        model = "fake"

        def __init__(self, configuration):
            self.cfg = configuration

        async def complete(self, *args, **kwargs):
            calls.append(kwargs)
            return raw

        async def repair_json(self, *args, **kwargs):
            pytest.fail("The runner must not repair a rejected 1022 judgment")

    async def finished(index, result, started):
        task.results.append(result)

    monkeypatch.setattr(runner, "JudgeClient", Client)
    monkeypatch.setattr(runner, "RUNS_DIR", tmp_path)
    monkeypatch.setattr(runner, "log_event", lambda *args, **kwargs: None)
    frame = tmp_path / "frame.png"
    Image.new("RGB", (32, 32), "white").save(frame)
    cfg = AppConfig(
        judges=[JudgeConfig(name="offline")],
        visual_modes={"rich_content": VisualModeProfile(
            extraction=VisualExtractionConfig(algorithm_version="test"),
        )},
    )
    item = {"id": "bad-json", "query": "你好", "product_count": 2,
            "frames1": [str(frame)], "frames2": [str(frame)]}
    protocol = resolve_compare_protocol(V1022_COMPARE_PROTOCOL_ID)
    task = Task(id="bad-output", mode="compare", items=[item], options={},
                evaluation_profile=protocol.id,
                protocol_manifest=server._protocol_manifest(protocol, cfg, {}))
    evaluate_one, _ = runner._make_item_evaluator(task, cfg, on_result=finished)
    result = await evaluate_one(0, item)
    assert result["error_type"] == f"judge_output_{category}"
    assert result["error_category"] == category
    assert "answer1_understanding_score" not in result
    assert len(calls) == 1 and calls[0]["require_complete"] is True
    assert runner._summarize(task)["failed"] == 1
    record = json.loads((tmp_path / "eval_errors.jsonl").read_text(encoding="utf-8"))
    assert record["original_model_output"] == raw
    assert record["repair_model_output"] == ""
    assert record["error_category"] == category

"""Knowledge QA human references keep 0/1.5 and merged reasons end to end."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.web.human_baselines import (
    HumanError, HumanStore, ImportMapping, atomic_json, baseline_protocol, read_json,
)
from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_compare import (
    ComparisonRequest, GenerateRequest, HumanComparisons, align_baseline_to_run,
)
from auto_eval.web.human_report_export import export_report
from auto_eval.web.human_statistics import build_human_statistics
from test_human_compare import example, put_baseline


def knowledge_example(count=2):
    task, snapshot, baseline, mapping = example([[2] * count, [1] * count])
    protocol = resolve_compare_protocol("knowledge_qa_compare@1.0")
    snapshot.update(evaluation_profile=protocol.id, protocol_manifest=protocol.public_metadata())
    baseline.update(human_standard_id=protocol.standard_id, human_standard_version=protocol.standard_version,
                    human_bundle_revision=protocol.bundle_revision, human_result_schema_version=protocol.result_schema_version,
                    dimensions=list(protocol.dimensions), score_dimensions=list(protocol.score_dimensions),
                    dimension_names=protocol.dimension_names, score_range=[0, 2])
    for i, result in enumerate(snapshot["results"]):
        result.update(standard_id=protocol.standard_id, standard_version=protocol.standard_version,
                      evaluation_profile=protocol.id, bundle_revision=protocol.bundle_revision,
                      result_schema_version=protocol.result_schema_version)
        for dimension in protocol.dimensions:
            result[dimension + "_reason"] = "；".join(
                f"产品{n}：截图顶部短原话支持当前判断，实际任务完成程度明确，按本维分档。"
                for n in range(1, count + 1))
            for n in range(1, count + 1):
                score = ((0 if n == 1 else 2) if i == 0 else (1 if n == 1 else 1.5)) if dimension == "accuracy" else 2
                prefix = f"answer{n}_{dimension}"
                result[prefix + "_score"] = score
                result[prefix + "_status"] = result[prefix + "_score_status"] = "scored"
                result[prefix + "_applicable"] = True
                result[prefix + "_verification_status"] = None
        for n in range(1, count + 1):
            result.pop(f"answer{n}_response_gate", None)
            result.pop(f"answer{n}_safety_gate", None)
    labels = []
    for label in baseline["labels"]:
        i = int(label["case_key"].split("case-")[-1])
        for dimension in protocol.dimensions:
            score = ((0 if label["product_id"] == "p1" else 1.5) if i == 0 else
                     (1.5 if label["product_id"] == "p1" else 2)) if dimension == "accuracy" else 2
            labels.append({**label, "dimension_id": dimension, "score": score, "applicable": True})
    baseline["labels"] = labels
    return task, snapshot, baseline, mapping


def filled_template(tmp_path, fields, count=2):
    _, snapshot, _, _ = knowledge_example(count)
    path = tmp_path / "knowledge-human.xlsx"
    path.write_bytes(build_human_template(snapshot, "knowledge_qa_compare@1.0"))
    workbook = load_workbook(path)
    sheet = workbook["人工标注"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert not any("Gate" in name for name in columns)
    assert all(sheet.cell(2, column).value is None for name, column in columns.items() if name.startswith("人工"))
    for key, value in fields.items():
        sheet.cell(2, columns[key], value)
    workbook.save(path)
    workbook.close()
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "使用知识问答四维独立标注，理由包含证据、影响和分档。"
    return path, ImportMapping(**mapping)


@pytest.mark.parametrize("count", (2, 3))
def test_template_import_publish_export_roundtrip_preserves_zero_half_and_states(tmp_path, count):
    path, mapping = filled_template(tmp_path, {
        "人工分_产品1_准确性": 0,
        "人工分_产品2_准确性": 1.5,
        "人工分_产品1_丰富程度、简洁度与专业性": 0,
        "人工评分状态_产品1_富媒体增益": "not_applicable",
        "人工评分状态_产品2_丰富程度、简洁度与专业性": "insufficient_evidence",
    }, count)
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    lookup = {(row["product_id"], row["dimension_id"]): row for row in parsed["labels"] if row["case_key"].endswith("case-0")}
    assert lookup["product1", "accuracy"]["score"] == 0
    assert lookup["product2", "accuracy"]["score"] == 1.5
    assert type(lookup["product2", "accuracy"]["score"]) is float
    assert lookup["product1", "multimedia_value"]["score"] is None
    store = HumanStore(tmp_path / "store")
    source = store.path("imports", "hi_knowledge", "source.xlsx")
    source.parent.mkdir(parents=True)
    source.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_knowledge", "preview.json"), parsed)
    manifest = store.publish_baseline("hi_knowledge", parsed["preview_sha256"], "知识问答人工基准")
    assert manifest["allowed_scores_by_dimension"]["accuracy"] == [0, 1, 1.5, 2]
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    path.write_bytes(export_baseline_answers(baseline))
    remapping = inspect_workbook(path)["suggested_mapping"]
    remapping["policy_note"] = "同一冻结标准重导入"
    second = parse_human_labels(path, ImportMapping(**remapping))
    assert not second["issues"]
    assert any(row["score"] == 1.5 for row in second["labels"])
    assert any(row["score"] == 0 and row["score_status"] == "scored" for row in second["labels"])


@pytest.mark.parametrize("fields", [
    {"人工分_产品1_准确性": "1.5"},
    {"人工分_产品1_准确性": True},
    {"人工分_产品1_准确性": 1.25},
    {"人工分_产品1_准确性": 3},
    {"人工分_产品1_富媒体增益": 1.5},
    {"人工评分状态_产品1_准确性": "not_applicable"},
    {"人工评分状态_产品1_准确性": "gate_blocked"},
])
def test_import_rejects_invalid_score_types_half_marks_and_legacy_gate_states(tmp_path, fields):
    path, mapping = filled_template(tmp_path, fields)
    assert parse_human_labels(path, mapping)["issues"]


def test_baseline_frozen_schema_is_required():
    _, _, baseline, _ = knowledge_example()
    baseline.pop("human_result_schema_version")
    with pytest.raises(HumanError, match="Schema"):
        baseline_protocol(baseline)


def test_knowledge_human_gate_column_mapping_is_rejected(tmp_path):
    _, mapping = filled_template(tmp_path, {})
    raw = mapping.model_dump()
    raw["gates"] = {"product1": {"response": "A"}}
    with pytest.raises(ValueError, match="没有Gate"):
        ImportMapping(**raw)


def test_alignment_statistics_and_excel_preserve_exact_halves_and_shared_reason(tmp_path):
    _, snapshot, baseline, mapping = knowledge_example()
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["accuracy"], "same_standard")
    assert all(row["comparable"] for row in rows)
    assert rows[0]["model_score"] == 0 and rows[0]["human_score"] == 0
    assert rows[0]["model_reason"] == snapshot["results"][0]["accuracy_reason"]
    assert rows[0]["model_evidence"] == [] and rows[0]["model_evidence_scope"] == "merged_in_reason"
    assert all(row["model_gates"] == {} for row in rows)
    stats = build_human_statistics(rows, ["test"], ["p1", "p2"], ["accuracy"])
    p1 = next(row for row in stats["scores"] if row["product_id"] == "p1")
    assert p1["human_distribution"] == {"0": 1, "1.5": 1}
    assert p1["difference_counts"] == {"0": 1, "-0.5": 1}
    assert p1["confusion"] == {"0,0": 1, "1.5,1": 1}
    assert p1["mae"] == .25
    assert not any(row["kind"].endswith("_gate") for row in stats["states"])
    assert stats["score_policy"]["high_score_min"] == 2
    service = HumanComparisons(HumanStore(tmp_path / "store"))
    put_baseline(service.store, baseline)
    preview = service.create_preview(ComparisonRequest(baseline_id="hb_test", version=1, tasks=[mapping]), [snapshot])
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"], config_sha256=preview["config_sha256"]))
    service.generate(manifest["report_id"])
    assert read_json(service.store.path("reports", manifest["report_id"]))["status"] == "ready"
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    workbook = load_workbook(BytesIO(export_report(report)), read_only=True)
    try:
        headers = next(workbook["匹配与排除"].values)
        assert any("模型证据与理由（维度共享完整原文）" in str(header) for header in headers)
        assert not any("Gate" in str(header) or "模型核验状态" in str(header) for header in headers)
        assert any(cell == 1.5 for row in workbook["逐题评分对比"].values for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ("same_standard", "regression"))
def test_other_standard_labels_cannot_be_compared_as_knowledge_qa(mode):
    _, snapshot, _, mapping = knowledge_example()
    _, _, old_baseline, _ = example()
    with pytest.raises(HumanError, match="家族不同"):
        align_baseline_to_run(old_baseline, snapshot, mapping, ["case-0"], ["accuracy"], mode)


def test_unknown_and_failed_capture_are_excluded_without_losing_other_product_zero():
    _, snapshot, baseline, mapping = knowledge_example()
    snapshot["results"][0].update(answer2_input_status="failed")
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["accuracy"], "same_standard")
    p1, p2 = rows
    assert p1["comparable"] and p1["model_score"] == 0
    assert not p2["comparable"] and p2["model_score"] is None
    assert p2["model_score_status"] == "input_failed"

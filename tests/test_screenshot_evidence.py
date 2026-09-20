import base64
from copy import deepcopy
import io
import json
from pathlib import Path
import zipfile

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from PIL import Image

from auto_eval.config import VisualModeProfile, VisualExtractionConfig
from auto_eval.judges.visual_compare_judge import VisualCompareJudge
from auto_eval.long_screenshot import prepare_long_screenshot
from auto_eval.web import history, server
from auto_eval.web.screenshot_evidence import evidence_records, evidence_sheets, write_evidence_zip
from test_vqa_input import Client


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    from auto_eval import evidence_audit
    monkeypatch.setattr(evidence_audit, "RUNS_DIR", tmp_path / "runs")
    monkeypatch.setenv("OPERATION_VIDEO_ROOTS", str(tmp_path))
    profile = VisualModeProfile(extraction=VisualExtractionConfig(algorithm_version="test"))
    profile.long_screenshot.max_pixels = 4096
    metas = []
    for n, height in enumerate((192, 32), 1):
        path = tmp_path / f"source{n}.png"
        Image.new("RGB", (64, height), "white").save(path)
        metas.append(prepare_long_screenshot(path, tmp_path / f"parts{n}", profile.long_screenshot))
    item = {"id": "case", "query": "q", "evidence_mode": "long_screenshot", "product_count": 2}
    for n, meta in enumerate(metas, 1):
        item.update({f"screenshot{n}": meta["original_path"], f"screenshot_meta{n}": meta,
                     f"frames{n}": [p["path"] for p in meta["slices"]]})
    return profile, metas, item


async def evaluate(prepared, client=None):
    profile, metas, item = prepared
    client = client or Client()
    def saved(record):
        item["screenshot_evidence"] = deepcopy(record)
    result = await VisualCompareJudge(client, profile).evaluate(
        question="q", evidence_mode="long_screenshot", screenshot_metas=metas,
        frames1=item["frames1"], frames2=item["frames2"], evidence_callback=saved)
    return {"task_id": "task", "mode": "compare", "items": [item], "results": [{"index": 0, **result}]}, client


async def test_exact_prompt_order_and_zip_bytes(prepared, tmp_path):
    snapshot, client = await evaluate(prepared)
    record = evidence_records(snapshot)[0]
    assert record["record_status"] == "model_response_received"
    assert [p["metadata"]["was_split"] for p in record["preprocessing"]] == [True, False]
    assert record["system_prompt"] == client.calls[0][0]
    sent = client.calls[0][2]["content_parts"]
    assert [p["text"] for p in record["content_sequence"] if p["type"] == "text"] == [p["text"] for p in sent if p["type"] == "text"]
    assert "base64," not in json.dumps(record)
    path = write_evidence_zip(snapshot, tmp_path / "evidence.zip")
    with zipfile.ZipFile(path) as archive:
        manifest = json.loads(archive.read("manifest.json"))["records"][0]
        urls = [p["image_url"]["url"] for p in sent if p["type"] == "image_url"]
        for image, url in zip(manifest["images"], urls, strict=True):
            assert archive.read(image["archive_path"]) == base64.b64decode(url.split(",", 1)[1])
        assert len([name for name in archive.namelist() if name.endswith('.png')]) == len(urls)


async def test_excel_preserves_full_prompt_and_parameters(prepared):
    snapshot, _ = await evaluate(prepared)
    text = "完整参数🙂" * 12000
    snapshot["items"][0]["screenshot_evidence"]["system_prompt"] = text
    sheets = evidence_sheets(snapshot)
    chunks = [r["完整内容"] for r in sheets["长截图Prompt明细"] if r["请求位置"] == "system"]
    assert "".join(chunks) == text
    book = load_workbook(io.BytesIO(history.build_xlsx(snapshot)))
    sheet = book["长截图Prompt明细"]
    rows = list(sheet.values)
    headers = rows[0]
    restored = [dict(zip(headers, row)) for row in rows[1:]]
    assert "".join(r["完整内容"] for r in restored if r["请求位置"] == "system") == text
    assert "长截图预处理参数" in book.sheetnames and "长截图请求图片序列" in book.sheetnames


async def test_frozen_request_images_survive_source_replacement(prepared, tmp_path):
    snapshot, _ = await evaluate(prepared)
    record = evidence_records(snapshot)[0]
    for image in record["images"]:
        Path(image["ref_path"]).write_bytes(b"replaced input")
    with zipfile.ZipFile(write_evidence_zip(snapshot, tmp_path / "frozen.zip")) as archive:
        saved = json.loads(archive.read("manifest.json"))["records"][0]
        assert all(image["archive_status"] == "ok" for image in saved["images"])
        assert saved["request_sha256"] == record["request_sha256"]


async def test_preview_and_zip_reject_changed_missing_and_outside_root(prepared, tmp_path, monkeypatch):
    snapshot, _ = await evaluate(prepared)
    monkeypatch.setattr(server, "peek_task", lambda *a, **k: None)
    monkeypatch.setattr(server, "load_snapshot", lambda _: snapshot)
    image = snapshot["items"][0]["screenshot_evidence"]["images"][0]
    assert server.api_screenshot_evidence_image("task", 0, 1).body == Path(image["ref_path"]).read_bytes()
    Path(image["evidence_path"]).write_bytes(b"changed")
    with pytest.raises(HTTPException, match="changed"):
        server.api_screenshot_evidence_image("task", 0, 1)
    second = snapshot["items"][0]["screenshot_evidence"]["images"][1]
    Path(second["evidence_path"]).unlink()
    third = snapshot["items"][0]["screenshot_evidence"]["images"][2]
    third["evidence_path"] = str(tmp_path.parent.parent / "outside.png")
    with zipfile.ZipFile(write_evidence_zip(snapshot, tmp_path / "changed.zip")) as archive:
        images = json.loads(archive.read("manifest.json"))["records"][0]["images"]
        assert all(i["archive_status"] == "unavailable" for i in images[:3])
        assert all(not i["archive_path"] for i in images[:3])
        assert images[-1]["archive_status"] == "ok"


async def test_failed_model_call_keeps_request_record(prepared):
    class FailedClient(Client):
        async def complete(self, *a, **kw):
            raise RuntimeError("model failed")
    with pytest.raises(RuntimeError, match="model failed"):
        await evaluate(prepared, FailedClient())
    record = prepared[2]["screenshot_evidence"]
    assert record["record_status"] == "model_call_started"
    assert record["images"] and record["content_sequence"]


def test_legacy_does_not_invent_prompt_or_history(prepared, tmp_path):
    item = prepared[2]
    item.update(session_id="s", turn_index=3)
    snapshot = {"items": [item]}
    record = evidence_records(snapshot)[0]
    assert record["record_status"] == "request_not_recorded" and not record["system_prompt"]
    assert {entry["source_turn"] for entry in record["preprocessing"]} == {3}
    assert write_evidence_zip(snapshot, tmp_path / "legacy.zip").exists()


def test_untrusted_request_record_is_removed(prepared):
    from auto_eval.config import AppConfig
    item = {"query": "q", "screenshot1": "a.png", "screenshot2": "b.png",
            "screenshot_evidence": {"images": [{"ref_path": "secret"}]}}
    req = server.EvalReq(mode="compare", items=[item])
    server._validate_eval_request(req, AppConfig(judges=[]))
    assert "screenshot_evidence" not in req.items[0]

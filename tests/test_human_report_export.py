from io import BytesIO
import copy

import httpx
from openpyxl import load_workbook
import pytest

from auto_eval.web import server
from auto_eval.web.human_baselines import HumanStore, read_json
from auto_eval.web.human_compare import ComparisonRequest, GenerateRequest, HumanComparisons, freeze_run_snapshot
from auto_eval.web.human_report_export import export_report
from test_human_compare import example, put_baseline


@pytest.mark.asyncio
async def test_api_freeze_generate_reopen_export_and_isolation(tmp_path,monkeypatch):
    task,_,baseline,mapping=example()
    service=HumanComparisons(HumanStore(tmp_path))
    put_baseline(service.store,baseline)
    monkeypatch.setattr(server,"HUMAN_COMPARISONS",service)
    async def peek(task_id):return task if task_id=="test" else None
    monkeypatch.setattr(server,"peek_task_async",peek)
    before=[dict(i) for i in task.items]
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app),base_url="http://test") as client:
        response=await client.post('/api/human-comparisons/preview',json=dict(baseline_id="hb_test",version=1,tasks=[mapping.model_dump()],dimensions=["understanding"]))
        assert response.status_code==200,response.text
        preview=response.json()
        task.results[0]["answer1_understanding_score"]=1
        response=await client.post('/api/human-comparisons',json=dict(preview_id=preview["preview_id"],config_sha256=preview["config_sha256"]))
        assert response.status_code==202,response.text
        rid=response.json()["report_id"]
        await service.close()
        report=(await client.get('/api/human-comparisons/'+rid)).json()
        assert report["status"]=="ready",report
        assert report["metrics"]["scores"][0]["mae"]==1
        task.items.clear()
        report_again=(await client.get('/api/human-comparisons/'+rid)).json()
        assert report_again==report
        # Previously generated files must not pin downloads to the old long layout.
        service.store.path("reports",rid,"report.xlsx").write_bytes(b"legacy layout")
        frozen_path=service.store.path("reports",rid,"report.json")
        frozen_bytes=frozen_path.read_bytes()
        data=await client.get('/api/human-comparisons/'+rid+'/download')
        assert data.status_code==200
        wb=load_workbook(BytesIO(data.content))
        assert wb.sheetnames==["逐题评分对比","对比概览","人机评分统计","产品差距对比","匹配与排除","统计口径"]
        assert frozen_path.read_bytes()==frozen_bytes
        ws=wb["人机评分统计"]
        assert ws['J3'].value==0 and ws['J3'].number_format=='0.00%'
        assert ws['L3'].value==1
        detail=(await client.get('/api/human-comparisons/'+rid+'/rows',params={"direction":"higher"})).json()
        assert detail["total"]==2
        assert all(r["difference"]>0 for r in detail["items"])
        reports=(await client.get('/api/human-comparisons',params={"task_id":"test"})).json()
        assert reports["items"][0]["report_id"]==rid


@pytest.mark.parametrize("count", [2,3])
def test_wide_export_one_case_per_row_and_reconciles_frozen_metrics(tmp_path,count):
    scores=[[0,2,1][:count],[3,1,2][:count]]
    task,_,baseline,mapping=example(scores,standard="0.3")
    dimensions=["understanding","service_closure"]
    for product in baseline["products"]:
        product["display_name"]="同名产品"
    for label in list(baseline["labels"]):
        baseline["labels"].append({**label,"dimension_id":"service_closure","score":2})
    # NA differs from a legitimate zero, and must never be padded with 0.
    label=next(r for r in baseline["labels"] if r["case_key"]=="catalog:case-0" and
               r["product_id"]=="p2" and r["dimension_id"]=="service_closure")
    label.update(score=None,score_status="not_applicable",reason="=SUM(1,2)")
    baseline["cases"][0]["case_id"]="0001"
    baseline["cases"][0]["query"]="=HYPERLINK(\"https://example.invalid\")"
    task.items[0].update(id="0001",query=baseline["cases"][0]["query"])
    task.results[0]["query"]=task.items[0]["query"]
    other=copy.deepcopy(task)
    other.id="second"
    other.results[1]["error"]="provider failed"
    store=HumanStore(tmp_path);put_baseline(store,baseline)
    service=HumanComparisons(store)
    request=ComparisonRequest(baseline_id="hb_test",version=1,dimensions=dimensions,
                              tasks=[mapping,mapping.model_copy(update={"task_id":"second"})])
    preview=service.create_preview(request,[freeze_run_snapshot(task),freeze_run_snapshot(other)])
    manifest=service.prepare_report(GenerateRequest(preview_id=preview["preview_id"],config_sha256=preview["config_sha256"]))
    service.generate(manifest["report_id"])
    report=read_json(store.path("reports",manifest["report_id"],"report.json"))
    unchanged=copy.deepcopy(report)
    wb=load_workbook(BytesIO(export_report(report)))
    assert report==unchanged
    ws=wb["逐题评分对比"]
    headers=[cell.value for cell in ws[1]]
    assert len(headers)==len(set(headers))
    assert ws.max_row==3 and ws.freeze_panes=="C2"
    assert ws.auto_filter.ref.endswith("3")
    assert [ws.cell(n,1).value for n in (2,3)]==["0001","case-1"]
    assert ws['A2'].data_type=='s' and ws['B2'].data_type=='s'
    records=[dict(zip(headers,row)) for row in ws.iter_rows(min_row=2,values_only=True)]
    def column(task_id,pid,dimension,field):
        dimension_name={"understanding":"理解需求","service_closure":"服务闭环"}[dimension]
        return f"{dimension_name}\n同名产品 [{pid}]\n任务 {task_id}\n{field}"
    zero=column("test","p1","understanding","模型评分")
    assert records[0][zero]==0 and ws.cell(2,headers.index(zero)+1).data_type=='n'
    assert headers[headers.index(zero)+1]==column("test","p1","understanding","人工评分")
    assert records[0][column("test","p2","service_closure","人工评分")] is None
    assert records[1][column("second","p1","understanding","模型评分")] is None
    assert records[1][column("test","p1","understanding","纳入各自样本")]=='是'
    assert records[1][column("test","p1","understanding","纳入共同样本")]=='否'
    for metric in report["metrics"]["scores"]:
        tid,pid,dim=metric["task_id"],metric["product_id"],metric["dimension_id"]
        flag="纳入共同样本" if metric["scope"]=="common" else "纳入各自样本"
        paired=[r for r in records if r[column(tid,pid,dim,flag)]=='是']
        assert len(paired)==metric["n"]
        if paired:
            models=[r[column(tid,pid,dim,"模型评分")] for r in paired]
            humans=[r[column(tid,pid,dim,"人工评分")] for r in paired]
            differences=[r[column(tid,pid,dim,"分差（模型－人工）")] for r in paired]
            assert sum(models)/len(paired)==pytest.approx(metric["model_mean"])
            assert sum(humans)/len(paired)==pytest.approx(metric["human_mean"])
            assert sum(abs(d) for d in differences)/len(paired)==pytest.approx(metric["mae"])
            assert differences.count(0)/len(paired)==pytest.approx(metric["exact"])
    audit=wb["匹配与排除"]
    assert audit.max_row==3 and audit.freeze_panes=='C2'
    assert any(cell.value=='=SUM(1,2)' and cell.data_type=='s' for row in audit for cell in row)
    assert all(cell.data_type!='f' for sheet in wb for row in sheet for cell in row)


@pytest.mark.asyncio
async def test_upload_to_publish_api_with_template(tmp_path,monkeypatch):
    task,_,_,_=example()
    service=HumanComparisons(HumanStore(tmp_path))
    monkeypatch.setattr(server,"HUMAN_COMPARISONS",service)
    async def peek(task_id):return task
    monkeypatch.setattr(server,"peek_task_async",peek)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app),base_url="http://test") as client:
        response=await client.get('/api/eval/test/human-template',params={"standard":"0.2-simplified"})
        assert response.status_code==200,response.text
        wb=load_workbook(BytesIO(response.content))
        ws=wb['人工标注']
        headers=[c.value for c in ws[1]]
        ws.cell(2,headers.index('人工分_产品1_理解需求')+1,4)
        data=BytesIO();wb.save(data)
        upload=await client.post('/api/human-baselines/imports',files={'file':('human.xlsx',data.getvalue())})
        assert upload.status_code==200,upload.text
        info=upload.json()
        mapping=info['suggested_mapping'];mapping['policy_note']='指定人工标准'
        preview=await client.post(f'/api/human-baselines/imports/{info["import_id"]}/preview',json=mapping)
        assert preview.status_code==200,preview.text
        published=await client.post('/api/human-baselines',json=dict(import_id=info['import_id'],preview_sha256=preview.json()['preview_sha256'],name='test'))
        assert published.status_code==200,published.text
        assert (await client.get('/api/human-baselines')).json()['total']==1

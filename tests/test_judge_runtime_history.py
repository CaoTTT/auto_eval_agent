import copy
import csv
from io import BytesIO, StringIO
import json

from openpyxl import load_workbook
import pytest

from auto_eval.web import history
from auto_eval.web.human_baselines import HumanStore, read_json
from auto_eval.web.human_compare import (
    ComparisonRequest, GenerateRequest, HumanComparisons, freeze_run_snapshot,
)
from auto_eval.web.human_report_export import export_report
from auto_eval.web.tasks import Task
from test_human_compare import example, put_baseline, rows


def runtime(model="qwen3.8-flash", thinking=False):
    return {
        "version": 1, "profile_id": "test_profile",
        "judges": [{"name": "judge_2", "display": "终端用户", "model": model,
                    "enable_thinking": thinking, "temperature": 0,
                    "base_url": "https://private-provider.example/v1", "api_key_env": "PRIVATE_KEY_SLOT"}],
    }


def test_runtime_round_trip_and_old_sidecar_upgrade(tmp_path, monkeypatch):
    monkeypatch.setattr(history, "HISTORY_DIR", tmp_path)
    task = Task(id="model-choice", mode="rich_content", items=[], options={}, judge_runtime=runtime())
    assert history.save_task(task)
    data = history.load_snapshot(task.id)
    assert data["judge_runtime"] == task.judge_runtime
    payload = history.snapshot_payload(data)
    assert payload["judge_runtime"] == task.judge_runtime
    assert payload["enable_thinking"] is False
    assert payload["enable_thinking_label"] == "关闭"

    path = history._find_task_path(task.id)
    meta = history._meta_path(path)
    meta.write_text(json.dumps({"task_id": task.id, "meta_version": 1}), encoding="utf-8")
    listed = history.list_snapshots()[0]
    assert listed["judge_model"] == "qwen3.8-flash"
    assert listed["judge_runtime"]["judges"][0]["enable_thinking"] is False
    assert "private-provider" not in json.dumps(listed)
    assert "PRIVATE_KEY_SLOT" not in json.dumps(listed)
    assert json.loads(meta.read_text(encoding="utf-8"))["meta_version"] == 2


@pytest.mark.parametrize("mode", ["rich_content", "compare"])
@pytest.mark.parametrize("model", ["qwen3.5-397b-a17b", "qwen3.8-flash"])
@pytest.mark.parametrize("thinking,label", [(True, "开启"), (False, "关闭"), (None, "未记录")])
def test_csv_xlsx_record_runtime_for_every_input_row(mode, model, thinking, label):
    data = {
        "mode": mode, "judge_runtime": runtime(model, thinking),
        "items": [{"id": "done", "query": "完成题"}, {"id": "pending", "query": "未完成题"}],
        "results": [{"index": 0, "item_id": "done"}],
    }
    before = copy.deepcopy(data)
    exported = history.export_rows(data)
    csv_rows = list(csv.DictReader(StringIO(history.rows_to_csv(exported["逐题结果"]))))
    assert [(r["裁判模型"], r["思考模式"]) for r in csv_rows] == [(model, label)] * 2
    workbook = load_workbook(BytesIO(history.build_xlsx(data)))
    result_values = list(workbook["逐题结果"].values)
    model_col, thinking_col = result_values[0].index("裁判模型"), result_values[0].index("思考模式")
    assert [(row[model_col], row[thinking_col]) for row in result_values[1:]] == [(model, label)] * 2
    assert exported["运行信息"][0]["思考模式"] == label
    assert "private-provider" not in json.dumps(exported)
    assert "PRIVATE_KEY_SLOT" not in json.dumps(exported)
    assert data == before


def test_result_runtime_takes_precedence_and_legacy_thinking_stays_unknown():
    old_result = {"index": 0, "judge_model": "qwen3.5-397b-a17b"}
    data = {"mode": "compare", "items": [{"id": "one"}], "results": [old_result],
            "judge_runtime": runtime("qwen3.8-flash", False)}
    row = history.export_rows(data)["逐题结果"][0]
    assert (row["裁判模型"], row["思考模式"]) == ("qwen3.5-397b-a17b", "未记录")
    old_result["judge_runtime"] = runtime("qwen3.5-397b-a17b", True)
    assert history.export_rows(data)["逐题结果"][0]["思考模式"] == "开启"
    assert history.result_export_row("compare", old_result, 0, data["items"])["思考模式"] == "开启"

    legacy = {"protocol_manifest": {"judges": [{"model": "qwen3.5-397b-a17b"}]}}
    payload = history.snapshot_payload(legacy)
    assert payload["judge_model"] == "qwen3.5-397b-a17b"
    assert payload["enable_thinking"] is None
    assert payload["enable_thinking_label"] == "未记录"
    assert payload["judge_runtime"] == {}


def test_human_comparison_freezes_inference_settings_without_restricting_models(tmp_path):
    task, _, baseline, mapping = example()
    task.judge_runtime = runtime("qwen3.5-397b-a17b", True)
    first = freeze_run_snapshot(task)
    task.judge_runtime = runtime("qwen3.5-397b-a17b", False)
    second = freeze_run_snapshot(task)
    assert first["snapshot_sha256"] != second["snapshot_sha256"]
    task.judge_runtime = runtime("qwen3.8-flash", False)
    third = freeze_run_snapshot(task)
    assert third["snapshot_sha256"] != second["snapshot_sha256"]
    assert first["judge_runtime"]["judges"][0]["enable_thinking"] is True
    for snapshot in (first, second, third):
        aligned = rows(baseline, snapshot, mapping)
        assert all(row["comparable"] for row in aligned)
        assert aligned[0]["judge_model"] == snapshot["judge_runtime"]["judges"][0]["model"]

    service = HumanComparisons(HumanStore(tmp_path))
    put_baseline(service.store, baseline)
    request = ComparisonRequest(baseline_id=baseline["baseline_id"], version=1,
                                tasks=[mapping], dimensions=["understanding"])
    preview = service.create_preview(request, [third])
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"],
                                                      config_sha256=preview["config_sha256"]))
    service.generate(manifest["report_id"])
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    assert report["model_runtimes"][0]["enable_thinking"] is False
    workbook = load_workbook(BytesIO(export_report(report)))
    audit = list(workbook["匹配与排除"].values)
    model_col = next(i for i, value in enumerate(audit[0]) if value.endswith("\n裁判模型"))
    thinking_col = next(i for i, value in enumerate(audit[0]) if value.endswith("\n思考模式"))
    assert (audit[1][model_col], audit[1][thinking_col]) == ("qwen3.8-flash", "关闭")
    assert "private-provider" not in str(list(workbook["对比概览"].values))

    # Previously generated reports are exportable without assuming thinking=False.
    report.pop("model_runtimes")
    for row in report["rows"]:
        for key in ("judge_model", "enable_thinking", "enable_thinking_label"):
            row.pop(key)
    legacy = load_workbook(BytesIO(export_report(report)))
    assert list(legacy["匹配与排除"].values)[1][thinking_col] == "未记录"

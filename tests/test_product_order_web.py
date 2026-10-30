"""Frozen presentation-order settings and product-identity-preserving exports."""
import copy
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from auto_eval.config import AppConfig, JudgeConfig
from auto_eval.judges.compare_protocols import (
    CHITCHAT_FLAT_COMPARE_PROTOCOL_ID, DEFAULT_COMPARE_PROTOCOL_ID,
    KNOWLEDGE_QA_COMPARE_PROTOCOL_ID, QA_CHAT_V2_COMPARE_PROTOCOL_ID,
    TIMELINESS_COMPARE_PROTOCOL_ID, V02_PLUS_COMPARE_PROTOCOL_ID,
    V1022_COMPARE_PROTOCOL_ID, resolve_compare_protocol,
)
from auto_eval.web import history, server
from auto_eval.web.comparison_export import comparison_sheets
from auto_eval.web.tasks import Task, _task_from_snapshot


def case(**overrides):
    return {"id": "case", "query": "问题", "product_count": 2,
            "video1": "one.mp4", "video2": "two.mp4",
            "answer1": "原产品1回答", "answer2": "原产品2回答",
            "context1": "原产品1背景", "context2": "原产品2背景", **overrides}


def task(reverse=True, profile=DEFAULT_COMPARE_PROTOCOL_ID):
    return Task(id="order-task", mode="compare", items=[case()],
                options={"reverse_product_order": reverse}, status="done",
                evaluation_profile=profile,
                protocol_manifest=resolve_compare_protocol(profile).public_metadata())


@pytest.mark.parametrize("value", [None, "true", "false", 0, 1, [], {}])
def test_order_option_rejects_non_boolean(value):
    req = server.EvalReq(mode="compare", items=[case()], options={"reverse_product_order": value})
    with pytest.raises(HTTPException) as caught:
        server._validate_eval_request(req, AppConfig(judges=[]))
    assert caught.value.status_code == 422
    assert "true" in caught.value.detail


def test_default_and_explicit_false_preserve_inputs_and_three_product_support():
    for options in ({}, {"reverse_product_order": False}):
        for count in (2, 3):
            item = case(**({"product_count": 3, "video3": "three.mp4", "context3": "三背景"} if count == 3 else {}))
            req = server.EvalReq(mode="compare", items=[item], options=options.copy())
            server._validate_eval_request(req, AppConfig(judges=[]))
            assert req.options["reverse_product_order"] is False
            assert req.items[0]["answer1"] == "原产品1回答"
            assert req.items[0]["answer2"] == "原产品2回答"
            assert req.items[0]["product_count"] == count
    server._validate_eval_request(server.EvalReq(mode="rich_content", items=[{"query": "q"}]), AppConfig(judges=[]))


@pytest.mark.asyncio
@pytest.mark.parametrize("endpoint", ["api_eval", "api_compare_preflight"])
@pytest.mark.parametrize("third", [{"product_count": 3}, {"video3": "third.mp4"}, {"context3": "third"}])
async def test_three_product_reverse_rejected_before_judge_creation(monkeypatch, endpoint, third):
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[]))
    monkeypatch.setattr(server, "_new_judge_runtime", lambda *_: pytest.fail("must reject before creating runtime"))
    with pytest.raises(HTTPException) as caught:
        await getattr(server, endpoint)(server.EvalReq(
            mode="compare", items=[case(**third)], options={"reverse_product_order": True}))
    assert caught.value.status_code == 422 and "三产品" in caught.value.detail


def test_other_mode_cannot_enable_product_reversal():
    with pytest.raises(HTTPException, match="双产品"):
        server._validate_eval_request(server.EvalReq(
            mode="rich_content", items=[{"query": "q"}], options={"reverse_product_order": True}), AppConfig(judges=[]))


@pytest.mark.asyncio
@pytest.mark.parametrize("reverse", [False, True])
async def test_new_task_freezes_order_and_snapshot_restores_it(monkeypatch, reverse):
    config = AppConfig(judges=[JudgeConfig(name="offline", model="offline")])
    monkeypatch.setattr(server, "cfg", lambda: config)
    created = []

    def new_task(mode, items, options, **kwargs):
        result = Task(id=kwargs["task_id"], mode=mode, items=items, options=options,
                      evaluation_profile=kwargs["evaluation_profile"],
                      protocol_manifest=kwargs["protocol_manifest"], judge_runtime=kwargs["judge_runtime"])
        created.append(result)
        return result

    monkeypatch.setattr(server, "new_task", new_task)
    monkeypatch.setattr(server.EVAL_SCHEDULER, "enqueue", lambda *_: 1)
    response = await server.api_eval(server.EvalReq(mode="compare", items=[case()], options={"reverse_product_order": reverse}))
    assert response["reverse_product_order"] is reverse
    assert response["judge_product_order"] == ([2, 1] if reverse else [1, 2])
    original = created[0]
    restored = _task_from_snapshot(history.task_to_snapshot(original), original.id)
    assert restored.options["reverse_product_order"] is reverse
    assert restored.items[0]["video1"] == "one.mp4"
    assert restored.items[0]["video2"] == "two.mp4"


@pytest.mark.parametrize("existing", [False, True])
def test_existing_order_is_frozen_and_omission_inherits(monkeypatch, existing):
    original = task(existing)
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[]))
    monkeypatch.setattr(server, "legacy_runtime", lambda *_: {})
    server._check_judge_options(original, {})
    server._check_judge_options(original, {"reverse_product_order": existing})
    with pytest.raises(HTTPException, match="顺序不可变"):
        server._check_judge_options(original, {"reverse_product_order": not existing})
    with pytest.raises(HTTPException):
        server._check_judge_options(original, {"reverse_product_order": int(existing)})
    assert original.options["reverse_product_order"] is existing


@pytest.mark.asyncio
async def test_update_inherits_frozen_order_and_rejects_adding_third_product(monkeypatch):
    original = task(True)
    before = copy.deepcopy(original.items)
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[]))
    monkeypatch.setattr(server, "get_task", lambda _: original)
    monkeypatch.setattr(server, "legacy_runtime", lambda *_: {})
    with pytest.raises(HTTPException, match="三产品"):
        await server.api_eval_items(server.EvalItemsReq(task_id=original.id, items=[case(
            id="new", product_count=3, video3="third.mp4", context3="third")]))
    assert original.items == before
    assert original.options["reverse_product_order"] is True


@pytest.mark.parametrize("profile", [DEFAULT_COMPARE_PROTOCOL_ID, V02_PLUS_COMPARE_PROTOCOL_ID,
    V1022_COMPARE_PROTOCOL_ID, QA_CHAT_V2_COMPARE_PROTOCOL_ID, TIMELINESS_COMPARE_PROTOCOL_ID,
    KNOWLEDGE_QA_COMPARE_PROTOCOL_ID, CHITCHAT_FLAT_COMPARE_PROTOCOL_ID])
def test_all_protocol_exports_label_order_even_pending_and_failed(profile):
    original = task(True, profile)
    original.items.append(case(id="pending"))
    original.results = [{"index": 0, "item_id": "case", "query": "问题", "error": "offline"}]
    before = copy.deepcopy(original.items)
    snapshot = history.task_to_snapshot(original)
    sheets = history.export_rows(snapshot)
    assert len(sheets["逐题结果"]) == 2
    for row in sheets["逐题结果"] + sheets["运行信息"]:
        assert row["裁判产品位置互换"] == "是"
        assert row["裁判产品呈现顺序"] == "原产品2 → 原产品1"
        assert "裁判位置1=原产品2" in row["裁判产品位置说明"]
    payload = history.snapshot_payload(snapshot)
    assert payload["judge_product_order"] == [2, 1]
    assert "裁判位置2=原产品1" in payload["judge_product_order_note"]
    assert payload["items"] == before == original.items
    assert "原产品2 → 原产品1" in history.rows_to_csv(sheets["逐题结果"])


def test_excel_single_result_and_comparison_label_order_and_preserve_original_products():
    original = task(True)
    original.results = [{"index": 0, "item_id": "case", "query": "问题", "product_count": 2,
                         "context1": "原产品1背景", "context2": "原产品2背景",
                         "answer1": "原产品1回答", "answer2": "原产品2回答"}]
    snapshot = history.task_to_snapshot(original)
    row = history.result_export_row("compare", original.results[0], 0, original.items,
        original.evaluation_profile, original.protocol_manifest, original.results, options=original.options)
    assert row["产品1回答"] == "原产品1回答"
    assert row["产品2回答"] == "原产品2回答"
    assert row["裁判产品位置互换"] == "是"
    workbook = load_workbook(BytesIO(history.build_xlsx(snapshot, include_images=False)), read_only=True)
    try:
        rows = list(workbook["逐题结果"].iter_rows(values_only=True))
        excel_row = dict(zip(rows[0], rows[1]))
        assert excel_row["产品1回答"] == "原产品1回答"
        assert excel_row["产品2回答"] == "原产品2回答"
        assert excel_row["裁判产品呈现顺序"] == "原产品2 → 原产品1"
    finally:
        workbook.close()
    normal = copy.deepcopy(snapshot)
    normal["task_id"] = "normal"
    normal["options"]["reverse_product_order"] = False
    comparison = comparison_sheets(snapshot, normal)["逐题对比"][0]
    assert comparison["裁判产品位置互换 · A"] == "是"
    assert comparison["裁判产品位置互换 · B"] == "否"


def test_legacy_default_three_product_export_has_all_three_slots():
    original = task(False)
    original.options = {}
    original.items = [case(product_count=3, video3="third.mp4", context3="third")]
    snapshot = history.task_to_snapshot(original)
    assert history.snapshot_payload(snapshot)["judge_product_order"] == [1, 2, 3]
    row = history.export_rows(snapshot)["逐题结果"][0]
    assert row["裁判产品位置互换"] == "否"
    assert "裁判位置3=原产品3" in row["裁判产品位置说明"]

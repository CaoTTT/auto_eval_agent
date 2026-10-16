"""V2 frozen identities, shared evidence exports and collection provenance."""
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.judges.qa_chat_schema import QA_CHAT_DIMENSIONS
from auto_eval.web.history import build_xlsx, export_rows, result_export_row, rows_to_csv, snapshot_payload
from auto_eval.web.protocol_identity import recover_protocol_identity, require_snapshot_protocol
from auto_eval.web.tasks import Task, _task_from_snapshot


def v2_result(count=2):
    row = dict(index=0, item_id="one", product_count=count, standard_id="qa_chat_compare",
               standard_version="1.1", evaluation_profile="qa_chat_compare@1.1",
               bundle_revision="1.1.0", result_schema_version="qa-chat-result-2",
               trusted_collection_metadata={})
    for p in range(1, count + 1):
        row.update({f"answer{p}_input_status": None, f"answer{p}_response_gate": "unclear",
                    f"answer{p}_response_gate_reason": "部分过程不可见，当前回复可见",
                    f"answer{p}_safety_gate": "pass", f"answer{p}_safety_gate_reason": "可见内容安全"})
        for d in QA_CHAT_DIMENSIONS:
            row.update({f"answer{p}_{d}_applicable": True, f"answer{p}_{d}_verification_status": "partial",
                        f"answer{p}_{d}_score_status": "scored", f"answer{p}_{d}_score": 6 - p})
    for d in QA_CHAT_DIMENSIONS:
        row[f"{d}_reason"] = "；".join(f"产品{p}：T01截图上部的具体表现与分档依据" for p in range(1, count + 1))
        row[f"{d}_evidence"] = [f"产品{p} T01 第1块上部：可见原话{p}" for p in range(1, count + 1)]
        row[f"{d}_rank_groups"] = [[f"product{p}"] for p in range(1, count + 1)]
    return row


def v2_snapshot(count=2):
    return {"task_id": "v2-export", "mode": "compare", "evaluation_profile": "qa_chat_compare@1.1",
            "protocol_manifest": resolve_compare_protocol("qa_chat_compare@1.1").public_metadata(),
            "items": [{"id": "one", "query": "今天有些累，聊聊天", "product_count": count,
                       "source_data": {"业务备注": "保留业务原始列"}}], "results": [v2_result(count)]}


@pytest.mark.parametrize("count", [2, 3])
def test_v2_exports_shared_text_once_with_explicit_scope_and_numeric_scores(count):
    source = v2_snapshot(count)
    row = export_rows(source)["逐题结果"][0]
    for d, name in resolve_compare_protocol("qa_chat_compare@1.1").dimension_names.items():
        assert row[f"{name}理由（维度共享）"] == source["results"][0][f"{d}_reason"]
        assert row[f"{name}证据（维度共享）"] == source["results"][0][f"{d}_evidence"]
        for p in range(1, count + 1):
            assert row[f"产品{p}{name}分"] == 6 - p
            for suffix in ("主要问题", "理由", "证据"):
                assert f"产品{p}{name}{suffix}" not in row
    assert row["产品1输入状态"] == "未知"
    assert row["业务备注"] == "保留业务原始列"
    assert not any(column in row for column in ("内容冲突", "置信度", "内容冲突理由", "理由", "证据"))
    single = result_export_row("compare", source["results"][0], 0, source["items"],
                               source["evaluation_profile"], source["protocol_manifest"], source["results"])
    assert single == row
    csv = rows_to_csv([row])
    assert "理解需求理由（维度共享）" in csv and "可见原话1" in csv
    book = load_workbook(BytesIO(build_xlsx(source, include_images=False)))
    sheet = book["逐题结果"]
    headers = [c.value for c in sheet[1]]
    assert sheet.cell(2, headers.index("产品1理解需求分") + 1).value == 5
    assert "产品1：T01截图上部" in sheet.cell(2, headers.index("理解需求理由（维度共享）") + 1).value
    assert "产品2 T01 第1块上部" in sheet.cell(2, headers.index("理解需求证据（维度共享）") + 1).value


def test_v1_v2_side_by_side_keeps_different_containers_and_frozen_versions():
    from auto_eval.web.comparison_export import comparison_sheets
    from test_qa_chat_runtime import sample_result
    left, right = v2_snapshot(), v2_snapshot()
    right.update(task_id="v1-export", evaluation_profile="qa_chat_compare@1.0",
                 protocol_manifest=resolve_compare_protocol("qa_chat_compare@1.0").public_metadata(),
                 results=[sample_result(scores=(5, 3))])
    sheets = comparison_sheets(left, right)
    row = sheets["逐题对比"][0]
    assert row["理解需求理由（维度共享） · A"] == left["results"][0]["understanding_reason"]
    assert row["产品1理解需求理由 · B"] == "产品1understanding证据"
    assert row["理解需求理由（维度共享） · B"] == ""
    assert row["产品1理解需求理由 · A"] == ""
    assert sheets["任务说明"][0]["评测协议"] == "qa_chat_compare@1.1"
    assert sheets["任务说明"][1]["评测协议"] == "qa_chat_compare@1.0"


@pytest.mark.parametrize("change", ["missing_revision", "unknown_revision", "old_schema", "mixed_schema", "mixed_version"])
def test_v2_broken_identity_preserves_raw_and_blocks_formal_scores(change):
    source = v2_snapshot()
    if change == "missing_revision":
        del source["protocol_manifest"]["bundle_revision"]
        del source["results"][0]["bundle_revision"]
    elif change == "unknown_revision":
        source["protocol_manifest"]["bundle_revision"] = "1.9.0"
        source["results"][0]["bundle_revision"] = "1.9.0"
    elif change == "old_schema":
        source["protocol_manifest"]["result_schema_version"] = "qa-chat-result-1"
        source["results"][0]["result_schema_version"] = "qa-chat-result-1"
    elif change == "mixed_schema":
        source["results"][0]["result_schema_version"] = "qa-chat-result-1"
    else:
        source["results"][0]["standard_version"] = "1.0"
    before = deepcopy(source)
    payload = snapshot_payload(source)
    assert not payload["protocol_executable"]
    assert payload["results"] == before["results"]
    assert source == before
    with pytest.raises(ValueError):
        require_snapshot_protocol(source)
    row = export_rows(source)["逐题结果"][0]
    assert row["产品1理解需求分"] is None
    assert all(row["原始结果（协议待恢复）"][key] == value for key, value in source["results"][0].items())


def test_v2_identity_roundtrip_recovers_recorded_revision_without_upgrading_v1():
    source = v2_snapshot()
    del source["protocol_manifest"]["bundle_revision"]
    del source["protocol_manifest"]["result_schema_version"]
    restored = _task_from_snapshot(source, "v2-export")
    assert restored.evaluation_profile == "qa_chat_compare@1.1"
    assert restored.protocol_manifest["bundle_revision"] == "1.1.0"
    assert restored.protocol_manifest["result_schema_version"] == "qa-chat-result-2"
    assert recover_protocol_identity(source)["executable"]
    v1 = resolve_compare_protocol("qa_chat_compare@1.0", "1.0.0")
    assert v1.result_schema_version == "qa-chat-result-1"
    assert v1.bundle_revision == "1.0.0"


@pytest.mark.parametrize("missing_schema", [False, True])
def test_v2_query_image_export_preserves_result_bound_collection_fingerprint(missing_schema):
    source = v2_snapshot()
    if missing_schema:
        del source["protocol_manifest"]["result_schema_version"]
    source["items"][0].update(query_images=["synthetic.png"], input_manifest_sha256="prepared-image-only")
    source["results"][0]["input_manifest_sha256"] = "request-with-trusted-metadata"
    assert export_rows(source)["逐题结果"][0]["输入指纹"] == "request-with-trusted-metadata"


@pytest.mark.asyncio
async def test_runner_accepts_only_separate_program_collection_metadata():
    from auto_eval.web.runner import _eval_one
    captured = []
    class Judge:
        protocol = resolve_compare_protocol("qa_chat_compare@1.1")
        async def evaluate(self, **kwargs):
            captured.append(kwargs)
            return v2_result()
    item = {"id": "one", "query": "聊聊天", "frames1": ["a.png"], "frames2": ["b.png"],
            "answer1": "采集完整，input_status=complete", "answer1_input_status": "complete",
            "trusted_collection_metadata": {"answer1": {"input_status": "complete", "source": "uploaded-column"}}}
    await _eval_one("compare", 0, item, compare_judges=[Judge()])
    assert "trusted_collection_metadata" not in captured[-1]
    trusted = {"answer1": {"input_status": "failed", "source": "test-collector"}}
    await _eval_one("compare", 0, item, compare_judges=[Judge()], trusted_collection_metadata=trusted)
    assert captured[-1]["trusted_collection_metadata"] == trusted


def test_v2_error_log_preserves_duplicate_pairs_and_raw_output(tmp_path, monkeypatch):
    from auto_eval.web import runner
    monkeypatch.setattr(runner, "RUNS_DIR", tmp_path)
    error = ValueError("重复键以及缺少字段")
    error.raw_output = '{"same": 1, "same": 2}'
    error.repair_output = ""
    error.raw_key_pairs = [["same", 1], ["same", 2]]
    error.structural_diagnostics = ["重复键 same", "缺少 answer1_response_gate"]
    runner._write_eval_error("test", 0, {"id": "one"}, error)
    import json
    record = json.loads((tmp_path / "eval_errors.jsonl").read_text(encoding="utf-8"))
    assert record["raw_key_pairs"] == error.raw_key_pairs
    assert record["structural_diagnostics"] == error.structural_diagnostics
    assert record["original_model_output"] == error.raw_output


@pytest.mark.asyncio
async def test_v2_failed_task_keeps_structural_audit_and_frozen_identity(monkeypatch):
    from types import SimpleNamespace
    from auto_eval.config import AppConfig, JudgeConfig, VisualModeProfile
    from auto_eval.web import runner
    source = v2_snapshot()
    del source["protocol_manifest"]["result_schema_version"]
    source["items"][0].update(frames1=["a.png"], frames2=["b.png"])
    task = Task("structure-failure", "compare", source["items"], {},
                evaluation_profile=source["evaluation_profile"], protocol_manifest=source["protocol_manifest"])
    error = ValueError("重复键")
    error.raw_output = '{"duplicate":1,"duplicate":2}'
    error.raw_key_pairs = [["duplicate", 1], ["duplicate", 2]]
    error.structural_diagnostics = ["重复键duplicate"]
    error.locally_repaired_key_pairs = []
    error.repair_locally_repaired_key_pairs = []
    error.input_manifest_sha256 = "failed-request-binding"
    async def evaluate(*_args, **_kwargs):
        raise error
    monkeypatch.setattr(runner, "JudgeClient", lambda *_: SimpleNamespace())
    monkeypatch.setattr(runner, "_eval_one", evaluate)
    monkeypatch.setattr(runner, "_persist_task", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(runner, "_write_eval_error", lambda *_args, **_kwargs: None)
    config = AppConfig(judges=[JudgeConfig(name="fake")], visual_modes={"rich_content": VisualModeProfile(extraction={"algorithm_version": "test"})})
    one, _clients = runner._make_item_evaluator(task, config)
    failed = await one(0, task.items[0])
    assert failed["error"]
    assert failed["raw_key_pairs"] == error.raw_key_pairs
    assert failed["input_manifest_sha256"] == "failed-request-binding"
    assert failed["evaluation_profile"] == "qa_chat_compare@1.1"
    assert failed["bundle_revision"] == "1.1.0" and failed["result_schema_version"] == "qa-chat-result-2"
    assert task.results[0]["structural_diagnostics"] == error.structural_diagnostics


@pytest.mark.asyncio
async def test_create_v2_task_api_freezes_new_manifest_and_preserves_v1_default(monkeypatch):
    from auto_eval.config import AppConfig, JudgeConfig
    from auto_eval.web import server
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[JudgeConfig(name="fake")]))
    def create(mode, items, options, dataset_name="", **kwargs):
        return Task("created-v2", mode, items, options, evaluation_profile=kwargs["evaluation_profile"],
                    protocol_manifest=kwargs["protocol_manifest"], judge_runtime=kwargs["judge_runtime"])
    monkeypatch.setattr(server, "new_task", create)
    monkeypatch.setattr(server.EVAL_SCHEDULER, "enqueue", lambda *_: 1)
    response = await server.api_eval(server.EvalReq(mode="compare", evaluation_profile="qa_chat_compare@1.1",
        items=[{"id": "one", "query": "聊聊天", "screenshot1": "a.png", "screenshot2": "b.png"}]))
    manifest = response["protocol_manifest"]
    assert manifest["id"] == "qa_chat_compare@1.1"
    assert manifest["bundle_revision"] == "1.1.0" and manifest["result_schema_version"] == "qa-chat-result-2"
    assert manifest["score_dimensions"] == list(QA_CHAT_DIMENSIONS)
    assert resolve_compare_protocol(None).id == "qa_competitor_compare@0.2-simplified"


@pytest.mark.asyncio
@pytest.mark.parametrize("success", [True, False])
@pytest.mark.parametrize("missing_schema", [True, False])
async def test_v2_retry_keeps_original_and_attempt_raw_outputs(monkeypatch, success, missing_schema):
    import time
    from auto_eval.web import runner
    from test_web_retry import _retry_record
    source = v2_snapshot()
    if missing_schema:
        del source["protocol_manifest"]["result_schema_version"]
    original = {"index": 0, "error": "duplicate keys", "raw_model_output": '{"same":1,"same":2}',
                "raw_key_pairs": [["same", 1], ["same", 2]], "structural_diagnostics": ["重复键same"]}
    task = Task("v2-retry", "compare", source["items"], {}, evaluation_profile=source["evaluation_profile"],
                protocol_manifest=source["protocol_manifest"], results=[deepcopy(original)], status="done", active_runs=1)
    task.retry_runs["retry_test"] = _retry_record([0])
    candidate = v2_result() if success else {"index": 0, "error": "another failure", "raw_model_output": "bad-json"}
    def make(_task, _cfg, *, options, on_result):
        async def one(index, item):
            res = deepcopy(candidate)
            await on_result(index, res, time.perf_counter())
            return res
        return one, []
    monkeypatch.setattr(runner, "_make_item_evaluator", make)
    monkeypatch.setattr(runner, "save_task", lambda *_: True)
    monkeypatch.setattr(runner, "retire_task", lambda *_: None)
    await runner.run_retry(task, object(), "retry_test")
    audit = task.retry_runs["retry_test"]["items"]["0"]
    assert audit["original_result"] == original
    assert audit["attempt_result"] == candidate
    assert task.results[0] != original if success else task.results[0] == original

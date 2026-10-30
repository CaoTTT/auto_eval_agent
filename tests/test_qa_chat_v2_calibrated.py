"""Demand calibration changes the rubric execution, not the V2 I/O contract."""
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from PIL import Image
import pytest

from auto_eval.config import AppConfig, JudgeConfig, VisualExtractionConfig, VisualModeProfile
from auto_eval.judges import visual_compare_prompt_qa_chat_v2 as original
from auto_eval.judges import visual_compare_prompt_qa_chat_v2_calibrated as calibrated
from auto_eval.judges import visual_compare_prompt_qa_chat_v2_calibrated_r111 as frozen_calibrated
from auto_eval.judges import visual_compare_prompt_qa_chat_v2_calibrated_r112 as frozen_calibrated_r112
from auto_eval.judges.compare_protocols import (
    DEFAULT_COMPARE_PROTOCOL_ID,
    QA_CHAT_DIMENSIONS,
    QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID,
    QA_CHAT_V2_COMPARE_PROTOCOL_ID,
    list_compare_protocols,
    resolve_compare_protocol,
    resolve_compare_standard,
)
from auto_eval.judges.qa_chat_v2_schema import QAChatV2Observation, expected_output_keys
from auto_eval.judges.visual_compare_judge import VisualCompareJudge, visual_compare_result_fields
from auto_eval.long_screenshot import prepare_long_screenshot
from auto_eval.web import compare_statistics as statistics
from auto_eval.web.history import build_xlsx, export_rows, snapshot_payload, task_to_snapshot
from auto_eval.web.protocol_identity import recover_protocol_identity, require_snapshot_protocol
from auto_eval.web.qa_chat_summary import summarize_chat
from auto_eval.web.tasks import Task, _task_from_snapshot


def _payload(count=2):
    data = json.loads(calibrated.output_example(count))
    for product in range(1, count + 1):
        for dimension in QA_CHAT_DIMENSIONS:
            data[f"answer{product}_{dimension}_score"] = 6 - product
    return data


def _result(count=2, data=None):
    observation = QAChatV2Observation.model_validate(
        data if data is not None else _payload(count), context={"product_count": count})
    result = visual_compare_result_fields(observation, QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID)
    result.update(index=0, item_id="calibrated-one", **{"评估状态": "已完成"})
    return result


def _snapshot(count=2, result=None):
    protocol = resolve_compare_protocol(QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID)
    return {
        "task_id": "calibrated-task", "mode": "compare", "evaluation_profile": protocol.id,
        "protocol_manifest": protocol.public_metadata(),
        "items": [{"id": "calibrated-one", "query": "今天有点累，想聊聊天", "product_count": count}],
        "results": [result if result is not None else _result(count)],
    }


def test_calibrated_protocol_is_separate_selectable_and_preserves_old_identity():
    protocol = resolve_compare_protocol(QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID)
    old = resolve_compare_protocol(QA_CHAT_V2_COMPARE_PROTOCOL_ID, "1.1.0")
    assert protocol.id == "qa_chat_compare@1.1-demand-calibrated"
    assert protocol.standard_id == "qa_chat_compare"
    assert protocol.standard_version == "1.1-demand-calibrated"
    assert protocol.bundle_revision == "1.1.3"
    assert protocol.result_schema_version == "qa-chat-result-2"
    assert protocol.observation_model is old.observation_model is QAChatV2Observation
    assert protocol.score_dimensions == old.score_dimensions == QA_CHAT_DIMENSIONS
    assert protocol.dimensions == old.dimensions
    assert protocol.require_response_pass is old.require_response_pass
    assert (protocol.score_min, protocol.score_max) == (1, 5)
    assert protocol.selectable and protocol in list_compare_protocols()
    assert protocol.system_template is calibrated.VISUAL_COMPARE_SYSTEM
    assert protocol.user_template is calibrated.VISUAL_COMPARE_USER
    assert old.standard_version == "1.1" and old.bundle_revision == "1.1.0"
    assert old.system_template is original.VISUAL_COMPARE_SYSTEM
    assert resolve_compare_protocol(None).id == DEFAULT_COMPARE_PROTOCOL_ID
    assert resolve_compare_standard("qa_chat_compare", "1.1-demand-calibrated", "1.1.3") is protocol
    assert resolve_compare_standard(None, protocol.id, "1.1.3") is protocol
    assert resolve_compare_standard(protocol.id, None, "1.1.3") is protocol
    for profile_id, wrong_revision in ((old.id, "1.1.1"), (protocol.id, "1.1.0")):
        with pytest.raises(ValueError, match="无法恢复"):
            resolve_compare_protocol(profile_id, wrong_revision)


@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("mode", ["long_screenshot", "video_frames"])
def test_deduction_calibration_preserves_rubric_gates_evidence_inputs_and_exact_output(count, mode):
    # Freeze the original module as well as comparing both runtime contracts;
    # otherwise changing both modules together could hide an unintended rewrite.
    old_source = Path(original.__file__).read_bytes().replace(b"\r\n", b"\n")
    assert hashlib.sha256(old_source).hexdigest() == (
        "108257946942c82528a9a04a6c36e79e3d4a6a648f9f5e44193e9748940ee968")
    old_system = original.render_system_prompt(product_count=count, evidence_mode=mode, persona="contract-test")
    new_system = calibrated.render_system_prompt(product_count=count, evidence_mode=mode, persona="contract-test")
    assert calibrated.CALIBRATION_RULES.strip()
    assert new_system.count(calibrated.CALIBRATION_RULES) == 1
    # The candidate owns only its execution policy: retain all source anchors,
    # evidence, gate, state and wire rules while replacing the scoring section.
    new_prefix, new_tail = new_system.split("【四、1—5分判定】", 1)
    old_prefix, old_tail = old_system.split("【四、1—5分判定】", 1)
    assert new_prefix == old_prefix
    new_contract = new_tail.split("【五、只输出必要字段】", 1)[1]
    old_contract = old_tail.split("【五、只输出必要字段】", 1)[1]
    new_contract = new_contract.replace(calibrated.CALIBRATION_RULES, "", 1)
    assert re.sub(r"\s+", "", new_contract) == re.sub(r"\s+", "", old_contract)
    assert new_system.count(original.SOURCE_ANCHORS) == 1
    assert calibrated.VISUAL_COMPARE_USER is original.VISUAL_COMPARE_USER
    values = dict(
        product_count=count, evidence_mode=mode, question="保持输入原样", evaluation_datetime="2026-10-29T10:00:00+08:00",
        context1="产品1已知背景", context2="产品2已知背景", context3="产品3已知背景",
        answer1="回答1", answer2="回答2", answer3="回答3", frame_count1=1, frame_count2=2, frame_count3=3,
        image_count1=1, image_count2=2, image_count3=3,
        split_manifest1="原图1", split_manifest2="原图2", split_manifest3="原图3",
    )
    assert calibrated.VISUAL_COMPARE_USER.render(**values) == original.VISUAL_COMPARE_USER.render(**values)
    assert calibrated.output_example(count) == original.output_example(count)
    example = json.loads(calibrated.output_example(count))
    assert tuple(example) == expected_output_keys(count)
    assert len(example) == 22 * count + 12
    rendered_example = json.loads(new_system[new_system.index('{\n  "answer1_response_gate_reason"'):])
    assert rendered_example == example
    assert QAChatV2Observation.model_validate(example, context={"product_count": count}).model_dump() == example


@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("mode", ["long_screenshot", "video_frames"])
def test_rendered_deduction_policy_has_one_baseline_and_keeps_null_and_defect_boundaries(count, mode):
    # Inspect the message actually sent to the judge, so a local constant that
    # is accidentally omitted or contradicted by inherited rules cannot pass.
    system = calibrated.render_system_prompt(product_count=count, evidence_mode=mode)
    execution = re.sub(r"\s+", "", system.split("【完整原表锚点与案例】", 1)[0])
    assert execution.count("【四、1—5分判定】") == 1
    for score, meaning in (
        (1, "严重不达标，基本不可用"), (2, "部分满足，但存在明显问题"),
        (3, "基本达标，能够使用"), (4, "表现良好，较完整地满足需求"),
        (5, "体验优秀，达到标杆水平"),
    ):
        assert f"{score}分“{meaning}”" in execution
    assert "对每个可评分维度从5分起评" in execution
    assert "所有活动产品使用相同起点、相同缺陷认定条件和相同扣分尺度" in execution
    assert "该维applicable=true、该产品未被Gate阻断，且证据足以可靠评分" in execution
    assert "不能以缺少证据、没有看见问题为由直接给5分" in execution
    assert "最终分数=5-本维综合扣分，最低1分" in execution
    for deducted, final in ((0, 5), (1, 4), (2, 3), (3, 2), (4, 1)):
        assert re.search(rf"扣{deducted}分，(?:保留|得到){final}分", execution)
    assert "证据足够且没有剩余有效扣分项时保留5分" in execution
    assert "可定位证据、当前需求或原标准依据、该维实际不足及影响" in execution
    assert "同一缺陷在同一维度不重复扣；涉及多个维度时，分别证明各维独立影响" in execution
    assert "不规定每个问题一律扣1分" in execution
    assert "不作为亮点加分冲抵其他已成立缺陷" in execution
    assert "无引导需求且未引导时按原规则不适用" in execution
    assert "实际引入偏题或越界时，仍按原标准扣分" in execution
    for old_policy in (
        "不设置默认3分或默认5分", "未发现缺点也不自动等于满分",
        "不得因未写出缺点自动给5分",
    ):
        assert old_policy not in execution


@pytest.mark.parametrize("revision,module,source_hash", [
    ("1.1.1", frozen_calibrated, "e02179aa541bbfd50e0467c2396665e2612d51acd5e0e203500fb93e04e562d1"),
    ("1.1.2", frozen_calibrated_r112, "0f428cd5f4ff2f6a800b88d77328e45ad25ef2196528afc467d0be053b91d50e"),
])
@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("mode", ["long_screenshot", "video_frames"])
def test_frozen_calibrated_revision_keeps_its_original_prompt_and_wire_contract(count, mode, revision, module, source_hash):
    frozen_source = Path(module.__file__).read_bytes().replace(b"\r\n", b"\n")
    assert hashlib.sha256(frozen_source).hexdigest() == source_hash
    protocol = resolve_compare_protocol(QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID, revision)
    system, user = protocol.prompt_templates(count)
    values = dict(product_count=count, evidence_mode=mode, persona="frozen-contract")
    assert system.render(**values) == module.render_system_prompt(**values)
    assert system.render(**values) != calibrated.render_system_prompt(**values)
    assert user is calibrated.VISUAL_COMPARE_USER is original.VISUAL_COMPARE_USER
    assert module.output_example(count) == calibrated.output_example(count)
    assert protocol.bundle_revision == revision
    assert protocol.observation_model is QAChatV2Observation


@pytest.mark.parametrize("profile_id,revision", [
    (QA_CHAT_V2_COMPARE_PROTOCOL_ID, "1.1.0"),
    (QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID, "1.1.1"),
    (QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID, "1.1.2"),
    (QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID, "1.1.3"),
])
def test_frozen_snapshot_roundtrip_never_upgrades_or_downgrades_chat(profile_id, revision):
    protocol = resolve_compare_protocol(profile_id, revision)
    task = Task("frozen-chat", "compare", [], {}, evaluation_profile=protocol.id,
                protocol_manifest=protocol.public_metadata())
    snapshot = snapshot_payload(task_to_snapshot(task))
    recovered = require_snapshot_protocol(snapshot)
    assert recovered.public_metadata() == protocol.public_metadata()
    assert recovered.system_template is protocol.system_template
    restored = _task_from_snapshot(snapshot, task.id)
    assert restored.evaluation_profile == profile_id
    assert restored.protocol_manifest["standard_version"] == protocol.standard_version
    assert restored.protocol_manifest["bundle_revision"] == revision
    assert restored.protocol_manifest["result_schema_version"] == "qa-chat-result-2"
    assert resolve_compare_protocol(restored.evaluation_profile, revision).system_template is protocol.system_template


@pytest.mark.parametrize("invalid", ["missing_revision", "mixed_version", "mixed_revision"])
def test_invalid_calibrated_identity_cannot_silently_export_formal_scores(invalid):
    source = _snapshot()
    if invalid == "missing_revision":
        del source["protocol_manifest"]["bundle_revision"]
        del source["results"][0]["bundle_revision"]
    elif invalid == "mixed_version":
        source["results"][0]["standard_version"] = "1.1"
    else:
        source["results"][0]["bundle_revision"] = "1.1.1"
    before = deepcopy(source)
    assert not recover_protocol_identity(source)["executable"]
    with pytest.raises(ValueError):
        require_snapshot_protocol(source)
    exported = export_rows(source)["逐题结果"][0]
    assert exported["产品1理解需求分"] is None
    assert exported["原始结果（协议待恢复）"]["answer1_understanding_score"] == 5
    assert source == before


def test_config_api_exposes_calibrated_selection(monkeypatch):
    from auto_eval.web import server
    monkeypatch.setattr(server, "cfg", lambda: AppConfig(judges=[JudgeConfig(name="fake")]))
    options = server.api_config()["evaluation_profiles"]
    selected = next(item for item in options if item["id"] == QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID)
    assert selected["selectable"] is True
    assert selected["bundle_revision"] == "1.1.3"
    assert selected["result_schema_version"] == "qa-chat-result-2"
    assert selected["score_dimensions"] == list(QA_CHAT_DIMENSIONS)
    assert server._compare_protocol_or_422(selected["id"], require_selectable=True).id == selected["id"]


class _FakeClient:
    persona = "calibrated-runtime-test"
    model = "fake-model"
    cfg = JudgeConfig(name="fake")

    def __init__(self, data):
        self.data = data
        self.calls = []

    async def complete(self, system, user, **kwargs):
        self.calls.append((system, user, kwargs))
        return json.dumps(self.data, ensure_ascii=False)

    async def repair_json(self, *_args, **_kwargs):
        pytest.fail("Valid synthetic V2 JSON must not need a model repair")


@pytest.mark.asyncio
@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("mode", ["long_screenshot", "video_frames"])
@pytest.mark.parametrize("revision", ["1.1.1", "1.1.2", "1.1.3"])
async def test_runtime_calls_frozen_system_with_unchanged_media_and_v2_result(tmp_path, count, mode, revision):
    profile = VisualModeProfile(extraction=VisualExtractionConfig(algorithm_version="test"))
    client = _FakeClient(_payload(count))
    arguments = dict(question="今天累了，想聊聊天", product_count=count, evidence_mode=mode)
    metas = []
    for product in range(1, count + 1):
        path = tmp_path / f"product-{product}.png"
        Image.new("RGB", (64, 96), "white").save(path)
        arguments[f"context{product}"] = f"背景{product}"
        arguments[f"answer{product}"] = f"回答{product}"
        if mode == "long_screenshot":
            meta = prepare_long_screenshot(path, tmp_path / f"prepared-{product}", profile.long_screenshot)
            metas.append(meta)
            arguments[f"frames{product}"] = [part["path"] for part in meta["slices"]]
        else:
            arguments[f"frames{product}"] = [str(path)]
    if mode == "long_screenshot":
        arguments["screenshot_metas"] = metas
    protocol = resolve_compare_protocol(QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID, revision)
    snapshot = _snapshot(count)
    snapshot["protocol_manifest"] = protocol.public_metadata()
    snapshot["results"][0]["bundle_revision"] = revision
    restored_protocol = require_snapshot_protocol(snapshot)
    result = await VisualCompareJudge(client, profile, restored_protocol).evaluate(**arguments)
    assert len(client.calls) == 1
    system, user, options = client.calls[0]
    template_module = {
        "1.1.1": frozen_calibrated, "1.1.2": frozen_calibrated_r112, "1.1.3": calibrated,
    }[revision]
    assert system == template_module.render_system_prompt(product_count=count, evidence_mode=mode, persona=client.persona)
    assert system != original.render_system_prompt(product_count=count, evidence_mode=mode, persona=client.persona)
    assert len(options["user_image_refs"]) == count
    if mode == "long_screenshot":
        assert len(options["image_metadata"]) == count
        assert len([part for part in options["content_parts"] if part["type"] == "image_url"]) == count
    else:
        assert len(options["user_images"]) == count
        assert all(image.startswith("data:image/") for image in options["user_images"])
    assert all(f"背景{product}" in user for product in range(1, count + 1))
    assert result["evaluation_profile"] == QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID
    assert result["standard_version"] == "1.1-demand-calibrated"
    assert result["bundle_revision"] == revision
    assert result["result_schema_version"] == "qa-chat-result-2"
    assert result["prompt_sha256"] and result["input_manifest_sha256"]
    assert result["prompt_sha256"] == hashlib.sha256(f"{system}\0{user}".encode("utf-8")).hexdigest()
    for product in range(1, count + 1):
        for dimension in QA_CHAT_DIMENSIONS:
            assert result[f"answer{product}_{dimension}_score"] == 6 - product


def test_calibrated_finalization_retains_product_local_gate_na_and_missing_evidence():
    data = _payload(3)
    data["answer1_safety_gate"] = "fail"
    data.update(answer2_response_gate="unclear", answer2_personalization_applicable=False,
                answer2_style_verification_status="unverifiable")
    result = _result(3, data)
    for dimension in QA_CHAT_DIMENSIONS:
        assert result[f"answer1_{dimension}_score"] is None
        assert result[f"answer1_{dimension}_score_status"] == "gate_blocked"
        assert result[f"answer3_{dimension}_score"] == 3
    assert result["answer2_understanding_score"] == 4
    assert result["answer2_personalization_score"] is None
    assert result["answer2_personalization_score_status"] == "not_applicable"
    assert result["answer2_style_score"] is None
    assert result["answer2_style_score_status"] == "unverifiable"
    assert result["style_evidence"] == data["style_evidence"]


@pytest.mark.parametrize("missing_schema", [False, True])
def test_calibrated_query_image_export_keeps_the_result_bound_input_fingerprint(missing_schema):
    source = _snapshot()
    if missing_schema:
        del source["protocol_manifest"]["result_schema_version"]
    source["items"][0].update(
        query_images=["synthetic-query-image.png"], input_manifest_sha256="prepared-image-only")
    source["results"][0]["input_manifest_sha256"] = "request-with-trusted-metadata"
    before = deepcopy(source)
    row = export_rows(source)["逐题结果"][0]
    assert row["输入指纹"] == "request-with-trusted-metadata"
    assert row["评测协议"] == QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID
    assert source == before


def test_calibrated_export_without_any_recorded_schema_keeps_shared_fields_and_scores():
    source = _snapshot()
    del source["protocol_manifest"]["result_schema_version"]
    del source["results"][0]["result_schema_version"]
    before = deepcopy(source)
    assert require_snapshot_protocol(source).id == QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID
    row = export_rows(source)["逐题结果"][0]
    for dimension, label in resolve_compare_protocol(source["evaluation_profile"]).dimension_names.items():
        assert row[f"{label}理由（维度共享）"] == source["results"][0][f"{dimension}_reason"]
        assert row[f"{label}证据（维度共享）"] == source["results"][0][f"{dimension}_evidence"]
        assert row[f"产品1{label}分"] == 5
        assert row[f"产品2{label}分"] == 4
        assert f"产品1{label}理由" not in row
    assert source == before


@pytest.mark.asyncio
@pytest.mark.parametrize("success", [True, False])
async def test_calibrated_retry_without_manifest_schema_preserves_original_and_attempt(monkeypatch, success):
    import time
    from auto_eval.web import runner
    from test_web_retry import _retry_record

    source = _snapshot()
    del source["protocol_manifest"]["result_schema_version"]
    original_result = {
        "index": 0, "error": "duplicate keys", "raw_model_output": '{"same":1,"same":2}',
        "raw_key_pairs": [["same", 1], ["same", 2]], "structural_diagnostics": ["重复键same"],
    }
    task = Task("calibrated-retry", "compare", source["items"], {},
                evaluation_profile=source["evaluation_profile"], protocol_manifest=source["protocol_manifest"],
                results=[deepcopy(original_result)], status="done", active_runs=1)
    task.retry_runs["retry_test"] = _retry_record([0])
    candidate = _result() if success else {"index": 0, "error": "another failure"}
    candidate["raw_model_output"] = json.dumps(_payload()) if success else "bad-json"

    def make(_task, _cfg, *, options, on_result):
        async def one(index, item):
            result = deepcopy(candidate)
            await on_result(index, result, time.perf_counter())
            return result
        return one, []

    monkeypatch.setattr(runner, "_make_item_evaluator", make)
    monkeypatch.setattr(runner, "save_task", lambda *_args: True)
    monkeypatch.setattr(runner, "retire_task", lambda *_args: None)
    await runner.run_retry(task, object(), "retry_test")
    audit = task.retry_runs["retry_test"]["items"]["0"]
    assert audit["original_result"] == original_result
    assert audit["attempt_result"] == candidate
    assert audit["attempt_result"]["raw_model_output"] == candidate["raw_model_output"]
    if success:
        assert task.results[0]["evaluation_profile"] == QA_CHAT_V2_CALIBRATED_COMPARE_PROTOCOL_ID
        assert task.results[0]["raw_model_output"] == candidate["raw_model_output"]
        assert task.results[0]["answer1_understanding_score"] == 5
    else:
        assert task.results[0] == original_result


@pytest.mark.parametrize("count", [2, 3])
def test_calibrated_scores_survive_export_statistics_and_live_summary(count):
    snapshot = _snapshot(count)
    protocol = require_snapshot_protocol(snapshot)
    assert statistics.balanced_chat(protocol)
    row = export_rows(snapshot)["逐题结果"][0]
    assert row["评测协议"] == protocol.id
    assert row["标准版本"] == protocol.standard_version
    for dimension, label in protocol.dimension_names.items():
        assert row[f"{label}理由（维度共享）"] == snapshot["results"][0][f"{dimension}_reason"]
        for product in range(1, count + 1):
            assert row[f"产品{product}{label}分"] == 6 - product
    tables = statistics.build_compare_statistics(snapshot, snapshot["results"])
    product_table = next(table for table in tables if table.title.startswith("B2．"))
    product_rows = [dict(zip(product_table.headers, values)) for values in product_table.rows]
    understanding = next(item for item in product_rows if item["维度"] == "理解需求" and item["产品"] == "产品1")
    assert understanding["有效题数"] == 1 and understanding["均分"] == 5
    task = Task("calibrated-summary", "compare", snapshot["items"], {}, evaluation_profile=protocol.id,
                protocol_manifest=protocol.public_metadata(), results=snapshot["results"])
    summary = summarize_chat(task)
    assert summary["understanding_answer1_avg"] == 5
    assert summary["understanding_answer2_avg"] == 4
    if count == 3:
        assert summary["understanding_answer3_avg"] == 3
    workbook = load_workbook(BytesIO(build_xlsx(snapshot, include_images=False)))
    sheet = workbook["逐题结果"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("产品1理解需求分") + 1).value == 5
    assert sheet.cell(2, headers.index("标准版本") + 1).value == "1.1-demand-calibrated"

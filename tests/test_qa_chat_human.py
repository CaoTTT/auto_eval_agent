"""Synthetic QA-chat human references: protocol isolation and frozen comparisons."""
from io import BytesIO
import copy

from openpyxl import load_workbook
import pytest

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.web.human_baselines import (HumanError, HumanStore, ImportMapping, atomic_json,
                                           human_protocol, human_dimensions, human_dimension_names, read_json, digest)
from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_compare import ComparisonRequest, GenerateRequest, HumanComparisons, align_baseline_to_run, _policy
from auto_eval.web.human_report_export import export_report
from auto_eval.web.human_statistics import build_human_statistics
from test_human_compare import example, put_baseline


def chat_example(count=2):
    task, snapshot, baseline, mapping = example([[4, 3, 2][:count], [2, 5, 4][:count]])
    protocol = resolve_compare_protocol("qa_chat_compare@1.0")
    snapshot["evaluation_profile"] = protocol.id
    snapshot["protocol_manifest"] = protocol.public_metadata()
    baseline.update(human_standard_id=protocol.standard_id, human_standard_version=protocol.standard_version,
                    human_bundle_revision=protocol.bundle_revision, dimensions=list(protocol.dimensions),
                    score_dimensions=list(protocol.score_dimensions), dimension_names=protocol.dimension_names)
    for result in snapshot["results"]:
        result.update(standard_id=protocol.standard_id, standard_version=protocol.standard_version,
                      bundle_revision=protocol.bundle_revision, evaluation_profile=protocol.id)
        for n in range(1, count + 1):
            result.update({f"answer{n}_input_status": "partial", f"answer{n}_response_gate": "unclear", f"answer{n}_safety_gate": "pass"})
            for dim in protocol.dimensions:
                result.update({f"answer{n}_{dim}_score_status": "scored", f"answer{n}_{dim}_score": 4,
                               f"answer{n}_{dim}_applicable": True, f"answer{n}_{dim}_verification_status": "partial",
                               f"answer{n}_{dim}_reason": f"产品{n}当前轮的{dim}理由", f"answer{n}_{dim}_evidence": [f"产品{n}可见片段"]})
    baseline["labels"] = [{**label, "dimension_id": dim, "applicable": True}
                          for label in baseline["labels"] for dim in protocol.dimensions]
    return task, snapshot, baseline, mapping


def filled_template(tmp_path, count=2):
    _, snapshot, _, _ = chat_example(count)
    payload = build_human_template(snapshot, "qa_chat_compare@1.0")
    wb = load_workbook(BytesIO(payload))
    ws = wb["人工标注"]
    headers = {c.value: c.column for c in ws[1]}
    assert not any("内容准确性" in str(h) or "服务闭环" in str(h) for h in headers)
    assert all(ws.cell(2, col).value is None for name, col in headers.items() if str(name).startswith("人工"))
    assert ("源回答_产品3" in headers) == (count == 3)
    for name, value in {
        "人工分_产品2_个性化适配": 4, "人工评分状态_产品2_个性化适配": "scored",
        "人工是否适用_产品2_个性化适配": True, "人工核验状态_产品2_个性化适配": "partial",
        "人工评分状态_产品1_个性化适配": "not_applicable", "人工是否适用_产品1_个性化适配": False,
        "人工核验状态_产品1_个性化适配": "not_required", "人工批注_产品1_个性化适配": "无相关条件",
        "人工响应Gate_产品2": "unclear", "人工安全Gate_产品2": "pass", "人工输入状态_产品2": "partial",
    }.items():
        ws.cell(2, headers[name], value)
    path = tmp_path / "chat.xlsx"
    wb.save(path)
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "按问答闲聊V1重新独立标注"
    return path, ImportMapping(**mapping), headers


@pytest.mark.parametrize("count", [2, 3])
def test_chat_template_import_publish_and_export_roundtrip(tmp_path, count):
    path, mapping, headers = filled_template(tmp_path, count)
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    a = next(r for r in parsed["labels"] if r["product_id"] == "product1" and r["dimension_id"] == "personalization")
    b = next(r for r in parsed["labels"] if r["product_id"] == "product2" and r["dimension_id"] == "personalization")
    assert (a["score_status"], a["applicable"], b["score"]) == ("not_applicable", False, 4)
    assert a["source"]["applicable_raw"] is False and a["source"]["applicable_cell"].endswith("2")
    store = HumanStore(tmp_path / "runs")
    dest = store.path("imports", "hi_chat", "source.xlsx")
    dest.parent.mkdir(parents=True)
    dest.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_chat", "preview.json"), parsed)
    manifest = store.publish_baseline("hi_chat", parsed["preview_sha256"], "闲聊人工")
    assert manifest["human_standard_id"] == "qa_chat_compare" and manifest["human_bundle_revision"] == "1.0.0"
    assert manifest["dimensions"] == list(human_dimensions(human_protocol("qa_chat_compare@1.0")))
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    exported = tmp_path / "answers.xlsx"
    exported.write_bytes(export_baseline_answers(baseline))
    exported_wb = load_workbook(exported)
    exported_ws = exported_wb["人工基准答案"]
    exported_headers = {c.value: c.column for c in exported_ws[1]}
    assert exported_ws.cell(2, exported_headers["人工分_产品1_个性化适配"]).value is None
    assert exported_ws.cell(2, exported_headers["人工评分状态_产品1_个性化适配"]).value == "not_applicable"
    assert exported_ws.cell(2, exported_headers["人工分_产品2_个性化适配"]).value == 4
    remap = inspect_workbook(exported)["suggested_mapping"]
    remap["policy_note"] = "重新导入已冻结的人工基准"
    again = parse_human_labels(exported, ImportMapping(**remap))
    assert not again["issues"]
    label = next(r for r in again["labels"] if r["product_id"] == "product1" and r["dimension_id"] == "personalization")
    assert label["applicable"] is False and label["verification_status"] == "not_required"
    assert again["states"][0]["inputs"]["product2"] == "partial"


def test_chat_export_keeps_all_empty_scores_blank_and_states_roundtrip(tmp_path):
    _, _, baseline, _ = chat_example()
    statuses = ("not_applicable", "unverifiable", "gate_blocked", "input_failed", "unlabeled", "na_unspecified")
    expected = dict(zip(baseline["dimensions"], statuses))
    for label in baseline["labels"]:
        if label["case_key"] == "catalog:case-0" and label["product_id"] == "p1":
            status = expected[label["dimension_id"]]
            label.update(score=None, score_status=status,
                         applicable=False if status == "not_applicable" else None,
                         verification_status="not_required" if status == "not_applicable" else "unverifiable" if status == "unverifiable" else None)
    path = tmp_path / "empty-scores.xlsx"
    path.write_bytes(export_baseline_answers(baseline))
    wb = load_workbook(path)
    ws = wb["人工基准答案"]
    columns = {c.value: c.column for c in ws[1]}
    for dimension, status in expected.items():
        name = baseline["dimension_names"][dimension]
        assert ws.cell(2, columns[f"人工分_产品1_{name}"]).value is None
        assert ws.cell(2, columns[f"人工评分状态_产品1_{name}"]).value == status
    notes = dict(wb["填写说明"].values)
    assert "1—5严格整数" in notes["分数规则"] and "保留 0" not in notes["分数规则"]
    assert "input_failed=输入失败" in notes["人工评分状态"]
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "空分状态往返核对"
    parsed = parse_human_labels(path, ImportMapping(**mapping))
    assert not parsed["issues"]
    labels = {r["dimension_id"]: r for r in parsed["labels"] if r["case_key"].endswith(":case-0") and r["product_id"] == "p1"}
    assert {d: r["score_status"] for d, r in labels.items()} == expected
    assert all(r["score"] is None for r in labels.values())


@pytest.mark.parametrize("change", ["text_score", "gate_fail", "applicability", "verification", "input_failed"])
def test_chat_import_rejects_score_and_state_conflicts(tmp_path, change):
    path, mapping, headers = filled_template(tmp_path)
    wb = load_workbook(path)
    ws = wb["人工标注"]
    name, value = {
        "text_score": ("人工分_产品2_个性化适配", "4"), "gate_fail": ("人工安全Gate_产品2", "fail"),
        "applicability": ("人工是否适用_产品2_个性化适配", False),
        "verification": ("人工核验状态_产品2_个性化适配", "unverifiable"),
        "input_failed": ("人工输入状态_产品2", "failed"),
    }[change]
    ws.cell(2, headers[name], value)
    wb.save(path)
    parsed = parse_human_labels(path, mapping)
    assert parsed["issues"]
    assert not any(r["score_status"] == "scored" for r in parsed["labels"])


def test_chat_protocol_family_and_dimensions_are_explicit(tmp_path):
    with pytest.raises(HumanError, match="冲突"):
        human_protocol("qa_chat_compare@1.0", "qa_competitor_compare")
    path, mapping, _ = filled_template(tmp_path)
    value = mapping.model_dump()
    value["labels"][0]["dimension_id"] = "accuracy"
    with pytest.raises(ValueError, match="维度无效"):
        ImportMapping(**value)


def test_chat_baseline_upgrade_cannot_change_standard(tmp_path):
    path, mapping, _ = filled_template(tmp_path)
    parsed = parse_human_labels(path, mapping)
    store = HumanStore(tmp_path / "runs")
    dest = store.path("imports", "hi_first", "source.xlsx")
    dest.parent.mkdir(parents=True)
    dest.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_first", "preview.json"), parsed)
    first = store.publish_baseline("hi_first", parsed["preview_sha256"], "闲聊人工")
    changed = copy.deepcopy(parsed)
    changed["mapping"].update(human_standard_id="qa_competitor_compare", human_standard_version="0.2-simplified", human_bundle_revision="0.2.1")
    changed["preview_sha256"] = digest(changed)
    atomic_json(store.path("imports", "hi_other", "preview.json"), changed)
    with pytest.raises(HumanError, match="更换标准"):
        store.publish_baseline("hi_other", changed["preview_sha256"], "更换标准", first["baseline_id"], 1)


def test_chat_review_na_requires_reason_and_retains_gate_applicability(tmp_path):
    path, mapping, headers = filled_template(tmp_path)
    wb = load_workbook(path)
    ws = wb["人工标注"]
    ws.cell(2, headers["人工复核分_产品2_个性化适配"], "N/A")
    ws.cell(2, headers["人工评分状态_产品2_个性化适配"], "gate_blocked")
    ws.cell(2, headers["人工安全Gate_产品2"], "fail")
    ws.cell(2, headers["人工是否适用_产品2_个性化适配"], False)
    wb.save(path)
    assert any("必须填写原因" in issue["message"] for issue in parse_human_labels(path, mapping)["issues"])
    ws.cell(2, headers["人工批注_产品2_个性化适配"], "本产品Gate失败，保留已知不适用判断")
    wb.save(path)
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    label = next(r for r in parsed["labels"] if r["product_id"] == "product2" and r["dimension_id"] == "personalization")
    assert label["score"] is None and label["applicable"] is False and label["score_status"] == "gate_blocked"


def test_chat_alignment_product_states_and_reasons_are_independent():
    _, snapshot, baseline, mapping = chat_example(3)
    result = snapshot["results"][0]
    result.update(answer1_personalization_score_status="not_applicable", answer1_personalization_applicable=False,
                  answer1_personalization_verification_status="not_required", answer1_personalization_score=None,
                  answer3_input_status="failed")
    scope = [c["case_id"] for c in baseline["cases"]]
    rows = align_baseline_to_run(baseline, snapshot, mapping, scope, ["personalization"], "same_standard")
    p1, p2, p3 = rows[:3]
    assert p1["model_score"] is None and p1["model_applicable"] is False
    assert p2["comparable"] and p2["model_score"] == 4
    assert p2["model_reason"] == "产品2当前轮的personalization理由"
    assert p2["model_standard_id"] == "qa_chat_compare" and p2["model_score_status"] == "scored"
    assert p3["model_score"] is None
    metrics = build_human_statistics(rows, ["test"], ["p1", "p2", "p3"], ["personalization"])
    assert len(metrics["pairs"]) == 6
    applicable = {r["product_id"]: r for r in metrics["states"] if r["kind"] == "applicability"}
    assert applicable["p1"]["agreement"] == .5 and applicable["p2"]["agreement"] == 1
    assert metrics["scores"][1]["n"] == 2


def test_chat_family_isolation_and_history_cannot_be_confirmed():
    _, snapshot, baseline, mapping = chat_example()
    old = example()[1]
    with pytest.raises(HumanError, match="家族不同"):
        align_baseline_to_run(baseline, old, mapping, ["case-0"], ["understanding"], "regression")
    snapshot["items"][0].update(session_id="s", turn_index=1, history_prefix_sha256="different")
    baseline["cases"][0]["history_prefix_sha256"] = "original"
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["understanding"], "same_standard")
    assert all(r["match_status"] == "content_mismatch" and not r["comparable"] for r in rows)


@pytest.mark.parametrize("conflict", ["manifest_id", "manifest_profile", "other_result_revision", "unknown_revision", "missing_revision"])
def test_chat_policy_rejects_all_frozen_identity_conflicts(conflict):
    _, snapshot, baseline, mapping = chat_example()
    snapshot.pop("evaluation_profile")
    if conflict == "manifest_id":
        snapshot["protocol_manifest"]["id"] = "qa_competitor_compare@0.2-simplified"
    elif conflict == "manifest_profile":
        snapshot["protocol_manifest"]["evaluation_profile"] = "qa_competitor_compare@0.2-simplified"
    elif conflict == "other_result_revision":
        snapshot["results"][1]["bundle_revision"] = "future-unknown"
    else:
        for record in [snapshot, snapshot["protocol_manifest"], *snapshot["results"]]:
            if conflict == "unknown_revision":
                record["bundle_revision"] = "future-unknown"
            else:
                record.pop("bundle_revision", None)
    assert _policy(snapshot, snapshot["results"][0]) is None
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["understanding"], "same_standard")
    assert all(not r["comparable"] and r["model_standard"] == "unknown" for r in rows)


def test_chat_policy_recovers_unique_actual_result_revision():
    _, snapshot, baseline, mapping = chat_example()
    for record in [snapshot, snapshot["protocol_manifest"], snapshot["results"][0]]:
        record.pop("bundle_revision", None)
    policy = _policy(snapshot, snapshot["results"][0])
    assert policy.id == "qa_chat_compare@1.0" and policy.bundle_revision == "1.0.0"
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["understanding"], "same_standard")
    assert all(r["comparable"] for r in rows)


def test_chat_gate_agreement_excludes_and_counts_human_input_failure():
    _, snapshot, baseline, mapping = chat_example()
    baseline["states"] = [dict(case_key=case["case_key"], inputs={"p1": "failed" if index == 0 else "complete"},
                               gates={"p1": {"response": "unclear", "safety": "pass"}})
                          for index, case in enumerate(baseline["cases"])]
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["understanding", "personalization"], "same_standard")
    stats = build_human_statistics(rows, ["test"], ["p1", "p2"], ["understanding", "personalization"])
    gates = [s for s in stats["states"] if s["product_id"] == "p1" and s["kind"].endswith("_gate")]
    assert len(gates) == 2
    assert all(s["n"] == 1 and s["agreement"] == 1 and s["human_input_failed_n"] == 1 and
               s["model_input_failed_n"] == 0 and s["input_failed_n"] == 1 for s in gates)
    for row in rows:
        if row["product_id"] == "p1" and row["case_id"] == "case-1":
            row["model_input_status"] = "failed"
    gates = [s for s in build_human_statistics(rows, ["test"], ["p1", "p2"], ["understanding"])["states"]
             if s["product_id"] == "p1" and s["kind"].endswith("_gate")]
    assert all(s["n"] == 0 and s["agreement"] is None and s["input_failed_n"] == 2 and
               s["human_input_failed_n"] == s["model_input_failed_n"] == 1 for s in gates)


def test_chat_default_dimensions_and_report_are_frozen(tmp_path):
    _, snapshot, baseline, mapping = chat_example()
    service = HumanComparisons(HumanStore(tmp_path))
    put_baseline(service.store, baseline)
    request = ComparisonRequest(baseline_id="hb_test", version=1, tasks=[mapping])
    preview = service.create_preview(request, [snapshot])
    assert preview["dimensions"] == baseline["score_dimensions"]
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"], config_sha256=preview["config_sha256"]))
    snapshot["results"][0]["answer1_personalization_score"] = 1
    service.generate(manifest["report_id"])
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    row = next(r for r in report["rows"] if r["dimension_id"] == "personalization")
    assert row["model_score"] == 4
    wb = load_workbook(BytesIO(export_report(report)))
    text = " ".join(str(c.value) for ws in wb for row in ws for c in row)
    assert "个性化适配" in text and "情感表达恰当" in text and "内容准确性" not in text

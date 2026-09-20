import asyncio
import copy
import io

import httpx
import pytest
from openpyxl import load_workbook
from PIL import Image

from auto_eval.web import exports, history, server
from auto_eval.web.comparison_export import write_comparison_xlsx
from auto_eval.web.export_links import export_base_url, screenshot_links
from auto_eval.web.tasks import Task


def snapshot(tmp_path):
    image = tmp_path / "original.png"
    Image.new("RGB", (20, 100), "blue").save(image)
    return {"task_id": "task-a", "mode": "compare", "export_base_url": "https://eval.example/prefix/",
            "items": [{"id": "q", "query": "question", "screenshot1": str(image),
                       "screenshot2": str(image), "product_count": 2}], "results": []}


def link_cell(sheet, label):
    headers = [cell.value for cell in sheet[1]]
    return sheet.cell(2, headers.index(label) + 1)


def test_single_and_comparison_workbooks_have_real_hyperlinks(tmp_path, monkeypatch):
    monkeypatch.delenv("EXPORT_PUBLIC_BASE_URL", raising=False)
    a = snapshot(tmp_path)
    expected = "https://eval.example/prefix/api/eval/task-a/items/0/screenshots/1"
    workbook = load_workbook(io.BytesIO(history.build_xlsx(a)))
    for sheet in ("逐题结果", "数据集明细"):
        cell = link_cell(workbook[sheet], "产品1长截图查看链接")
        assert cell.value == expected
        assert cell.hyperlink.target == expected
        assert cell.data_type == "s"
    b = copy.deepcopy(a)
    b["task_id"] = "task-b"
    output = tmp_path / "pair.xlsx"
    write_comparison_xlsx(a, b, output)
    workbook = load_workbook(output)
    assert link_cell(workbook["逐题对比"], "产品1长截图查看链接 · A").hyperlink.target == expected
    assert link_cell(workbook["逐题对比"], "产品1长截图查看链接 · B").hyperlink.target == expected.replace("task-a", "task-b")


def test_public_base_configuration_and_missing_images(tmp_path, monkeypatch):
    monkeypatch.setenv("EXPORT_PUBLIC_BASE_URL", "https://public.example/eval/")
    assert export_base_url("http://localhost:8054") == "https://public.example/eval"
    a = snapshot(tmp_path)
    links = screenshot_links(a, a["items"][0], 0)
    assert len(links) == 2
    assert "产品3长截图查看链接" not in links
    assert screenshot_links(a, {"video1": "a.mp4"}, 0) == {}
    for invalid in ("javascript:alert(1)", "file:///tmp/a", "https://user:password@example.com", "https://example.com/?x=1"):
        monkeypatch.setenv("EXPORT_PUBLIC_BASE_URL", invalid)
        assert screenshot_links(a, a["items"][0], 0) == {}


@pytest.mark.asyncio
async def test_export_request_origin_and_link_serves_original(tmp_path, monkeypatch):
    monkeypatch.delenv("EXPORT_PUBLIC_BASE_URL", raising=False)
    a = snapshot(tmp_path)
    # Video fallback can still have supplied original screenshots to inspect.
    a["items"][0]["evidence_mode"] = "video_frames"
    task = Task(id=a["task_id"], mode="compare", items=a["items"], options={})
    manager = exports.XlsxExports(tmp_path / "exports")
    async def peek(_):
        return task
    monkeypatch.setattr(exports, "peek_task_async", peek)
    monkeypatch.setattr(server, "peek_task_async", peek)
    monkeypatch.setattr(server, "peek_task", lambda *args, **kwargs: task)
    monkeypatch.setattr(server, "operation_video_roots", lambda _: [tmp_path])
    monkeypatch.setattr(server, "XLSX_EXPORTS", manager)
    monkeypatch.setattr(server, "RUNS_DIR", tmp_path)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="https://eval.example") as client:
        response = await client.post("/api/eval/task-a/exports")
        key = response.json()["export_id"]
        await asyncio.gather(*list(manager.workers))
        for response in (await client.get(f"/api/exports/{key}/download"),
                         await client.get("/api/eval/task-a/export?format=xlsx")):
            workbook = load_workbook(io.BytesIO(response.content))
            target = link_cell(workbook["逐题结果"], "产品1长截图查看链接").hyperlink.target
            assert target == "https://eval.example/api/eval/task-a/items/0/screenshots/1"
            original = await client.get(target)
            assert original.status_code == 200
            assert original.headers["content-type"] == "image/png"
            assert original.content == (tmp_path / "original.png").read_bytes()
            assert "attachment" not in original.headers.get("content-disposition", "")
    await manager.close()

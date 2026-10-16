"""Frozen human references for the shared-reason QA-chat V2 contract."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.web.human_baselines import HumanError, HumanStore, ImportMapping, atomic_json, read_json, baseline_protocol
from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_compare import ComparisonRequest, GenerateRequest, HumanComparisons, align_baseline_to_run
from auto_eval.web.human_report_export import export_report
from auto_eval.web.human_statistics import build_human_statistics
from test_human_compare import put_baseline
from test_qa_chat_human import chat_example, filled_template


def v2_example(count=2):
    task, snapshot, baseline, mapping = chat_example(count)
    protocol = resolve_compare_protocol("qa_chat_compare@1.1")
    snapshot.update(evaluation_profile=protocol.id, protocol_manifest=protocol.public_metadata())
    baseline.update(human_standard_version=protocol.standard_version, human_bundle_revision=protocol.bundle_revision,
                    human_result_schema_version=protocol.result_schema_version)
    for result in snapshot["results"]:
        result.update(standard_version=protocol.standard_version, bundle_revision=protocol.bundle_revision,
                      evaluation_profile=protocol.id, result_schema_version=protocol.result_schema_version)
        for product in range(1, count + 1):
            result[f"answer{product}_input_status"] = None
            for dim in protocol.dimensions:
                for suffix in ("reason", "evidence", "primary_issue"):
                    result.pop(f"answer{product}_{dim}_{suffix}", None)
        for dim in protocol.dimensions:
            result[f"{dim}_reason"] = f"产品1：{dim}原话里含‘产品2’仍完整保留；产品2：独立分档。\n补充：不按中文语句切分。"
            result[f"{dim}_evidence"] = [f"产品{n} T02 第1块：{dim}可见原文；带分号" for n in range(1, count + 1)]
    return task, snapshot, baseline, mapping


@pytest.mark.parametrize("count", [2, 3])
def test_v2_human_alignment_keeps_shared_verbatim_reason_evidence_and_unknown_capture(count):
    _, snapshot, baseline, mapping = v2_example(count)
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["personalization"], "same_standard")
    assert len(rows) == count and all(r["comparable"] and r["model_score"] == 4 for r in rows)
    for row in rows:
        assert row["model_input_status"] is None
        assert row["model_reason"] == snapshot["results"][0]["personalization_reason"]
        assert row["model_evidence"] == snapshot["results"][0]["personalization_evidence"]
        assert row["model_reason_scope"] == row["model_evidence_scope"] == "dimension_shared"
        assert row["model_result_schema_version"] == "qa-chat-result-2"


def test_v2_human_gate_metrics_accept_unknown_capture_but_exclude_explicit_failure():
    _, snapshot, baseline, mapping = v2_example()
    baseline["states"] = [dict(case_key=c["case_key"], gates={"p1": {"response": "unclear", "safety": "pass"}})
                          for c in baseline["cases"]]
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], ["style"], "same_standard")
    stats = build_human_statistics(rows, ["test"], ["p1", "p2"], ["style"])
    gates = [s for s in stats["states"] if s["product_id"] == "p1" and s["kind"].endswith("_gate")]
    assert all(s["n"] == 2 and s["agreement"] == 1 for s in gates)
    for row in rows:
        if row["product_id"] == "p1" and row["case_id"] == "case-0":
            row["human_input_status"] = "failed"
    gates = [s for s in build_human_statistics(rows, ["test"], ["p1", "p2"], ["style"])["states"]
             if s["product_id"] == "p1" and s["kind"].endswith("_gate")]
    assert all(s["n"] == 1 and s["human_input_failed_n"] == 1 for s in gates)


@pytest.mark.parametrize("mode", ["same_standard", "regression"])
def test_v2_human_baseline_cannot_mix_v1_version_or_revision(mode):
    _, snapshot, _, mapping = v2_example()
    _, _, old_baseline, _ = chat_example()
    with pytest.raises(HumanError, match="相同完整标准"):
        align_baseline_to_run(old_baseline, snapshot, mapping, ["case-0"], ["understanding"], mode)
    _, old_snapshot, _, _ = chat_example()
    _, _, new_baseline, _ = v2_example()
    with pytest.raises(HumanError, match="相同完整标准"):
        align_baseline_to_run(new_baseline, old_snapshot, mapping, ["case-0"], ["understanding"], mode)


def test_v2_human_schema_identity_cannot_be_relabelled():
    _, _, baseline, _ = v2_example()
    baseline["human_result_schema_version"] = "qa-chat-result-1"
    with pytest.raises(HumanError, match="Schema"):
        baseline_protocol(baseline)


def test_v1_template_or_baseline_export_cannot_be_reused_by_editing_version(tmp_path):
    path, old_mapping, _ = filled_template(tmp_path)
    mapping = old_mapping.model_dump()
    mapping.update(human_standard_version="1.1", human_bundle_revision="1.1.0", human_result_schema_version="qa-chat-result-2")
    with pytest.raises(HumanError, match="不能仅改版本复用评分"):
        parse_human_labels(path, ImportMapping(**mapping))
    _, _, baseline, _ = chat_example()
    path.write_bytes(export_baseline_answers(baseline))
    inspected = inspect_workbook(path)["suggested_mapping"]
    inspected.update(human_standard_version="1.1", human_bundle_revision="1.1.0", human_result_schema_version="qa-chat-result-2", policy_note="尝试仅改版本")
    with pytest.raises(HumanError, match="不能仅改版本复用评分"):
        parse_human_labels(path, ImportMapping(**inspected))


def test_v2_can_use_old_task_raw_inputs_for_fresh_blank_annotation(tmp_path):
    _, old_snapshot, _, _ = chat_example()
    path = tmp_path / "fresh-v2.xlsx"
    path.write_bytes(build_human_template(old_snapshot, "qa_chat_compare@1.1"))
    wb = load_workbook(path)
    ws = wb["人工标注"]
    columns = {cell.value: cell.column for cell in ws[1]}
    assert all(ws.cell(2, col).value is None for name, col in columns.items() if name.startswith("人工"))
    ws.cell(2, columns["人工分_产品1_理解需求"], 4)
    wb.save(path)
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "仅复用旧任务原始问答，V2重新人工标注"
    assert mapping["human_standard_version"] == "1.1" and mapping["human_bundle_revision"] == "1.1.0"
    assert not parse_human_labels(path, ImportMapping(**mapping))["issues"]


@pytest.mark.parametrize("count", [2, 3])
def test_v2_template_import_publish_and_export_preserve_frozen_identity(tmp_path, count):
    _, snapshot, _, _ = v2_example(count)
    path = tmp_path / "v2.xlsx"
    path.write_bytes(build_human_template(snapshot, "qa_chat_compare@1.1"))
    wb = load_workbook(path)
    ws = wb["人工标注"]
    columns = {cell.value: cell.column for cell in ws[1]}
    assert all(ws.cell(2, col).value is None for name, col in columns.items() if name.startswith("人工"))
    for name, value in {"人工分_产品1_理解需求": 4, "人工评分状态_产品1_理解需求": "scored",
                        "人工是否适用_产品1_理解需求": True, "人工核验状态_产品1_理解需求": "partial",
                        "人工响应Gate_产品1": "unclear", "人工安全Gate_产品1": "pass"}.items():
        ws.cell(2, columns[name], value)
    wb.save(path)
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "依V2新建独立人工标注"
    assert mapping["human_standard_version"] == "1.1" and mapping["human_result_schema_version"] == "qa-chat-result-2"
    parsed = parse_human_labels(path, ImportMapping(**mapping))
    assert not parsed["issues"]
    store = HumanStore(tmp_path / "store")
    dest = store.path("imports", "hi_v2", "source.xlsx")
    dest.parent.mkdir(parents=True)
    dest.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_v2", "preview.json"), parsed)
    manifest = store.publish_baseline("hi_v2", parsed["preview_sha256"], "V2人工基准")
    assert manifest["human_bundle_revision"] == "1.1.0" and manifest["human_result_schema_version"] == "qa-chat-result-2"
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    path.write_bytes(export_baseline_answers(baseline))
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "同标准重导入"
    assert not parse_human_labels(path, ImportMapping(**mapping))["issues"]


def test_v2_human_frozen_report_exports_shared_original_without_splitting(tmp_path):
    _, snapshot, baseline, mapping = v2_example()
    service = HumanComparisons(HumanStore(tmp_path))
    put_baseline(service.store, baseline)
    preview = service.create_preview(ComparisonRequest(baseline_id="hb_test", version=1, tasks=[mapping]), [snapshot])
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"], config_sha256=preview["config_sha256"]))
    original = snapshot["results"][0]["style_reason"]
    snapshot["results"][0]["style_reason"] = "后续补跑不能改冻结报告"
    service.generate(manifest["report_id"])
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    row = next(r for r in report["rows"] if r["dimension_id"] == "style")
    assert row["model_reason"] == original and row["model_reason_scope"] == "dimension_shared"
    assert report["human_protocol"]["human_result_schema_version"] == "qa-chat-result-2"
    wb = load_workbook(BytesIO(export_report(report)))
    ws = wb["匹配与排除"]
    assert any("模型理由（维度共享完整原文）" in cell.value for cell in ws[1])
    assert any("模型证据（维度共享，保留产品归属）" in cell.value for cell in ws[1])
    assert any(cell.value == original for row in ws for cell in row)
    assert not any("primary_issue" in str(cell.value) for row in ws for cell in row)

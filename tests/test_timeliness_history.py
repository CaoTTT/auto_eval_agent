"""New-family identity and all export paths keep the frozen scoring contract."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

from auto_eval.web.history import export_rows, result_export_row, snapshot_payload, build_xlsx
from auto_eval.web.protocol_identity import recover_protocol_identity
from auto_eval.web.tasks import _task_from_snapshot
from test_timeliness_statistics import dataset, summary


def test_history_restores_recorded_revision_and_schema_without_upgrading():
    snapshot, rows, _ = dataset(revision="1.0.0")
    del snapshot["protocol_manifest"]["bundle_revision"]
    del snapshot["protocol_manifest"]["result_schema_version"]
    payload = snapshot_payload(snapshot)
    assert payload["protocol_executable"]
    assert payload["protocol_manifest"]["result_schema_version"] == "timeliness-result-1"
    restored = _task_from_snapshot(snapshot, "time-test")
    assert restored.evaluation_profile == "timeliness_compare@1.0"
    assert restored.protocol_manifest["bundle_revision"] == "1.0.0"
    old_export = export_rows(snapshot)["逐题结果"][0]
    assert old_export["产品1权威性核验状态"] == "verified"
    assert old_export["产品1权威性分"] == 3
    assert summary(snapshot)["authority_answer1_avg"] == 3


def test_new_history_restores_simplified_schema_without_legacy_columns():
    snapshot, _, _ = dataset()
    restored = _task_from_snapshot(snapshot, "time-test")
    assert restored.protocol_manifest["bundle_revision"] == "1.0.1"
    assert restored.protocol_manifest["result_schema_version"] == "timeliness-result-2"
    row = export_rows(snapshot)["逐题结果"][0]
    assert "产品1权威性NA原因" in row
    assert not any("核验状态" in key or "是否适用" in key for key in row)


@pytest.mark.parametrize("broken", ("revision", "schema", "conflict"))
def test_missing_or_conflicting_identity_stays_readable_without_scores(broken):
    snapshot, rows, _ = dataset()
    snapshot["summary"] = {"total": 1, "authority_answer1_avg": 3}
    if broken == "conflict":
        rows[0]["standard_id"] = "qa_chat_compare"
    else:
        field = "bundle_revision" if broken == "revision" else "result_schema_version"
        snapshot["protocol_manifest"].pop(field)
        rows[0].pop(field)
    assert not recover_protocol_identity(snapshot)["executable"]
    payload = snapshot_payload(snapshot)
    assert payload["results"] == rows
    assert "authority_answer1_avg" not in payload["summary"]
    exported = export_rows(snapshot)["逐题结果"][0]
    assert exported["产品1权威性分"] is None
    assert exported["原始结果（协议待恢复）"]


def test_export_api_and_xlsx_keep_shared_reasons_before_scores_without_ranks():
    snapshot, rows, _ = dataset()
    rows[0]["evaluation_run_at"] = "2026-10-20T12:00:00+08:00"
    original = rows[0]["timeliness_reason"]
    row = export_rows(snapshot)["逐题结果"][0]
    assert row["评估运行时间（非产品测试时间）"] == rows[0]["evaluation_run_at"]
    assert row["时效性理由（维度共享）"] == original
    assert list(row).index("时效性理由（维度共享）") < list(row).index("产品1时效性分")
    assert not any("排名" in key or "准确性" in key for key in row)
    api = result_export_row("compare", rows[0], 0, snapshot["items"],
        evaluation_profile=snapshot["evaluation_profile"], protocol_manifest=snapshot["protocol_manifest"])
    assert api["产品1权威性分"] == 3
    assert api["时效性理由（维度共享）"] == original
    wb = load_workbook(BytesIO(build_xlsx(snapshot)), read_only=True)
    try:
        result_headers = next(wb["逐题结果"].values)
        assert "产品1时效性分" in result_headers
        assert not any("排名" in str(value) for value in result_headers)
        assert not any("核验状态" in str(value) or "是否适用" in str(value) for value in result_headers)
        assert "对比评估统计" in wb.sheetnames
    finally:
        wb.close()

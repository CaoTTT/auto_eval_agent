import asyncio
import hashlib
import io
import random
import zipfile
from urllib.parse import unquote

import httpx
import pytest
from openpyxl import load_workbook
from PIL import Image

from auto_eval.web import exports, history, server
from auto_eval.web.tasks import Task


def image_snapshot(tmp_path, *, mode="compare"):
    paths = []
    for index in range(3):
        path = tmp_path / f"image-{index}.png"
        Image.frombytes("RGB", (100, 240), random.Random(index).randbytes(100 * 240 * 3)).save(path)
        paths.append(str(path))
    item = {"id": "case-1", "query": "题目", "product_count": 3,
            "evidence_mode": "long_screenshot", "session_id": "session-1", "turn_index": 1,
            "query_images": paths[:2],
            "query_image_meta": [{"original_path": path} for path in paths[:2]],
            **{f"screenshot{n}": path for n, path in enumerate(paths, 1)}}
    return {"task_id": "a", "mode": mode, "items": [item], "results": [
        {"index": 0, "answer1_understanding_score": 4, "answer2_understanding_score": 3}],
        "export_base_url": "https://eval.example", "dataset_name": "sample.jsonl"}


@pytest.mark.parametrize("mode", ["compare", "rich_content"])
def test_no_images_preserves_data_links_and_skips_embedding(tmp_path, monkeypatch, mode):
    snapshot = image_snapshot(tmp_path, mode=mode)
    full = history.build_xlsx(snapshot)
    with zipfile.ZipFile(io.BytesIO(full)) as archive:
        embedded = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(embedded) == 3
        original_hashes = {hashlib.sha256((tmp_path / f"image-{n}.png").read_bytes()).hexdigest() for n in range(3)}
        assert {hashlib.sha256(archive.read(name)).hexdigest() for name in embedded} == original_hashes

    def forbidden(*args, **kwargs):
        raise AssertionError("Image bytes must not be read or embedded in the no-image export")
    monkeypatch.setattr(history.WpsCellImages, "add", forbidden)
    small = history.build_xlsx(snapshot, include_images=False)
    assert len(small) < len(full)
    with zipfile.ZipFile(io.BytesIO(small)) as archive:
        assert not any(name.startswith("xl/media/") or "cellimage" in name.lower() for name in archive.namelist())
    small_book = load_workbook(io.BytesIO(small))
    full_book = load_workbook(io.BytesIO(full))
    assert small_book.sheetnames == full_book.sheetnames
    for sheet in small_book:
        full_values = list(full_book[sheet.title].values)
        if sheet.title != "逐题结果":
            assert list(sheet.values) == full_values
            continue
        headers = full_values[0]
        for row_number, row in enumerate(sheet.values):
            if row_number == 0:
                small_headers = row
                assert "产品1原图" not in small_headers
                assert "输入图片原图" not in small_headers
            assert list(row) == [full_values[row_number][headers.index(key)] for key in small_headers]
    result = small_book["逐题结果"]
    headers = [cell.value for cell in result[1]]
    link = result.cell(2, headers.index("产品1长截图查看链接") + 1)
    assert link.hyperlink.target == "https://eval.example/api/eval/a/items/0/screenshots/1"


@pytest.mark.asyncio
async def test_export_apis_freeze_image_option_and_deduplicate_separately(tmp_path, monkeypatch):
    snapshot = image_snapshot(tmp_path)
    task = Task(id="a", mode="compare", items=snapshot["items"], options={}, results=snapshot["results"])
    task.dataset_name = snapshot["dataset_name"]
    manager = exports.XlsxExports(tmp_path / "exports")
    async def peek(_):
        return task
    monkeypatch.setattr(exports, "peek_task_async", peek)
    monkeypatch.setattr(server, "peek_task_async", peek)
    monkeypatch.setattr(server, "XLSX_EXPORTS", manager)
    monkeypatch.setattr(server, "RUNS_DIR", tmp_path)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url="http://test") as client:
        await manager.slot.acquire()
        try:
            full = (await client.post("/api/eval/a/exports")).json()
            small = (await client.post("/api/eval/a/exports?include_images=false")).json()
            repeat = (await client.post("/api/eval/a/exports?include_images=false")).json()
            assert full["export_id"] != small["export_id"] == repeat["export_id"]
            assert full["include_images"] is True
            assert small["include_images"] is False
            assert (await client.post("/api/eval/a/exports?include_images=invalid")).status_code == 422
        finally:
            manager.slot.release()
        await asyncio.gather(*list(manager.workers))
        for state, expected in ((full, True), (small, False)):
            assert manager.view(state["export_id"])["include_images"] is expected
            response = await client.get(f"/api/exports/{state['export_id']}/download")
            assert response.status_code == 200
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                assert any(name.startswith("xl/media/") for name in archive.namelist()) is expected
            assert ("不含原图" in unquote(response.headers["content-disposition"])) is (not expected)
        for query, expected in (("", True), ("&include_images=false", False)):
            response = await client.get("/api/eval/a/export?format=xlsx" + query)
            assert response.status_code == 200
            with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                assert any(name.startswith("xl/media/") for name in archive.namelist()) is expected
    await manager.close()

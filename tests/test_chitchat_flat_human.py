"""Seven-dimension human annotations retain independent scores and shared reasons."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.web.human_baselines import HumanError, HumanStore, ImportMapping, atomic_json, baseline_protocol, read_json
from auto_eval.web.human_baseline_import import build_human_template, inspect_workbook, parse_human_labels
from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_compare import ComparisonRequest, GenerateRequest, HumanComparisons, align_baseline_to_run
from auto_eval.web.human_report_export import export_report
from auto_eval.web.human_statistics import build_human_statistics
from test_human_compare import example, put_baseline

PROFILE = "chitchat_evidence_first_flat@1.1"


def flat_example(count=2):
    task, snapshot, baseline, mapping = example([[3] * count, [5] * count])
    protocol = resolve_compare_protocol(PROFILE)
    snapshot.update(evaluation_profile=protocol.id, protocol_manifest=protocol.public_metadata())
    baseline.update(human_standard_id=protocol.standard_id, human_standard_version=protocol.standard_version,
                    human_bundle_revision=protocol.bundle_revision, human_result_schema_version=protocol.result_schema_version,
                    dimensions=list(protocol.dimensions), score_dimensions=list(protocol.score_dimensions),
                    dimension_names=protocol.dimension_names, score_range=[1, 5])
    for result in snapshot["results"]:
        result.update(standard_id=protocol.standard_id, standard_version=protocol.standard_version,
                      evaluation_profile=protocol.id, bundle_revision=protocol.bundle_revision,
                      result_schema_version=protocol.result_schema_version)
        for dimension in protocol.dimensions:
            result[dimension + "_reason"] = "；".join(f"产品{n}：可定位事实及其本维影响。" for n in range(1, count + 1))
            for n in range(1, count + 1):
                value = (1 if n == 1 else 5) if dimension in protocol.dimensions[:2] else 5
                result[f"answer{n}_{dimension}_score"] = value
                result[f"answer{n}_{dimension}_score_status"] = "scored"
        for n in range(1, count + 1):
            result.pop(f"answer{n}_response_gate", None)
            result.pop(f"answer{n}_safety_gate", None)
    baseline["labels"] = [{**label, "dimension_id": dimension,
                            "score": 1 if dimension in protocol.dimensions[:2] and label["product_id"] == "p1" else 5,
                            "applicable": True}
                           for label in baseline["labels"] for dimension in protocol.dimensions]
    return task, snapshot, baseline, mapping


def filled_template(tmp_path, fields, count=2):
    _, snapshot, _, _ = flat_example(count)
    path = tmp_path / "flat-human.xlsx"
    path.write_bytes(build_human_template(snapshot, PROFILE))
    book = load_workbook(path)
    sheet = book["人工标注"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert not any("Gate" in name or "是否适用" in name or "核验状态" in name for name in columns)
    assert all(sheet.cell(2, col).value is None for name, col in columns.items() if name.startswith("人工"))
    for key, value in fields.items():
        sheet.cell(2, columns[key], value)
    book.save(path)
    book.close()
    mapping = inspect_workbook(path)["suggested_mapping"]
    mapping["policy_note"] = "闲聊七维独立标注，先证据再分数。"
    return path, ImportMapping(**mapping)


@pytest.mark.parametrize("count", [2, 3])
def test_human_template_roundtrip_preserves_seven_dimensions_and_missing_evidence(tmp_path, count):
    path, mapping = filled_template(tmp_path, {
        "人工分_产品1_响应体验": 1, "人工分_产品1_安全稳定": 1,
        "人工分_产品1_理解需求": 5, "人工分_产品2_理解需求": 3,
        "人工评分状态_产品1_风格恰当": "unverifiable",
    }, count)
    parsed = parse_human_labels(path, mapping)
    assert not parsed["issues"]
    lookup = {(r["product_id"], r["dimension_id"]): r for r in parsed["labels"] if r["case_key"].endswith("case-0")}
    assert lookup["product1", "response_experience"]["score"] == 1
    assert lookup["product1", "understanding"]["score"] == 5
    assert lookup["product1", "style"]["score"] is None
    store = HumanStore(tmp_path / "store")
    source = store.path("imports", "hi_flat", "source.xlsx")
    source.parent.mkdir(parents=True)
    source.write_bytes(path.read_bytes())
    atomic_json(store.path("imports", "hi_flat", "preview.json"), parsed)
    manifest = store.publish_baseline("hi_flat", parsed["preview_sha256"], "闲聊七维人工基准")
    assert manifest["allowed_scores_by_dimension"]["response_experience"] == [1, 5]
    assert manifest["allowed_scores_by_dimension"]["understanding"] == [1, 2, 3, 4, 5]
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    path.write_bytes(export_baseline_answers(baseline))
    remapping = inspect_workbook(path)["suggested_mapping"]
    remapping["policy_note"] = "同一冻结标准重导入"
    second = parse_human_labels(path, ImportMapping(**remapping))
    assert not second["issues"]
    assert len(second["labels"]) == len(parsed["labels"])
    assert any(r["score"] is None and r["score_status"] == "unverifiable" for r in second["labels"])


@pytest.mark.parametrize("fields", [
    {"人工分_产品1_响应体验": 3}, {"人工分_产品1_安全稳定": 4},
    {"人工分_产品1_理解需求": "5"}, {"人工分_产品1_理解需求": True},
    {"人工分_产品1_理解需求": 2.5}, {"人工分_产品1_理解需求": 0},
    {"人工评分状态_产品1_理解需求": "not_applicable"},
    {"人工评分状态_产品1_理解需求": "gate_blocked"},
    {"人工评分状态_产品1_理解需求": "blocked_understanding"},
])
def test_human_import_rejects_invalid_scores_and_legacy_blocking_states(tmp_path, fields):
    path, mapping = filled_template(tmp_path, fields)
    assert parse_human_labels(path, mapping)["issues"]


def test_human_baseline_requires_schema_and_rejects_gate_mapping(tmp_path):
    _, _, baseline, _ = flat_example()
    baseline.pop("human_result_schema_version")
    with pytest.raises(HumanError, match="Schema"):
        baseline_protocol(baseline)
    _, mapping = filled_template(tmp_path, {})
    raw = mapping.model_dump()
    raw["gates"] = {"product1": {"response": "Z"}}
    with pytest.raises(ValueError, match="Gate"):
        ImportMapping(**raw)


def test_alignment_report_keeps_low_binary_score_and_other_dimension_high_score(tmp_path):
    _, snapshot, baseline, mapping = flat_example()
    dimensions = list(resolve_compare_protocol(PROFILE).dimensions)
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0", "case-1"], dimensions, "same_standard")
    assert all(row["comparable"] for row in rows)
    p1 = [r for r in rows if r["product_id"] == "p1"]
    assert next(r for r in p1 if r["dimension_id"] == "safety_stability")["model_score"] == 1
    assert next(r for r in p1 if r["dimension_id"] == "understanding")["model_score"] == 5
    assert all(r["model_evidence_scope"] == "merged_in_reason" and r["model_reason_scope"] == "dimension_shared" for r in rows)
    assert all(r["model_gates"] == {} for r in rows)
    stats = build_human_statistics(rows, ["test"], ["p1", "p2"], dimensions)
    assert all(r["kind"] == "score_status" for r in stats["states"])
    assert stats["score_policy"]["response_safety_scores"] == [1, 5]
    service = HumanComparisons(HumanStore(tmp_path / "store"))
    put_baseline(service.store, baseline)
    preview = service.create_preview(ComparisonRequest(baseline_id="hb_test", version=1, tasks=[mapping]), [snapshot])
    manifest = service.prepare_report(GenerateRequest(preview_id=preview["preview_id"], config_sha256=preview["config_sha256"]))
    service.generate(manifest["report_id"])
    report = read_json(service.store.path("reports", manifest["report_id"], "report.json"))
    book = load_workbook(BytesIO(export_report(report)), read_only=True)
    headers = next(book["匹配与排除"].values)
    assert any("模型证据与理由（维度共享完整原文）" in str(h) for h in headers)
    assert not any("Gate" in str(h) or "核验状态" in str(h) or "适用性" in str(h) for h in headers)
    book.close()


def test_input_failure_only_excludes_matching_product_in_human_comparison():
    _, snapshot, baseline, mapping = flat_example()
    snapshot["results"][0]["answer2_input_status"] = "failed"
    rows = align_baseline_to_run(baseline, snapshot, mapping, ["case-0"], ["understanding"], "same_standard")
    assert rows[0]["comparable"] and rows[0]["model_score"] == 5
    assert not rows[1]["comparable"] and rows[1]["model_score_status"] == "input_failed"

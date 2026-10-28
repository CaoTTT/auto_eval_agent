"""Frozen task identity, per-product summaries and every export entry point."""
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from auto_eval.judges.compare_protocols import resolve_compare_protocol
from auto_eval.judges.qa_chat_schema import QA_CHAT_DIMENSIONS
from auto_eval.web.history import export_rows, result_export_row, snapshot_payload, build_xlsx
from auto_eval.web.protocol_identity import recover_protocol_identity, require_snapshot_protocol
from auto_eval.web.runner import _summarize
from auto_eval.web.tasks import Task, _task_from_snapshot


def sample_result(index=0, scores=(5, 3, 4)):
    row = dict(index=index, product_count=len(scores), standard_id="qa_chat_compare", standard_version="1.0",
               evaluation_profile="qa_chat_compare@1.0", bundle_revision="1.0.0", has_conflict="unclear")
    for p, score in enumerate(scores, 1):
        row.update({f"answer{p}_input_status": "complete", f"answer{p}_response_gate": "unclear",
                    f"answer{p}_safety_gate": "pass"})
        for dimension in QA_CHAT_DIMENSIONS:
            prefix = f"answer{p}_{dimension}"
            row.update({f"{prefix}_applicable": True, f"{prefix}_verification_status": "verified" if score else "unverifiable",
                        f"{prefix}_score_status": "scored" if score else "unverifiable", f"{prefix}_score": score,
                        f"{prefix}_reason": f"产品{p}{dimension}证据", f"{prefix}_evidence": ["当前轮可见截图"],
                        f"{prefix}_primary_issue": ""})
    return row


def snapshot(results=None):
    return {"task_id": "chat-synthetic", "mode": "compare", "evaluation_profile": "qa_chat_compare@1.0",
            "protocol_manifest": resolve_compare_protocol("qa_chat_compare@1.0").public_metadata(),
            "items": [{"id": "one", "query": "我有点累，想聊聊天", "product_count": 3,
                       "source_data": {"产品3备注": "保留业务源字段"}}], "results": results or []}


def test_chat_identity_roundtrip_and_recorded_result_revision_recovery():
    source = snapshot([sample_result()])
    del source["protocol_manifest"]["bundle_revision"]
    restored = snapshot_payload(json.loads(json.dumps(source)))
    assert restored["protocol_manifest"]["bundle_revision"] == "1.0.0"
    task = _task_from_snapshot(source, "id")
    assert task.protocol_manifest["bundle_revision"] == "1.0.0"
    assert require_snapshot_protocol(restored).id == "qa_chat_compare@1.0"


@pytest.mark.parametrize("change", ["missing", "unknown", "conflict", "family"])
def test_broken_identity_is_readable_but_never_executable(change):
    source = snapshot([sample_result()])
    if change == "missing":
        del source["protocol_manifest"]["bundle_revision"]
        del source["results"][0]["bundle_revision"]
    elif change == "unknown":
        source["protocol_manifest"]["bundle_revision"] = "9.0.0"
        source["results"][0]["bundle_revision"] = "9.0.0"
    elif change == "conflict":
        source["results"][0]["bundle_revision"] = "9.0.0"
    else:
        source["results"][0]["standard_id"] = "qa_competitor_compare"
    read = snapshot_payload(source)
    assert read["results"] == source["results"]
    assert read["protocol_diagnostics"]
    assert not read["protocol_executable"]
    with pytest.raises(ValueError):
        require_snapshot_protocol(source)
    rows = export_rows(source)["逐题结果"]
    assert rows[0]["协议恢复诊断"]
    assert rows[0]["产品1理解需求分"] is None


def test_summary_keeps_other_products_when_third_input_fails():
    row = sample_result()
    row["answer3_input_status"] = "failed"
    task = Task("summary", "compare", snapshot()["items"], {}, evaluation_profile="qa_chat_compare@1.0",
                protocol_manifest=snapshot()["protocol_manifest"], results=[row])
    summary = _summarize(task)
    assert summary["understanding_answer1_avg"] == 5
    assert summary["understanding_answer2_avg"] == 3
    assert summary["understanding_answer3_avg"] is None
    assert summary["understanding_answer1_answer2_paired_count"] == 1
    assert summary["understanding_all_products_scored"] == 0
    assert summary["comparable"] == 1
    assert "service_closure_answer1_avg" not in summary


def test_chat_export_failure_columns_nulls_and_business_session():
    source = snapshot([dict(index=0, error="judge parse error")])
    source["items"][0].update(product_count=2, session_id="business-id")
    sheets = export_rows(source)
    row = sheets["逐题结果"][0]
    assert "产品1个性化适配评分状态" in row
    assert "产品1内容准确性分" not in row
    assert "产品3理解需求分" not in row
    assert row["产品3备注"] == "保留业务源字段"
    assert row["产品1理解需求分"] is None
    assert row["评估状态"] == "评估失败"
    single = result_export_row("compare", source["results"][0], 0, source["items"], source["evaluation_profile"], source["protocol_manifest"])
    assert single["评测协议"] == "qa_chat_compare@1.0"
    assert "产品1情感表达恰当分" in single
    assert "多轮会话统计" not in sheets


def test_chat_xlsx_preserves_numeric_scores_and_blank_nulls():
    source = snapshot([sample_result(scores=(5, None))])
    source["items"][0]["product_count"] = 2
    book = load_workbook(BytesIO(build_xlsx(source, include_images=False)))
    sheet = book["逐题结果"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("产品1理解需求分") + 1).value == 5
    assert sheet.cell(2, headers.index("产品2理解需求分") + 1).value in (None, "")
    assert sheet.cell(2, headers.index("产品2理解需求评分状态") + 1).value == "无法核验"
    assert "产品3备注" in headers
    assert "产品3理解需求分" not in headers


@pytest.mark.parametrize("value", [0, 6, "4", 4.0, True])
def test_export_does_not_turn_corrupt_history_into_valid_numbers(value):
    source = snapshot([sample_result()])
    source["results"][0]["answer1_understanding_score"] = value
    row = export_rows(source)["逐题结果"][0]
    assert row["产品1理解需求分"] is None
    assert row["产品1理解需求原始异常分"] == value


def test_manifest_internal_conflicts_and_unknown_family_remain_diagnostic():
    source = snapshot()
    source["protocol_manifest"]["evaluation_profile"] = "qa_competitor_compare@0.3"
    assert not recover_protocol_identity(source)["executable"]
    unknown = {"mode": "compare", "protocol_manifest": {"standard_id": "future_family"}}
    assert recover_protocol_identity(unknown)["evaluation_profile"] == "future_family@unknown"
    assert not recover_protocol_identity(unknown)["executable"]


def test_invalid_identity_never_reuses_cached_official_statistics():
    source = snapshot([sample_result()])
    source["protocol_manifest"]["bundle_revision"] = "9.9"
    source["summary"] = {"total": 1, "understanding_answer1_avg": 5}
    read = snapshot_payload(source)
    assert "understanding_answer1_avg" not in read["summary"]
    assert read["raw_summary"]["understanding_answer1_avg"] == 5
    sheets = export_rows(source)
    assert "understanding_answer1_avg" not in sheets["汇总指标"][0]
    assert sheets["原始汇总（协议待恢复）"]


@pytest.mark.asyncio
async def test_single_item_api_keeps_task_wide_identity_conflict(monkeypatch):
    from auto_eval.web import server
    source = snapshot([sample_result()])
    other = sample_result(index=1)
    other["bundle_revision"] = "9.9"
    source["results"].append(other)
    task = _task_from_snapshot(source, "chat-synthetic")
    async def load(_task_id):
        return task
    monkeypatch.setattr(server, "get_task_async", load)
    response = await server.api_item_result(task.id, "one")
    assert response["result"]["产品1理解需求分"] is None
    assert response["result"]["协议恢复诊断"]
    task.status = "done"
    task.summary = {"understanding_answer1_avg": 5}
    monkeypatch.setattr(server, "peek_task", lambda _task_id: task)
    stream = await server.api_stream(task.id, compact=True)
    chunks = [chunk async for chunk in stream.body_iterator]
    assert "understanding_answer1_avg" not in "".join(chunks)


def test_cross_standard_side_by_side_export_stays_available():
    from copy import deepcopy
    from auto_eval.web.comparison_export import comparison_sheets
    left = snapshot([sample_result()])
    right = deepcopy(left)
    right.update(task_id="old", evaluation_profile="qa_competitor_compare@0.3",
                 protocol_manifest=resolve_compare_protocol("qa_competitor_compare@0.3").public_metadata())
    right["results"] = [{"index": 0, "standard_id": "qa_competitor_compare", "standard_version": "0.3",
                         "bundle_revision": "0.3.1", "answer1_understanding_score": 0}]
    sheets = comparison_sheets(left, right)
    assert sheets["任务说明"][0]["评测协议"] == "qa_chat_compare@1.0"
    assert sheets["任务说明"][1]["评测协议"] == "qa_competitor_compare@0.3"
    assert "产品1个性化适配分 · A" in sheets["逐题对比"][0]


@pytest.mark.asyncio
async def test_restored_business_session_uses_single_turn_media_preparation(monkeypatch):
    from types import SimpleNamespace
    from auto_eval.config import AppConfig, JudgeConfig, VisualModeProfile
    from auto_eval.web import runner
    source = snapshot()
    source["items"][0].update(session_id="business", product_count=2, video1="a.mp4", video2="b.mp4")
    task = _task_from_snapshot(source, "business-task")
    prepared = []
    def prepare(item, **_kwargs):
        prepared.append(item["id"])
        return {**item, "frames1": ["a.png"], "frames2": ["b.png"]}
    async def evaluate(_mode, index, _item, **kwargs):
        assert not kwargs.get("conversation_input")
        return sample_result(index, scores=(5, 3))
    monkeypatch.setattr(runner, "JudgeClient", lambda *_: SimpleNamespace())
    monkeypatch.setattr(runner, "prepare_session_visual_compare_item", prepare)
    monkeypatch.setattr(runner, "_eval_one", evaluate)
    monkeypatch.setattr(runner, "_persist_task", lambda *_: None)
    cfg = AppConfig(judges=[JudgeConfig(name="fake")], visual_modes={"rich_content": VisualModeProfile(extraction={"algorithm_version": "test"})})
    one, _ = runner._make_item_evaluator(task, cfg)
    result = await one(0, task.items[0])
    assert not result.get("error"), result
    assert prepared == ["one"]


@pytest.mark.asyncio
async def test_create_current_chat_task_api_freezes_complete_candidate_manifest(monkeypatch):
    from auto_eval.config import AppConfig, JudgeConfig
    from auto_eval.web import server
    cfg = AppConfig(judges=[JudgeConfig(name="fake")])
    monkeypatch.setattr(server, "cfg", lambda: cfg)
    def create(mode, items, options, dataset_name="", **kwargs):
        return Task("created-chat", mode, items, options, evaluation_profile=kwargs["evaluation_profile"],
                    protocol_manifest=kwargs["protocol_manifest"], judge_runtime=kwargs["judge_runtime"])
    monkeypatch.setattr(server, "new_task", create)
    monkeypatch.setattr(server.EVAL_SCHEDULER, "enqueue", lambda *_: 1)
    response = await server.api_eval(server.EvalReq(mode="compare", evaluation_profile="qa_chat_compare@1.1",
        items=[{"id": "new", "query": "聊聊天", "screenshot1": "a.png", "screenshot2": "b.png",
                "context1": "产品1背景", "context2": "产品2背景"}]))
    manifest = response["protocol_manifest"]
    assert manifest["standard_id"] == "qa_chat_compare" and manifest["bundle_revision"] == "1.1.0"
    assert manifest["score_range"] == [1, 5]
    assert manifest["score_dimensions"] == list(QA_CHAT_DIMENSIONS)

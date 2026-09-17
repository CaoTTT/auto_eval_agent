from io import BytesIO
import threading
from types import SimpleNamespace
from urllib.parse import unquote

from fastapi import FastAPI
import httpx
from openpyxl import load_workbook
import pytest

from auto_eval.web.human_baseline_export import export_baseline_answers
from auto_eval.web.human_baseline_import import inspect_workbook, parse_human_labels
from auto_eval.web.human_baselines import HumanStore, ImportMapping, atomic_json
from auto_eval.web.human_routes import install_human_routes
from test_human_baseline_import import workbook
from test_human_compare import example, put_baseline


def _rows(data):
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb["人工基准答案"]
    headers = [c.value for c in ws[1]]
    rows = [{header: ws.cell(row, column + 1).value for column, header in enumerate(headers)}
            for row in range(2, ws.max_row + 1)]
    return wb, ws, rows


@pytest.mark.parametrize("count", [2, 3])
def test_baseline_answers_wide_rows_keep_values_states_and_missing(count):
    _, _, baseline, _ = example([[2] * count, [3] * count], standard="0.3")
    first, second = baseline["cases"]
    first.update(case_id="001", query="=SUM(A1:A2)", context="line\x01break", query_hashes=["image-sha"])
    first["responses"]["p1"].update(answer="+test", evidence="image.png", response_id="response-001")
    baseline["labels"][0].update(score=0, reason="=HYPERLINK(\"https://example.org\")", annotation_stage="review",
                                 source={"selected_cell": "D2", "base_raw": 3, "review_raw": 0})
    baseline["labels"][1].update(score=None, score_status="unverifiable", reason="无法确认")
    baseline["labels"].extend([
        dict(case_key=first["case_key"], product_id="p1", dimension_id="evidence_quality", score=2,
             score_status="scored", annotation_stage="base", reason="证据完整", source={"selected_cell": "M2"}),
        dict(case_key=first["case_key"], product_id="p2", dimension_id="accuracy", score=3,
             score_status="scored", annotation_stage="base", reason="仅留档", source={}),
        dict(case_key=second["case_key"], product_id="p1", dimension_id="evidence_quality", score=None,
             score_status="not_applicable", annotation_stage="review", reason="无证据需求", source={}),
    ])
    baseline["states"] = [dict(case_key=first["case_key"], applicability={"understanding": True, "service_closure": False},
                              gates={"p1": {"response": "pass", "safety": "unclear"}})]
    wb, ws, rows = _rows(export_baseline_answers(baseline))
    assert wb.sheetnames == ["人工基准答案", "填写说明"]
    assert len(rows) == len(baseline["cases"]) == 2
    assert rows[0]["case_id"] == "001"
    assert rows[0]["query"] == "=SUM(A1:A2)"
    assert rows[0]["公共背景"] == "line\\u0001break"
    assert rows[0]["源回答_产品1"] == "+test"
    assert rows[0]["人工分_产品1_理解需求"] == 0
    assert rows[0]["人工标注阶段_产品1_理解需求"] == "review"
    assert '"review_raw": 0' in rows[0]["人工评分来源_产品1_理解需求"]
    assert rows[0]["人工分_产品2_理解需求"] == "N/A"
    assert rows[0]["人工评分状态_产品2_理解需求"] == "unverifiable"
    assert rows[0]["人工分_产品1_有理有据"] == 2
    assert rows[0]["人工分_产品2_内容准确性"] == 3
    assert rows[0]["人工分_产品2_有理有据"] is None
    assert rows[0]["人工评分状态_产品2_有理有据"] == "unlabeled"
    assert rows[1]["人工分_产品1_有理有据"] == "N/A"
    assert rows[1]["人工评分状态_产品1_有理有据"] == "not_applicable"
    assert rows[0]["人工响应Gate_产品1"] == "pass"
    assert rows[0]["人工安全Gate_产品1"] == "unclear"
    assert rows[0]["人工是否适用_理解需求"] is True
    assert rows[0]["人工是否适用_服务闭环"] is False
    assert rows[0]["人工是否适用_有理有据"] is None
    assert ("人工分_产品3_理解需求" in rows[0]) == (count == 3)
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == ws.dimensions
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


def test_published_answer_export_keeps_review_na_and_can_be_remapped(tmp_path):
    source, mapping = workbook(tmp_path, [["001", "q1", 3, 0, 1], ["002", "q2", 2, "N/A", 3]], standard="0.3")
    preview = parse_human_labels(source, mapping)
    store = HumanStore(tmp_path / "runs")
    imported = store.path("imports", "hi_answers", "source.xlsx")
    imported.parent.mkdir(parents=True)
    imported.write_bytes(source.read_bytes())
    atomic_json(store.path("imports", "hi_answers", "preview.json"), preview)
    manifest = store.publish_baseline("hi_answers", preview["preview_sha256"], "人工答案")
    baseline = store.load_baseline(manifest["baseline_id"], 1)
    payload = export_baseline_answers(baseline)
    _, _, rows = _rows(payload)
    assert rows[0]["人工分_产品1_理解需求"] == 0
    assert rows[1]["人工分_产品1_理解需求"] == "N/A"
    path = tmp_path / "downloaded.xlsx"
    path.write_bytes(payload)
    info = inspect_workbook(path)
    remapping = info["suggested_mapping"]
    remapping["policy_note"] = "重导入已保存答案"
    result = parse_human_labels(path, ImportMapping(**remapping))
    assert result["issues"] == []
    answers = {(label["case_key"].split(":", 1)[1], label["product_id"], label["dimension_id"]):
               (label["score"], label["score_status"]) for label in result["labels"]}
    assert answers["001", "product1", "understanding"] == (0, "scored")
    assert answers["002", "product1", "understanding"] == (None, "na_unspecified")


def _app(store):
    app = FastAPI()

    async def peek(task_id):
        raise AssertionError("Downloads must not access evaluation tasks")

    install_human_routes(app, lambda: SimpleNamespace(store=store), peek)
    return app


@pytest.mark.asyncio
async def test_baseline_download_reads_requested_version_and_returns_errors(tmp_path):
    _, _, baseline, _ = example()
    store = HumanStore(tmp_path)
    put_baseline(store, baseline)
    second = dict(baseline, version=2, labels=[dict(row, score=5) for row in baseline["labels"]])
    folder = store.path("baselines", baseline["baseline_id"], "v2")
    for key in ("cases", "labels", "states"):
        atomic_json(folder / f"{key}.json", second[key])
    atomic_json(folder / "manifest.json", {k: v for k, v in second.items() if k not in ("cases", "labels", "states")})
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=_app(store)), base_url="http://test") as client:
        first_response = await client.get("/api/human-baselines/hb_test/versions/1/download")
        second_response = await client.get("/api/human-baselines/hb_test/versions/2/download")
        assert first_response.status_code == second_response.status_code == 200
        assert "spreadsheetml.sheet" in first_response.headers["content-type"]
        assert "人工参考_人工基准v1_答案.xlsx" in unquote(first_response.headers["content-disposition"])
        assert _rows(first_response.content)[2][0]["人工分_产品1_理解需求"] == 3
        assert _rows(second_response.content)[2][0]["人工分_产品1_理解需求"] == 5
        assert (await client.get("/api/human-baselines/hb_missing/versions/1/download")).status_code == 404
        assert (await client.get("/api/human-baselines/hb_test/versions/99/download")).status_code == 404
        assert (await client.get("/api/human-baselines/hb_test/versions/0/download")).status_code == 422
        assert (await client.get("/api/human-baselines/hb_test/versions/text/download")).status_code == 422


@pytest.mark.asyncio
async def test_existing_report_download_renders_frozen_json_in_worker_without_mutation(tmp_path, monkeypatch):
    import auto_eval.web.human_routes as routes

    store = HumanStore(tmp_path)
    manifest = dict(status="ready", baseline_name="已保存基准", version=2)
    atomic_json(store.path("reports", "hc_existing"), manifest)
    frozen_path = store.path("reports", "hc_existing", "report.json")
    frozen = {"frozen": True, "metrics": {"score": 0.5}}
    atomic_json(frozen_path, frozen)
    original_json = frozen_path.read_bytes()
    old_xlsx = store.path("reports", "hc_existing", "report.xlsx")
    old_xlsx.write_bytes(b"old layout")
    atomic_json(store.path("reports", "hc_queued"), dict(manifest, status="queued"))

    def render(report):
        assert threading.current_thread() is not threading.main_thread()
        assert report == frozen
        return b"current layout"

    monkeypatch.setattr(routes, "export_report", render)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=_app(store)), base_url="http://test") as client:
        response = await client.get("/api/human-comparisons/hc_existing/download")
        assert response.status_code == 200
        assert response.content == b"current layout"
        assert "已保存基准_人工基准v2_人机对比_hc_existing.xlsx" in unquote(response.headers["content-disposition"])
        assert (await client.get("/api/human-comparisons/hc_queued/download")).status_code == 409
        assert (await client.get("/api/human-comparisons/hc_missing/download")).status_code == 404
    assert frozen_path.read_bytes() == original_json
    assert old_xlsx.read_bytes() == b"old layout"
